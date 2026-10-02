"""Check Comment 1 values, preserved figures, references and exported pagination."""
import argparse
import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".codex_tmp/public_dataset_paper"


def check(evidence, render=False):
    import pymupdf
    with (evidence / "summary.csv").open(newline="") as f:
        rows = list(csv.DictReader(f))
    lookup = {(r["scenario"], r["model"]): r for r in rows}
    plotted = json.loads((WORK / "plot_values.json").read_text())
    for point in plotted:
        r = lookup[(point["scenario"], point["model"])]
        for value, suffix in (("median", "median"), ("p25", "p25"), ("p75", "p75")):
            assert point[value] == float(r[f"{point['panel']}_{suffix}"])
        assert point["trials"] == int(r["trials"])
    integration = json.loads((WORK / "integration.json").read_text())
    checks = []
    for item in integration["outputs"]:
        source, pdf = Path(item["docx"]), Path(item["pdf"])
        backup = WORK / "backups" / source.relative_to(ROOT)
        doc, old = Document(source), Document(backup)
        assert doc.paragraphs[0].text == old.paragraphs[0].text
        assert [p.text for p in doc.paragraphs[1:6]] == [p.text for p in old.paragraphs[1:6]]
        assert len(doc.sections) == len(old.sections), "unintended section added"
        assert len(doc.inline_shapes) == len(old.inline_shapes) + 1
        new_image_paragraph = next(p for p in doc.paragraphs if p._p.xpath(".//w:drawing") and p._p.xpath(".//wp:extent")[-1].get("cy") == str(doc.inline_shapes[-1].height))
        assert new_image_paragraph.paragraph_format.line_spacing == 1, "new figure must not inherit exact text line height"
        with zipfile.ZipFile(backup) as before, zipfile.ZipFile(source) as after:
            for name in before.namelist():
                if name.startswith("word/media/"):
                    assert before.read(name) == after.read(name), f"original figure modified: {name}"
        bert = "BERT" in source.name
        has_reference = lambda p: any(b.get(qn("w:name"), "").startswith("ref_") for b in p._p.xpath(".//w:bookmarkStart"))
        refs = [p for p in doc.paragraphs if has_reference(p)]
        assert len(refs) == 31
        if not bert:
            assert [int(re.match(r"^\[(\d+)\]", p.text).group(1)) for p in refs] == list(range(1, 32))
        old_refs = [p.text for p in old.paragraphs if has_reference(p)]
        assert [p.text for p in refs[:30]] == old_refs, "original references changed"
        anchors = {e.get(qn("w:anchor")) for e in doc._element.xpath(".//w:hyperlink")}
        for p in refs:
            bookmark = next(e.get(qn("w:name")) for e in p._p.xpath(".//w:bookmarkStart") if e.get(qn("w:name"), "").startswith("ref_"))
            assert bookmark in anchors, f"uncited reference {bookmark}"
        assert len(doc.tables) == len(old.tables) + 2
        assert len(doc.tables[2].rows) == 10 and len(doc.tables[2].columns) == 2
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "synthetic correction histories" in text and "1,000,000 records" in text
        assert "2,964,624 rows" in text and "47,277 missing passenger counts" in text
        large = lookup[("tlc-n1000000-h10-c1000", "Revon-H")]
        for metric in ("incremental_commit_ms", "diff_ms"):
            reported = f"{float(large[metric+'_median']):.2f} [{float(large[metric+'_p25']):.2f}, {float(large[metric+'_p75']):.2f}] ms"
            assert reported in text, f"narrative does not match {metric}"
        for model in ("Revon-H", "Dolt", "Dolt (bulk import)"):
            r = lookup[("tlc-n1000000-h10-c1000", model)]
            reported = f"{float(r['initial_import_ms_median'])/1000:.2f} [{float(r['initial_import_ms_p25'])/1000:.2f}, {float(r['initial_import_ms_p75'])/1000:.2f}] s"
            assert reported in text, f"narrative import does not match {model}"
        ratio = float(large["storage_bytes_median"]) / float(lookup[("tlc-n1000000-h10-c1000", "Dolt")]["storage_bytes_median"])
        assert f"footprints was {ratio:.2f}" in text
        assert "We did not test alternate schemas, public multi-table datasets, or one-million-row datasets." not in text
        assert "Our results do not establish performance on public datasets" not in text
        abstract_text = next(p.text for p in doc.paragraphs if p.text.startswith("Abstract:"))
        assert "Versioning structured datasets requires durable history" not in text, "stale duplicate abstract"
        assert len(abstract_text.split()) < 250
        assert "public TLC record evaluation reaches 1,000,000 records" in abstract_text
        new_table = doc.tables[-1]
        assert [c.text for c in new_table.rows[0].cells] == ["H / density", "System", "Commit", "Diff", "MiB"]
        models = [("Snapshot", "Snapshot"), ("Log-only", "Log-only"), ("Revon-M (forced Merkle)", "Revon-M"), ("Revon-H", "Revon-H"), ("Dolt", "Dolt SQL"), ("Dolt (bulk import)", "Dolt CSV")]
        expected = []
        for scenario, label in (("tlc-n100000-h10-c100", "10 / 0.1%"), ("tlc-n100000-h50-c100", "50 / 0.1%"), ("tlc-n100000-h10-c10", "10 / 0.01%"), ("tlc-n100000-h10-c1000", "10 / 1%")):
            for model, short in models:
                r = lookup[(scenario, model)]
                expected.append([label, short, f"{float(r['incremental_commit_ms_median']):.2f}", f"{float(r['diff_ms_median']):.2f}", f"{float(r['storage_bytes_median'])/2**20:.2f}"])
        assert [[c.text for c in row.cells] for row in new_table.rows[1:]] == expected
        exported = pymupdf.open(pdf)
        pdftext = "\n".join(page.get_text() for page in exported)
        assert ("NYC TLC, 2024" if bert else "[31]") in pdftext and "Public Dataset and Scalability" in pdftext
        caption_pages = {}
        image_matches = {}
        for page_index, page in enumerate(exported):
            image_rectangles = [rect for image in page.get_images(full=True) for rect in page.get_image_rects(image[0])]
            for block in page.get_text("dict")["blocks"]:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    line_text = "".join(span["text"] for span in line["spans"])
                    match = re.match(r"^Fig\.\s*(\d+)\.", line_text.strip())
                    if match:
                        number = int(match.group(1))
                        caption_pages[number] = page_index + 1
                        box = pymupdf.Rect(line["bbox"])
                        candidates = [r for r in image_rectangles if r.y1 <= box.y0 + 12 and box.y0-r.y1 < 90 and r.x0 < box.x1 and r.x1 > box.x0]
                        assert candidates, f"figure {number} caption lacks preceding image on page {page_index+1}"
                        image_matches[number] = tuple(min(candidates, key=lambda r: abs(box.y0-r.y1)))
                        image_rect = pymupdf.Rect(image_matches[number])
                        assert image_rect.y0 >= 0 and image_rect.y1 <= page.rect.height, f"figure {number} outside page"
            if render:
                out = WORK / "qa" / source.stem
                out.mkdir(parents=True, exist_ok=True)
                page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(out / f"page-{page_index+1:02d}.png")
        assert set(caption_pages) == set(range(1, 8)), caption_pages
        tables = {number: [i+1 for i,p in enumerate(exported) if f"TABLE {number}." in p.get_text()] for number in ("I", "II", "III", "IV", "V")}
        assert all(len(pages) == 1 for pages in tables.values()), tables
        result_page = exported[tables["V"][0]-1].get_text()
        assert all(result_page.count(label) == 6 for label in ("10 / 0.1%", "50 / 0.1%", "10 / 0.01%", "10 / 1%")), "new results table split across pages"
        dataset_page = exported[tables["III"][0]-1].get_text()
        dataset_normalized = " ".join(re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", dataset_page).split())
        assert "physical row ordinal" in dataset_normalized and "Sampling" in dataset_normalized, "dataset table split across pages"
        checks.append({"docx": str(source), "pdf": str(pdf), "pages": len(exported),
            "figure_caption_pages": caption_pages, "figure_image_rectangles": image_matches,
            "table_caption_pages": tables, "all_new_table_values_match": True,
            "all_original_figures_preserved": True, "references_cited": 31,
            "reference_style": "author-year" if bert else "numbered",
            "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest()})
        exported.close()
    report = {"passed": True, "plot_values_checked": len(plotted), "manuscripts": checks,
              "visual_review": "PNG inspection still required; structural checks alone do not establish visual quality"}
    (WORK / "qa_report.json").write_text(json.dumps(report, indent=2))
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    check(args.evidence.resolve(), args.render)
