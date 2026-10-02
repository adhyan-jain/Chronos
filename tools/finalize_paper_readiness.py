"""Apply the final bounded-claim and workload-table fixes to the paper DOCX."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper" / "Revon_Final_Research_Paper.docx"
TITLE = "Revon: Evaluating Hybrid Diffing in a Versioned Key-Value Prototype"
ABSTRACT = (
    "Versioning structured datasets requires durable history, efficient sparse "
    "updates, historical reconstruction, and comparison without copying every "
    "record. Revon is a Python prototype built around an immutable fixed-depth "
    "Merkle hash trie; Revon-H adds addressed changesets and adaptively selects "
    "log aggregation or hash-pruned differencing. We compare Snapshot, Log-only, "
    "forced-Merkle Revon-M, Revon-H, and Dolt 2.3.1 on deterministic workloads up "
    "to 100,000 rows, using evidence from 540 successful executions. Within the "
    "declared API and CLI workflows, median paired Dolt/Revon-H diff-latency ratios "
    "range from 7.14 to 625.10, and all seven-trial 95% bootstrap intervals favor "
    "Revon-H latency. These are workflow-level results: the comparison includes "
    "different interfaces and does not isolate index or tree performance. A "
    "separate three-seed study found the 50-commit ablation result inconclusive. "
    "Experiments used one Windows host and WSL2 on the same physical computer. The "
    "results support workload-specific trade-offs for this prototype, not general "
    "or production-level superiority."
)


def set_table_widths(table, widths_dxa: list[int]) -> None:
    if any(width <= 0 for width in widths_dxa):
        raise ValueError("table column widths must be positive")
    if len(table.columns) != len(widths_dxa):
        raise ValueError("unexpected workload-matrix column count")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")

    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = table._tbl.tblGrid.gridCol_lst
    if len(grid_columns) != len(widths_dxa):
        raise ValueError("unexpected workload-matrix grid")
    for grid_column, width in zip(grid_columns, widths_dxa):
        grid_column.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)


def copy_row_format(source_row, destination_row) -> None:
    if source_row._tr.trPr is not None:
        old = destination_row._tr.trPr
        if old is not None:
            destination_row._tr.remove(old)
        destination_row._tr.insert(0, deepcopy(source_row._tr.trPr))

    for source_cell, destination_cell in zip(source_row.cells, destination_row.cells):
        source_properties = source_cell._tc.tcPr
        destination_properties = destination_cell._tc.tcPr
        source_width = destination_properties.find(qn("w:tcW"))
        for child in list(destination_properties):
            if child.tag != qn("w:tcW"):
                destination_properties.remove(child)
        for child in source_properties:
            if child.tag != qn("w:tcW"):
                destination_properties.append(deepcopy(child))
        if source_width is not None:
            destination_properties.remove(source_width)
            destination_properties.insert(0, source_width)

        for source_paragraph, destination_paragraph in zip(
            source_cell.paragraphs, destination_cell.paragraphs
        ):
            if source_paragraph._p.pPr is not None:
                old = destination_paragraph._p.pPr
                if old is not None:
                    destination_paragraph._p.remove(old)
                destination_paragraph._p.insert(
                    0, deepcopy(source_paragraph._p.pPr)
                )
            source_run = next(iter(source_paragraph.runs), None)
            if source_run is not None and source_run._r.rPr is not None:
                for destination_run in destination_paragraph.runs:
                    old = destination_run._r.rPr
                    if old is not None:
                        destination_run._r.remove(old)
                    destination_run._r.insert(0, deepcopy(source_run._r.rPr))


def set_cell_bottom_border(cell) -> None:
    properties = cell._tc.tcPr
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    old_bottom = borders.find(qn("w:bottom"))
    if old_bottom is not None:
        borders.remove(old_bottom)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)


def repair_evaluation_table(table, widths_dxa: list[int]) -> None:
    if len(table.rows) != 9 or len(table.columns) != 5:
        raise ValueError("unexpected evaluation-table shape")
    set_table_widths(table, widths_dxa)
    # The final three rows had been appended without cell formatting, and the
    # preceding body row incorrectly retained the old table-bottom border.
    for row_index in (5, 6, 7, 8):
        copy_row_format(table.rows[1], table.rows[row_index])
    for cell in table.rows[8].cells:
        set_cell_bottom_border(cell)


def main() -> None:
    document = Document(PAPER)
    if len(document.tables) < 3:
        raise ValueError("expected the literature, workload, and results tables")
    if len(document.paragraphs) <= 8:
        raise ValueError("unexpected manuscript paragraph structure")
    if len(document.tables[1].columns) != 5 or len(document.tables[1].rows) != 9:
        raise ValueError("workload matrix no longer matches the reviewed table")

    document.paragraphs[0].text = TITLE
    document.paragraphs[8].text = ABSTRACT

    # The previous 5.90-inch table gave the locality column only 1.25 inches.
    # Use the full 6.26-inch text width and reserve enough space for descriptions.
    # The literature matrix was needlessly narrow compared with the one-column
    # text region. Expanding it makes its comparisons readable without reducing
    # type size.
    set_table_widths(document.tables[0], [2100, 2300, 4620])
    dolt_position = document.tables[0].cell(11, 2).paragraphs[0]
    if dolt_position.runs:
        dolt_position.runs[0].text = "Mature production comparator"
        for run in dolt_position.runs[1:]:
            run.text = ""
    else:
        dolt_position.add_run("Mature production comparator")
    repair_evaluation_table(document.tables[1], [2100, 1100, 800, 900, 4120])
    repair_evaluation_table(document.tables[2], [2050, 750, 900, 1250, 4070])
    document.save(PAPER)
    print(f"Updated title, abstract, and workload-matrix geometry in {PAPER}")


if __name__ == "__main__":
    main()
