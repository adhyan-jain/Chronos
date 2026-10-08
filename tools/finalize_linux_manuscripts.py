"""Build both final manuscripts from audited, separately identified evidence.

Run with the Codex bundled Python. Optional plotting packages are in
.codex_tmp/dataset_deps and are used only for scientific figures.
"""
from __future__ import annotations
import copy
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil
import statistics
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.append(str(ROOT/'.codex_tmp/dataset_deps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import revise_discover_manuscripts as layout
import preview_diff_figure as diff_figure

OUT=ROOT/'output/docs/linux-final-20261008'
WORK=ROOT/'.codex_tmp/linux-final-paper'
MAIN=ROOT/'evidence/linux-validation-20261008'
SUPP=ROOT/'evidence/linux-final-supplement-20261008'
H='Revon-H'; M='Revon-M (forced Merkle)'; D='Dolt'
SCENARIOS=['small-sparse','medium-sparse','medium-dense','large-sparse',
 'large-application-key-local','large-range-local','large-repeated-key','large-hash-route-local',
 'linux-tlc-n100000-h10-c100']
LABELS=['S1','S2','S3','S4','S5','S6','S7','S8','P1']
COLORS={H:'#007C83',M:'#B75D36',D:'#355C99'}

def read(path):return list(csv.DictReader(path.open(encoding='utf-8',newline='')))
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def body(text):return dict(kind='body',text=text)
def heading(text,level=1):return dict(kind='h'+str(level),text=text)
def fig(number,caption):return dict(kind='figure',number=number,caption=caption)
def table(number,caption,headers,rows,widths=None):return dict(kind='table',number=number,caption=caption,headers=headers,rows=rows,widths=widths)
def fmt(value):return f'{float(value):.3f}'

def evidence():
    for folder in (MAIN,SUPP):
        a=json.loads((folder/'audit.json').read_text())
        assert a['passed'],a
    assert json.loads((SUPP/'geometry/audit.json').read_text())['passed']
    rows=read(MAIN/'raw_results.csv')
    assert len(rows)==516 and all(r['status']=='ok' and r['correctness']=='True' for r in rows)
    lookup={(r['scenario'],r['model']):r for r in read(MAIN/'summary.csv')}
    pairs={(r['scenario'],r['numerator_model'],r['metric']):r for r in read(MAIN/'paired_uncertainty.csv')}
    return rows,lookup,pairs

def plots(lookup):
    folder=OUT/'figures';folder.mkdir(parents=True,exist_ok=True)
    assets={}
    plt.rcParams.update({'font.family':'Arial','font.size':10,'axes.spines.top':False,
        'axes.spines.right':False,'axes.labelcolor':'#222222','text.color':'#222222',
        'pdf.fonttype':42,'ps.fonttype':42,'svg.fonttype':'none','savefig.facecolor':'white'})
    for n in (1,):
        for ext in ('png','pdf','svg'):
            shutil.copyfile(ROOT/f'.codex_tmp/seven_revisions/figures/Fig{n}.{ext}',folder/f'Fig{n}.{ext}')
        assets[n]=str(folder/f'Fig{n}.png')
    def save(figure,n):
        figure.tight_layout(pad=1.2)
        for ext in ('png','pdf','svg','eps'):
            figure.savefig(folder/f'Fig{n}.{ext}',dpi=600,bbox_inches='tight')
        plt.close(figure);assets[n]=str(folder/f'Fig{n}.png')
    from matplotlib.patches import FancyBboxPatch
    f,ax=plt.subplots(figsize=(7,2.55));ax.set_xlim(0,10);ax.set_ylim(0,4);ax.axis('off')
    def box(x,y,w,h,text,color):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.06,rounding_size=0.06',
            facecolor='white',edgecolor=color,linewidth=1.2))
        ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=9)
    def arrow(a,b):ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',color='#30434A',lw=1.1))
    box(.1,1.35,1.5,1.0,'Requested\nversion pair','#355C99')
    box(2,1.15,2,1.4,'Validate versions\nand ancestry;\ncount operations','#355C99')
    box(4.5,1.15,2,1.4,'Ancestor path\nexists and\noperations ≤ T?','#007C83')
    box(7,2.45,2.7,1,'Log aggregation\n(reverse if needed)','#007C83')
    box(7,.15,2.7,1,'Merkle comparison\n(prune equal hashes)','#B75D36')
    arrow((1.66,1.85),(1.94,1.85));arrow((4.06,1.85),(4.44,1.85))
    arrow((6.56,2.1),(7,2.9));arrow((6.56,1.6),(7,.7))
    ax.text(6.62,2.8,'Yes',fontsize=8);ax.text(6.6,.85,'No',fontsize=8)
    ax.text(4.1,.35,'Both paths return the same sorted semantic diff',ha='center',fontsize=9)
    save(f,2)
    def points(ax,models,metric,scenarios=SCENARIOS,labels=LABELS,log=True):
        x=np.arange(len(scenarios))
        for i,model in enumerate(models):
            rs=[lookup[s,model] for s in scenarios]
            med=np.array([float(r[metric+'_median']) for r in rs])
            lo=np.array([float(r[metric+'_p25']) for r in rs]);hi=np.array([float(r[metric+'_p75']) for r in rs])
            ax.errorbar(x+(i-(len(models)-1)/2)*.18,med,yerr=[med-lo,hi-med],fmt=['o','s','^'][i],
                markersize=5,elinewidth=1.2,capsize=3,color=COLORS[model],label={M:'Revon-M',H:'Revon-H',D:'Dolt SQL/CLI'}[model])
        ax.set_xticks(x,labels);ax.set_xlabel('Workload ID');ax.grid(axis='y',alpha=.2)
        if log:ax.set_yscale('log')
    with plt.rc_context():
        diff_figure.main(folder, 'Fig3', preview_only=False)
    assets[3]=str(folder/'Fig3.png')
    f,ax=plt.subplots(figsize=(7,3.25))
    threshold=['linux-threshold-h16-c512','linux-threshold-h64-c512']
    for i,model in enumerate(['Revon-log calibration',M,H]):
        rs=[lookup[s,model] for s in threshold]
        med=np.array([float(r['diff_ms_median']) for r in rs]);lo=np.array([float(r['diff_ms_p25']) for r in rs]);hi=np.array([float(r['diff_ms_p75']) for r in rs])
        color=['#756A99',COLORS[M],COLORS[H]][i]
        ax.errorbar(np.arange(2)+(i-1)*.15,med,yerr=[med-lo,hi-med],fmt=['^','s','o'][i],color=color,capsize=4,label=['Forced log','Forced Merkle','Hybrid T = 16,384'][i])
    ax.set_xticks([0,1],['8,192 operations\n16 commits × 512','32,768 operations\n64 commits × 512'])
    ax.set_ylabel('Diff latency (ms)');ax.set_ylim(bottom=0);ax.grid(axis='y',alpha=.2);ax.legend(frameon=False,ncols=3,loc='upper left',fontsize=9);save(f,4)
    f,axes=plt.subplots(1,2,figsize=(7,3.3))
    for ax,metric,label in zip(axes,['incremental_commit_ms','checkout_ms'],['Commit latency (ms, log scale)','Checkout latency (ms, log scale)']):
        points(ax,[H,D],metric);ax.set_ylabel(label)
    axes[0].legend(frameon=False,fontsize=9);save(f,5)
    f,axes=plt.subplots(1,2,figsize=(7,3.4))
    for ax,metric,label,divisor in [(axes[0],'storage_bytes','Repository footprint (MiB)',2**20),(axes[1],'process_tree_peak_rss_bytes','Whole-worker peak RSS (MiB)',2**20)]:
        for i,model in enumerate([M,H,D]):
            vals=[float(lookup[s,model][metric+'_median'])/divisor for s in SCENARIOS]
            ax.plot(LABELS,vals,['s--','o-','^:'][i],color=COLORS[model],label={M:'Revon-M',H:'Revon-H',D:'Dolt'}[model],markersize=4)
        ax.set_yscale('log');ax.set_ylabel(label);ax.set_xlabel('Workload ID');ax.grid(axis='y',alpha=.2)
    axes[0].legend(frameon=False,fontsize=9);save(f,6)
    geometry=read(SUPP/'geometry/summary.csv')
    f,axes=plt.subplots(1,2,figsize=(7,3.3))
    configs=[r['config'] for r in geometry]
    med=[float(r['diff_ms_median']) for r in geometry]
    lo=[float(r['diff_ms_p25']) for r in geometry];hi=[float(r['diff_ms_p75']) for r in geometry]
    axes[0].errorbar(configs,med,yerr=[np.subtract(med,lo),np.subtract(hi,med)],fmt='o',color=COLORS[M],capsize=4)
    axes[0].set_ylabel('Forced-Merkle diff latency (ms)');axes[0].set_ylim(bottom=0)
    axes[1].bar(configs,[float(r['leaf_entries_examined_median']) for r in geometry],color=COLORS[H],hatch='//',edgecolor='white')
    axes[1].set_ylabel('Leaf entries examined');axes[1].set_ylim(bottom=0)
    for ax in axes:ax.set_xlabel('Fixed trie geometry');ax.grid(axis='y',alpha=.2)
    save(f,7)
    return assets

def content(rows,lookup,pairs):
    old=json.loads((ROOT/'output/docs/seven_revision_audit/manuscript_model.json').read_text(encoding='utf-8'))
    out=copy.deepcopy(old[:58])
    # Keep the reviewed introduction, related work, design, JSON contract and algorithms.
    out[2]['text']=out[2]['text'].rstrip(';')
    out[6]['text']=('Abstract: Revon is a Python prototype for durable, linear-history key/value versioning. '
        'It combines a fixed-depth Merkle hash trie, addressed changesets and an operation-count policy that selects log aggregation or hash-pruned comparison for version differences. '
        'We validate the repaired JSON ownership and identity implementation on a separate Ubuntu cloud machine using 516 isolated workflow executions, including synthetic histories, public taxi records and threshold controls. '
        'Across eight synthetic scenarios and one 100,000-record public scenario, seven paired trials favour the hybrid variant over forced Merkle in eight diff-latency intervals; the dense case includes parity. '
        'A frozen 16,384-operation policy selects Merkle at 32,768 operations even though forced log is faster in every matched trial, so the evidence does not establish an optimal threshold or superiority over always-log. '
        'A separate 72-execution supplement verifies process telemetry and fixed-trie geometry. Million-record runs establish correctness and feasibility for one trial per system, without a comparative performance claim. '
        'Results concern complete Revon API and Dolt CLI workflows, rather than isolated storage engines. The contribution is a reproducible comparison of two correct diff paths over the same durable representation, with explicit costs and limits.')
    out[20]['text']='We compare five systems through six import/diff workflows on deterministic histories, retain per-trial semantic and historical-state checks, balance execution order, and validate the repaired source on a separate Linux cloud machine.'
    out[30]['text']='Table I summarizes the systems and studies most closely related to the implemented design.'
    out[47]['text']='The primary Linux workflow matrix retains T=4,096, matching the original workflow protocol. A separate Linux control tests the previously frozen candidate T=16,384. Bucket capacity b^d and the operation threshold are distinct parameters; neither is tuned on these Linux results.'
    out += [heading('IV. EXPERIMENTAL METHODOLOGY'),heading('A. Environment and evidence boundaries',2),
      body('The current evaluation runs on a separate Google Compute Engine n2-standard-4 VM in us-east1-c. Table II records the environment and separates the two source snapshots. Windows and WSL results from earlier revisions remain archived; they are not pooled with the Linux trials or used as repaired-source measurements. The new machine provides an additional hardware environment, not an independent research team or an experiment isolating the effect of Linux.'),
      table(2,'Linux environment and separately retained evidence',['Property','Recorded configuration'],[
        ['Machine','GCP n2-standard-4; Intel Cascade Lake; 4 vCPUs; 16 GB RAM; us-east1-c'],
        ['Operating system','Ubuntu 24.04.5 LTS; x86_64; kernel 7.0.0-1011-gcp'],
        ['Storage','100 GB balanced persistent disk; no swap configured'],
        ['Runtime','CPython 3.14.3, Clang 22.1.1 build; GIL enabled; JIT not enabled; SQLite 3.50.4'],
        ['Comparator and packages','Dolt 2.3.1; psutil 7.1.0; PyArrow 25.0.1'],
        ['Primary source','linux-validation-20261008: archived repaired source, fixed before measurement'],
        ['Supplement source','linux-final-supplement-20261008: same core; zero-aware I/O sampler and geometry history checks'],
        ['Background conditions','Google Ops Agent active; filesystem cache not cleared; monitoring included']],[.24,.76]),
      body('Each bundle contains the executed source, hashes, environment records, frozen workload definitions, execution order, raw rows and audits. PYTHONHASHSEED is fixed to 20261008. Source snapshots include the exact working revision rather than relying on the repository HEAD alone. Unit tests passed on the VM: 67 before the primary campaign and 69 before the supplement. An early supplement preflight failed because the transferred test_data directory was absent; the empty fixture directory was created and all tests rerun before any supplement timing began. This packaging failure is retained in the supplement records.'),
      heading('B. Variants and workloads',2),
      body('Snapshot stores a complete state per version; Log-only retains the initial state and operations; Revon-M forces Merkle diff; Revon-H selects its path; and Dolt uses a primary-key SQL table and native JSON diff. Dolt SQL INSERT and CSV table-import workflows are measured separately, giving six workflow variants for five systems. Subsequent operations have the same semantic contract, but the interfaces and storage features differ. Table III defines workload identifiers used in the figures.'),
      table(3,'Linux workloads; all primary histories have ten generated commits',['ID / scenario','Rows','Changes / commit','Locality'],[
        ['S1 Small sparse','1,000','10','Spread'],['S2 Medium sparse','10,000','10','Spread'],
        ['S3 Medium dense','10,000','1,000','Spread'],['S4 Large sparse','100,000','100','Spread'],
        ['S5 Application-key local','100,000','100','Lexically adjacent keys'],['S6 Range local','100,000','100','Contiguous key intervals'],
        ['S7 Repeated key','100,000','100','Repeated small key set'],['S8 Hash-route local','100,000','100','Shared 9-bit hash prefix'],
        ['P1 TLC public records','100,000','100','Generated fare corrections']],[.33,.14,.21,.32]),
      body('The primary synthetic matrix has eight scenarios, six variants, two warm-ups and seven measured trials: 432 executions. The public-record matrix adds 27 executions for Revon-M, Revon-H and Dolt SQL. Two threshold-control workloads add 54 executions. Three million-row diagnostics complete the 516-execution campaign: 114 warm-ups, 399 measured trials and three diagnostics. No failed timing execution is silently replaced or excluded; all 516 completed with verified outputs.'),
      heading('C. Public dataset and generated histories',2),
      copy.deepcopy(old[73]),
      body('The Linux campaign evaluates the 100,000-row subset with ten commits and 100 corrections per commit; the million-row diagnostic uses ten commits and 1,000 corrections per commit. Each correction increases fare and total by 0.01 while preserving the other retained fields. All systems receive identical records and mutations. These are generated histories over public records, not observed longitudinal edits. The inherited payload_bytes=8 configuration field is unused by the public-data builder; public values retain the eleven fields and are accounted for by their measured serialized bytes. The harness stores each canonical record as a JSON-encoded string in the common value column. These timings do not characterize native nested-object mutation or field-level query performance.'),
      table(4,'Checksum-pinned public data and Linux protocol',['Item','Definition'],[
        ['Source','NYC TLC yellow_tripdata_2024-01.parquet; 2,964,624 source rows'],
        ['Raw SHA-256','c4d59da7bbc8abaeeeb1727947ee93d9891a71acb42854bd80db1571b2030510'],
        ['Prepared million-row SHA-256','6608b7680e82432ca22f90a71d4994676ffd290c4fc85f17f91fab03d4c7b661'],
        ['P1','100,000 records; H=10; C=100; 2 warm-ups + 7 measured trials per system'],
        ['Scale diagnostic','1,000,000 records; H=10; C=1,000; one trial per system; no interval or ranking']],[.27,.73]),
      heading('D. Timing, correctness and uncertainty',2),
      body('Every primary execution uses a fresh worker and a new repository. Model order is deterministically shuffled and rotated within repetition blocks and retained in the protocol. Import timing reaches the first durable version. Incremental-commit latency is the median of the history\'s commit timers. Diff compares the first and last versions and includes sorted changed keys plus old/new value hashes serialized as canonical JSON. Checkout materializes the full sorted key/value state as canonical UTF-8 JSON. Dolt timings include CLI launch, JSON parsing and output normalization; Revon operates through its in-process API. No engine-only speed claim follows from this comparison.'),
      body('Outside the operation timers, every primary execution checks initial and final state hashes, the observed canonical semantic diff against an independently constructed workload oracle, and all H+1 historical states. Verification JSON retains observed and expected hashes, not only expected key counts. The post-run audit regenerates workloads, checks the exact matrix and selected paths, verifies source/data/protocol identity, and recomputes all summaries. A second local audit checks the downloaded evidence. These are internal integrity checks; they do not constitute independent timing reproduction.'),
      body('We report medians and interquartile ranges (IQRs) from seven measured repository trials, with warm-ups excluded. Matched-trial latency ratios are summarized by their median and a paired percentile bootstrap interval using 20,000 resamples. Ratios above one favour the denominator. These exploratory 95% intervals describe repeated runs of a fixed workload on one VM; they are not population confidence intervals and are not adjusted for multiple comparisons. A median ratio need not equal the ratio of two latency medians. All completed trials, including outlying timings, remain in the evidence.'),
      body('The parent samples worker and descendant RSS, CPU and I/O counters every 10 ms. Whole-worker peaks include workload generation, oracle state and verification, not solely the storage engine. Short-lived subprocess peaks or counters may be missed. The original sampler reports an all-zero read counter as missing, so no primary read-I/O claim is made. The separate supplement distinguishes observed zero from unavailable counters; it is not retroactively merged into the primary data. Physical read bytes can be zero for cached workloads and do not measure logical bytes examined.'),
      body('Repository footprint is the directory size before close and compaction, including indexes, metadata and representation costs. We also retain logical payload normalization and selected first-trial compaction probes. These are workflow measures, not a decomposition into pure tree overhead. CPU, RSS and I/O cover the complete worker lifecycle; operation timers and resource counters therefore have different scopes.'),
      heading('E. Frozen threshold and geometry controls',2),
      body('The earlier Windows threshold study tested 1,024, 2,048, 4,096, 8,192 and 16,384 operations using predefined, disjoint calibration and validation profiles. It selected T=16,384 before validation by minimizing normalized calibration latency. The held-out normalized score was 1.032, compared with 1.297 for T=4,096; always-log scored 1.000 on both splits. These historical cached-query scores remain tied to their archived source and environment, not the repaired Linux implementation. The largest tested candidate winning does not establish an optimum. The earlier long-history and three-seed pilots remain exploratory supplementary evidence.'),
      body('The Linux threshold control freezes T=16,384 and tests 10,000 rows with 512 changes per commit at H=16 and H=64, giving 8,192 and 32,768 accumulated operations. Forced log, forced Merkle and the hybrid run in separate workers with two warm-ups and seven measured trials. No Linux result is used to alter the threshold. This is a two-case transport check of the existing choice, not a new calibration/validation cycle.'),
      body('The geometry supplement uses one spread workload with 10,000 rows, ten commits, 100 changes per commit, 32-byte synthetic values and seed 20260824. It fixes forced Merkle and tests b4-d6, b8-d3, b8-d4, b8-d5 and b16-d3. Each geometry has two warm-ups and seven measured trials in a fresh worker. A rotating order limits configuration/time confounding. Every execution checks all eleven historical states and the endpoint diff outside timing. Geometry timings use the same adapter output contract, but resource peaks from its inherited in-process monitor are not pooled with external RSS measurements.'),
      heading('V. LINUX RESULTS'),heading('A. Hybrid versus forced Merkle',2),
      body('Revon-M and Revon-H retain the same persistent representation and changesets; the ablation changes the requested diff strategy. Figure 3 and Table V show all nine repeated-trial comparisons. The hybrid interval excludes parity in eight scenarios; the dense S3 interval includes one. Sparse and repeated-key results support access to log aggregation under the tested conditions. They do not establish a universal advantage from switching or a representation-level storage saving.'),
      fig(3,diff_figure.CAPTION),
    ]
    ablation=[];workflow=[]
    for sid,s in zip(LABELS,SCENARIOS):
        a=lookup[s,M];b=lookup[s,H];p=pairs[s,M,'diff_ms'];q=pairs[s,D,'diff_ms']
        ablation.append([sid,fmt(a['diff_ms_median']),fmt(b['diff_ms_median']),f"{float(p['median_paired_ratio']):.2f} [{float(p['ci95_low']):.2f}, {float(p['ci95_high']):.2f}]",b['strategy_selected']])
        workflow.append([sid,fmt(lookup[s,D]['diff_ms_median']),fmt(b['diff_ms_median']),f"{float(q['median_paired_ratio']):.2f} [{float(q['ci95_low']):.2f}, {float(q['ci95_high']):.2f}]"])
    out += [table(5,'Paired M/H diff ablation; milliseconds and median trial ratio [95% interval]',['ID','M (ms)','H (ms)','M/H ratio [interval]','H path'],ablation,[.08,.15,.15,.43,.19]),
      body('For P1, Revon-H has a 41.964 ms diff median versus 60.597 ms for forced Merkle. Both use the same eleven-field public records and generated updates. S3 accumulates 10,000 operations and uses Merkle in H; its near-parity result is consistent with both variants executing the same diff path plus selection work.'),
      heading('B. Threshold control and adverse result',2),
      body('Figure 4 and Table VI report both threshold workloads, including the result adverse to the frozen policy. At 8,192 operations, H uses log. At 32,768 operations, H selects Merkle, although forced log is faster in all seven matched measured trials. The operation count alone is insufficient to select the fastest path for these cases. The Linux control therefore completes the validation task while leaving the threshold optimality question unresolved.'),
      fig(4,'Linux control for the previously frozen T=16,384 policy; medians and IQRs from seven measured trials per path, with different history lengths labelled explicitly'),
      table(6,'Frozen threshold controls; diff medians in milliseconds',['Operations / history','Forced log','Forced Merkle','Hybrid','Selected path'],[
        [f'{n:,} / H={h}',fmt(lookup[f'linux-threshold-h{h}-c512','Revon-log calibration']['diff_ms_median']),fmt(lookup[f'linux-threshold-h{h}-c512',M]['diff_ms_median']),fmt(lookup[f'linux-threshold-h{h}-c512',H]['diff_ms_median']),lookup[f'linux-threshold-h{h}-c512',H]['strategy_selected']] for n,h in [(8192,16),(32768,64)]],[.26,.17,.20,.17,.20]),
      heading('C. Complete workflow comparison',2),
      body('Table VII compares canonical diff outputs from the Revon-H API and Dolt SQL/CLI workflows. All nine paired Dolt/H intervals exceed one. This is evidence for these complete workflows; process startup, language, serialization, query execution and resident state are confounded with the data structure. Figure 5 reports commit and checkout separately so the diff result is not generalized to every operation.'),
      table(7,'Dolt SQL/CLI versus Revon-H diff; milliseconds and paired ratio [95% interval]',['ID','Dolt (ms)','H (ms)','Dolt/H ratio [interval]'],workflow,[.09,.18,.18,.55]),
      fig(5,'Linux incremental-commit and full-checkout latency for Revon-H and Dolt SQL/CLI; points show medians and bars show IQRs of seven trials, using logarithmic latency'),
      heading('D. Storage and memory costs',2),
      body('Figure 6 and Table VIII retain repository footprint and complete-worker RSS alongside latency. M and H both store addressed changesets. Any small directory-size difference between them is a measured workflow difference, not evidence that selecting a log path removes trie data. Dolt uses different serialization, compression and indexing. A low diff latency can coexist with a substantially larger retained repository; the resource costs constrain the scope of the efficiency claim.'),
      fig(6,'Median repository directory size before compaction and median whole-worker peak RSS across seven Linux trials; logarithmic scales, including shared workload and verification memory in RSS'),
      table(8,'Repository footprint and worker RSS; medians in MiB',['ID','M disk','H disk','Dolt disk','H RSS','Dolt RSS'],[
        [sid]+[f'{float(lookup[s,m][metric+"_median"])/2**20:.1f}' for m,metric in [(M,'storage_bytes'),(H,'storage_bytes'),(D,'storage_bytes'),(H,'process_tree_peak_rss_bytes'),(D,'process_tree_peak_rss_bytes')]] for sid,s in zip(LABELS,SCENARIOS)],[.08,.18,.18,.20,.18,.18]),
      heading('E. Million-record feasibility',2),
      body('The million-row diagnostic verifies one fresh repository for each of M, H and Dolt, including all eleven historical states and the common diff oracle. Table IX reports these observations without IQRs, confidence intervals or a performance ranking. H uses Merkle because the 10,000 accumulated operations exceed T=4,096. These runs establish that the bounded workflow completed on this VM, not reliable relative scaling or an optimal configuration. Repeated million-row trials remain future work.'),
    ]
    diagnostic=[r for r in rows if r['trial_kind']=='diagnostic']
    out += [table(9,'Single-run million-record diagnostics; n=1 per system, no comparative inference',['System','Import (s)','Commit (ms)','Diff (ms)','Checkout (s)','RSS (GiB)'],[
      [{M:'Revon-M',H:H,D:D}[r['model']],f"{float(r['initial_import_ms'])/1000:.3f}",fmt(r['incremental_commit_ms']),fmt(r['diff_ms']),f"{float(r['checkout_ms'])/1000:.3f}",f"{float(r['process_tree_peak_rss_bytes'])/2**30:.3f}"] for r in sorted(diagnostic,key=lambda r:r['model'])],[.18,.15,.19,.17,.17,.14]),
      heading('F. Fixed-trie geometry and telemetry',2),
      body('Figure 7 and Table X describe the fresh geometry supplement. The equal-capacity b4-d6, b8-d4 and b16-d3 configurations each have 4,096 leaf buckets and equal expected occupancy; their contrast concerns routing shape. Increasing depth at b=8 changes bucket capacity and the traversal/leaf-scan trade-off. One workload and one VM do not establish a universal geometry optimum.'),
      fig(7,'Linux fixed-trie geometry supplement; left shows forced-Merkle diff medians and IQRs from seven trials, right shows median examined leaf entries'),
    ]
    geometry=read(SUPP/'geometry/summary.csv')
    out += [table(10,'Geometry work and footprint; one fixed spread workload',['Geometry','Buckets','Diff (ms)','Leaf entries','Node pairs','Disk (MiB)'],[
        [r['config'],r['leaf_buckets'],fmt(r['diff_ms_median']),f"{float(r['leaf_entries_examined_median']):.0f}",f"{float(r['nodes_compared_median']):.0f}",f"{float(r['storage_bytes_median'])/2**20:.2f}"] for r in geometry],[.18,.14,.17,.19,.16,.16]),
      body('The separate telemetry regression exercises valid all-zero I/O samples and unavailable counters through real isolated workers. Both tests pass. The 27-execution Linux supplement then checks M, H and Dolt on a fixed 10,000-row workload with two warm-ups and seven measured trials per system. These rows retain zero read bytes as observations rather than blanks. They validate counter handling; they do not establish cold-storage behavior, eliminate sampling loss or alter the primary latency summaries.'),
      heading('G. Representation baselines',2),
      body('Table XI provides the complete six-workflow comparison for the representative S4 spread workload. These medians are descriptive; all eight synthetic scenarios, trial dispersion, normalized storage and the separate SQL/CSV import variants are retained in Online Resource 1. Snapshot and Log-only isolate representation choices within the common harness, not production feature equivalence. Their results make the cost of retaining full states or replaying operations visible alongside Revon.'),
      table(11,'Representative S4 workflow medians; complete eight-scenario matrix in Online Resource 1',['Workflow','Import (s)','Commit (ms)','Diff (ms)','Checkout (ms)','Disk (MiB)'],[
        [label,f"{float(lookup['large-sparse',model]['initial_import_ms_median'])/1000:.3f}",fmt(lookup['large-sparse',model]['incremental_commit_ms_median']),fmt(lookup['large-sparse',model]['diff_ms_median']),fmt(lookup['large-sparse',model]['checkout_ms_median']),f"{float(lookup['large-sparse',model]['storage_bytes_median'])/2**20:.1f}"]
        for model,label in [('Snapshot','Snapshot'),('Log-only','Log-only'),(M,'Revon-M'),(H,'Revon-H'),(D,'Dolt SQL'),('Dolt (bulk import)','Dolt CSV')]],[.20,.14,.18,.16,.19,.13]),
      heading('VI. DISCUSSION'),
      body('RQ1: the Linux evidence supports selected diff-workflow advantages and documents storage, memory and commit costs. Snapshot, Log-only and both Dolt import paths remain in the complete raw matrix even where figures emphasize the central M/H/Dolt comparison. The comparison is over a bounded single-writer versioning workflow, with common materialized outputs. It does not establish superior SQL query processing, transaction concurrency, relational capabilities or general structured-data management.'),
      body('RQ2: access to log aggregation improves on forced Merkle in eight of nine paired repeated-trial cases, with one inconclusive dense case. Always-log remains a necessary static control. The frozen threshold\'s adverse Linux result shows that a successful M/H ablation does not validate adaptation itself. A richer selector should be trained on fresh calibration histories and tested on untouched validation histories; these results are not used to implement or tune one.'),
      body('RQ3: the geometry supplement supplies current-source Linux work counters and timings under a rotated execution order. Expected bucket occupancy explains only part of the cost: touched leaf serialization, ancestor reconstruction, output size and traversal shape also matter. Equal-capacity configurations must not be described as changing expected occupancy.'),
      body('The JSON repair is part of all current Linux experiments. Earlier API defects allowed retained mutable aliases and Python-equality collisions, including true versus 1 and 1 versus 1.0. Input/output detachment and v1 encoded-identity comparisons repair these semantics without changing the existing object encoding or on-disk schema. Historical pre-repair timings remain archived and are not relabelled. The Linux trials demonstrate repaired-source behavior on the supplied histories; unit tests additionally cover null/absence, nested ownership, unsupported values and rollback.'),
      heading('VII. LIMITATIONS AND FUTURE WORK'),
      body('External validity: the new cloud VM is separate from the original Windows/WSL host, but all Linux repetitions use one hardware configuration and are author-run. Hardware, operating system, interpreter build, virtualization and storage change together. We cannot attribute differences solely to Linux or infer performance across machines. Google Ops Agent and filesystem caching remain active; neither a perfectly idle machine nor cold I/O is claimed.'),
      body('Workload validity: the synthetic scenarios use one deterministic history per scenario. TLC supplies public record contents but generated corrections, one monthly release and nested affine samples. It supplies no natural edit history, relational queries, multi-table joins or diverse public datasets. Million-row Linux evidence has one execution per system and establishes feasibility only. Additional seeds, observed histories, datasets and repeated scale trials are required for broader inference.'),
      body('Implementation scope: the trie geometry is fixed and the repository eagerly loads retained nodes, changesets and commits into memory. It has no eviction policy. Revon supports linear history and a single writer; named branches, merge commits, conflict handling, schema evolution, remotes, ordered scans, crash fault injection and concurrent workloads are not evaluated. Atomic SQLite transactions do not prove all production durability or operational guarantees.'),
      body('Measurement and inference: API/CLI boundaries and resident-state asymmetry preclude a storage-engine speed claim. Ten-millisecond sampling may miss short peaks and child counters. Whole-worker RSS contains shared oracle data and cannot be read as product-only memory. Keys, log operations and trie-node pairs are different work units, and Dolt internal work is not exposed. Repetition intervals describe fixed histories, with no multiple-comparison correction. Source/hash audits check internal integrity, not independent performance reproduction.'),
      body('Selection: T=16,384 was the largest candidate in a bounded historical search; the cloud controls neither locate a crossover nor prove an optimum. Its slower choice at 32,768 operations must be retained. Future work should hold row count and history length constant, vary locality and density independently, fix process/hash-seed conditions and use independent histories as the uncertainty unit. A new policy needs separate calibration and held-out validation before replacing the frozen implementation.'),
      body('Storage improvements may include binary nodes, value chunks, packed objects, compression and reachability-based garbage collection. An ordered Merkle index or content-defined chunking could support range operations. These changes require separate semantic and performance studies rather than being inferred from the present fixed-trie results.'),
      heading('VIII. CONCLUSION'),
      body('Revon combines incremental fixed-trie updates, addressed changesets and atomic SQLite persistence with two semantically consistent version-diff paths. The repaired source passed 516 Linux workflow executions and a separate 72-execution geometry/telemetry supplement on a cloud machine. Repeated M/H comparisons support log-path benefits for selected workloads, while the dense interval includes parity. The frozen threshold can select a slower path than forced log, so the study does not establish optimal adaptation. Single-run million-record checks demonstrate bounded feasibility and expose resource costs without a comparative scaling claim. The contribution is an auditable prototype and trade-off study over complete versioning workflows, with broader generality and independent reproduction still open.'),
      heading('APPENDIX A. REPRODUCIBILITY'),
      body('Online Resource 1 (Revon_Linux_Reproducibility_Supplement.zip) contains the exact two executed source snapshots, frozen protocols, manifests, raw rows, paired intervals, per-trial semantic verification, audit records and geometry observations. The primary directory is linux-validation-20261008 and the additional directory is linux-final-supplement-20261008. They are separate experiments and their timing rows must not be pooled. README instructions describe data acquisition, pinned environment setup, commands and internal verification. TLC records and generated repositories are rebuilt instead of redistributed. This follows the principle of retaining a complete computational provenance chain [24].'),
      body('For the primary snapshot, run python -m experiments.linux_validation --output <new-directory> --dataset <prepared-TLC-directory> with PYTHONHASHSEED=20261008, pinned Python 3.14.3 and Dolt 2.3.1. For the second snapshot, create tests/test_data, run the 69-test suite, then python -m experiments.linux_final_supplement --output <new-directory>. Environment gates and unit tests run before measurement. Archived audits can be rerun from their corresponding source directories; public metadata and records must match the retained dataset.json. Raw timing reproduction will vary with environment even if semantic hashes agree.'),
      heading('DECLARATIONS'),
      copy.deepcopy(old[189]),copy.deepcopy(old[190]),
      body(old[191]['text'].replace(' All authors reviewed and approved the manuscript.','')),
      body('Ethics and consent to participate: No participants were recruited and no human or animal intervention was performed. This computational study uses synthetic histories and publicly released de-identified TLC records. The study does not claim institutional ethics approval or exemption.'),
      body('Consent for publication: Not applicable; no identifiable participant material is included.'),
      body('Data and code availability: The TLC source file is publicly accessible through NYC TLC [31]; the retained checksum and preparation procedure identify the exact input. The source snapshots and experimental evidence supporting the Linux results are supplied as Online Resource 1 with this manuscript. The public project repository is https://github.com/atharvasheersh/Revon; historical evidence is pinned at commit 37d27d465ffc922a3881b85bb415574f956a9a09. The new supplementary package is supplied with the manuscript and is not claimed to be present at that historical commit.'),
      heading('ACKNOWLEDGMENT'),copy.deepcopy(old[195]),heading('REFERENCES')]
    out.extend(copy.deepcopy(old[197:]))
    for block in out:
        if block['kind']=='reference' and 'https://doi.org/' in block['text']:
            block['text']=re.split(r'\s+(?:Open copy:|Open paper:|Extended report:)',block['text'])[0].rstrip()
    return out

def build_doc(blocks,assets,ieee):
    # Use the established journal/IEEE typography, with journal font >=12 pt.
    doc=layout.setup_document(ieee);body_started=False
    for b in blocks:
        kind=b['kind']
        if ieee and kind=='abstract' and not body_started:
            layout.new_section(doc,2);body_started=True
        if kind=='table':
            if ieee:layout.new_section(doc,1)
            t=layout.add_table(doc,b,ieee)
            # Every final table is short enough to keep its rows together.
            for i,row in enumerate(t.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:p.paragraph_format.keep_with_next=i<len(t.rows)-1
            if ieee:layout.new_section(doc,2)
            continue
        if kind=='figure':
            if ieee:layout.new_section(doc,1)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            width=6.9 if ieee else 6.7
            p.add_run().add_picture(assets[b['number']],width=Inches(width))
            caption=b['caption']
            if not ieee:caption=caption.replace('Table III','Table 3')
            c=doc.add_paragraph(f"Fig. {b['number']}"+('. ' if ieee else ' ')+caption+('.' if ieee else ''),style='Caption')
            c.paragraph_format.keep_together=True
            if ieee:layout.new_section(doc,2)
            continue
        text=b['text']
        if ieee and kind=='keywords':text=text.replace('Keywords:','Index Terms—')
        if not ieee:
            text=re.sub(r'Table ([IVX]+)\b',lambda m:'Table '+str(next(n for n,v in layout.ROMAN.items() if v==m[1])),text)
            if kind=='h1' and re.match(r'^[IVX]+\.',text):
                roman,text2=text.split('. ',1);text=str(next(n for n,v in layout.ROMAN.items() if v==roman))+' '+text2.title()
            elif kind=='h2':text=re.sub(r'^[A-Z]\. ','',text)
        style={'title':'Title','h1':'Heading 1','h2':'Heading 2'}.get(kind,'Normal')
        p=doc.add_paragraph(text,style=style)
        p.paragraph_format.widow_control=True
        if kind in ('title','author'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        elif kind not in ('h1','h2'):p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if kind=='reference':
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
            p.paragraph_format.space_after=Pt(3 if ieee else 7)
            if ieee:p.paragraph_format.line_spacing=1.0
            for r in p.runs:r.font.size=Pt(8 if ieee else 12)
        if kind in ('abstract','keywords'):
            for r in p.runs:r.font.size=Pt(10 if ieee else 12)
        if ieee and kind=='abstract':
            for r in p.runs:r.bold=True
        if ieee and kind=='h1':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if ieee and kind=='h2':
            for r in p.runs:r.italic=True;r.bold=False
    layout.sanitize(doc,'Times New Roman' if ieee else 'Arial')
    path=ROOT/'paper'/('Revon_Research_Paper_IEEE.docx' if ieee else 'Revon_Research_Paper_Discover_Computing.docx')
    doc.save(path);return path

def main():
    OUT.mkdir(parents=True,exist_ok=True);WORK.mkdir(parents=True,exist_ok=True)
    backup=WORK/'before-linux-final';backup.mkdir(exist_ok=True)
    for path in (ROOT/'paper').glob('Revon_Research_Paper_*'):
        if path.suffix in ('.docx','.pdf') and not (backup/path.name).exists():shutil.copy2(path,backup/path.name)
    rows,lookup,pairs=evidence();assets=plots(lookup);blocks=content(rows,lookup,pairs)
    (OUT/'manuscript_model.json').write_text(json.dumps(blocks,indent=2,ensure_ascii=False),encoding='utf-8')
    paths=[build_doc(blocks,assets,False),build_doc(blocks,assets,True)]
    abstract=next(b['text'] for b in blocks if b['kind']=='abstract')
    assert len(abstract.split())<250
    for p in paths:
        doc=Document(p)
        assert len(doc.tables)==11 and len(doc.inline_shapes)==7
        assert 'No independent Linux machine was available' not in '\n'.join(p.text for p in doc.paragraphs)
    audit=dict(abstract_words=len(abstract.split()),blocks=len(blocks),tables=11,figures=7,
        main_executions=516,supplement_executions=72,main_audit_sha256=sha(MAIN/'audit.json'),
        supplement_audit_sha256=sha(SUPP/'audit.json'),geometry_audit_sha256=sha(SUPP/'geometry/audit.json'),
        source_manuscript_model_sha256=sha(ROOT/'output/docs/seven_revision_audit/manuscript_model.json'),
        model_sha256=sha(OUT/'manuscript_model.json'),outputs={p.name:sha(p) for p in paths},
        layout_review='pending exported PDF page inspection')
    (OUT/'build-verification.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
    print(json.dumps(audit,indent=2))

if __name__=='__main__':main()
