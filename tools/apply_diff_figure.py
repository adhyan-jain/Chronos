"""Replace only Figure 3 and its caption in the two current manuscripts."""
from pathlib import Path
import copy
import hashlib
import json
import shutil
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
import preview_diff_figure as figure

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/docs/linux-final-20261008'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    backup=ROOT/'.codex_tmp/figure3-update-20261009'
    backup.mkdir(parents=True,exist_ok=True)
    assets=OUT/'figures'
    figure.main(assets,'Fig3',preview_only=False)
    records=[]
    for name in ['Discover_Computing','IEEE']:
        path=ROOT/f'paper/Revon_Research_Paper_{name}.docx'
        saved=backup/path.name
        if not saved.exists():
            shutil.copy2(path,saved)
            shutil.copy2(path.with_suffix('.pdf'),backup/path.with_suffix('.pdf').name)
        doc=Document(path)
        original_text=[p.text for p in doc.paragraphs]
        original_tables=[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
        original_images=[hashlib.sha256(doc.part.related_parts[s._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob).hexdigest() for s in doc.inline_shapes]
        matches=[(i,p) for i,p in enumerate(doc.paragraphs) if p.text.startswith(('Fig. 3 ','Fig. 3. '))]
        assert len(matches)==1
        i,caption=matches[0]
        drawing=doc.paragraphs[i-1]._p.xpath('.//wp:inline')
        assert len(drawing)==1
        inline=drawing[0]
        rid=inline.xpath('.//a:blip')[0].get(qn('r:embed'))
        part=doc.part.related_parts[rid]
        part._blob=(assets/'Fig3.png').read_bytes()
        width=Inches(6.9 if name=='IEEE' else 6.7)
        height=round(width*4.55/7.35)
        inline.extent.cx=width
        inline.extent.cy=height
        for ext in inline.xpath('.//a:xfrm/a:ext'):
            ext.set('cx',str(width));ext.set('cy',str(height))
        properties=copy.deepcopy(caption.runs[0]._r.rPr) if caption.runs and caption.runs[0]._r.rPr is not None else None
        caption.clear()
        text=figure.CAPTION if name=='IEEE' else figure.CAPTION.replace('Table III','Table 3')
        run=caption.add_run(('Fig. 3. ' if name=='IEEE' else 'Fig. 3 ')+text+'.')
        if properties is not None:
            run._r.insert(0,properties)
        caption.paragraph_format.keep_together=True
        doc.save(path)
        check=Document(path)
        changed=[j for j,(a,b) in enumerate(zip(original_text,[p.text for p in check.paragraphs])) if a!=b]
        assert len(check.paragraphs)==len(original_text) and changed==[i],changed
        assert original_tables==[[[c.text for c in row.cells] for row in t.rows] for t in check.tables]
        new_images=[hashlib.sha256(check.part.related_parts[s._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob).hexdigest() for s in check.inline_shapes]
        assert len(new_images)==len(original_images)==7
        assert [j for j,(a,b) in enumerate(zip(original_images,new_images)) if a!=b]==[2]
        records.append(dict(format=name,docx_sha256=sha(path),changed_paragraphs=changed,
                            unchanged_tables=len(check.tables),changed_image=3))
    model_path=OUT/'manuscript_model.json'
    model=json.loads(model_path.read_text(encoding='utf-8'))
    block=next(b for b in model if b['kind']=='figure' and b['number']==3)
    block['caption']=figure.CAPTION
    model_path.write_text(json.dumps(model,indent=2,ensure_ascii=False),encoding='utf-8')
    audit_path=OUT/'build-verification.json'
    audit=json.loads(audit_path.read_text(encoding='utf-8'))
    audit['model_sha256']=sha(model_path)
    audit['outputs']={f'Revon_Research_Paper_{name}.docx':sha(ROOT/f'paper/Revon_Research_Paper_{name}.docx') for name in ['Discover_Computing','IEEE']}
    audit['layout_review']='pending exported PDF page inspection after Figure 3 replacement'
    audit_path.write_text(json.dumps(audit,indent=2)+'\n',encoding='utf-8')
    report=dict(outputs=records,scope='Figure 3 image and caption only',data_unchanged=True,
                backups=str(backup),figure_sha256=sha(assets/'Fig3.png'))
    (OUT/'figure3-update-verification.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,indent=2))

if __name__=='__main__':
    main()
