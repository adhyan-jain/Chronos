"""Build synchronized journal/IEEE manuscripts from reviewed sources and evidence.

Use the Codex bundled Python; plotting dependencies can be supplied on PYTHONPATH.
The original source is backed up once before any replacement.
"""
from __future__ import annotations
import copy
import argparse
import csv
import hashlib
import json
import re
import shutil
import statistics
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from matplotlib.ticker import NullFormatter, FuncFormatter
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/".codex_tmp/seven_revisions"
AUDIT=ROOT/"output/docs/seven_revision_audit"
TITLE="Adaptive Merkle Trie Versioning for Efficient Structured Data Management"
PRIMARY=ROOT/"evidence/paper-corrected-issues13-17-counterbalanced-final-20260923"
PUBLIC=ROOT/"evidence/public-tlc-combined-20261002"
THRESH=ROOT/"evidence/threshold-heldout-20261002"
EXTENSION=ROOT/"evidence/crossover-diagnostic-20261002"
PILOT=ROOT/"evidence/threshold-robustness-pilot-20261002"
TEAL,BLUE,ORANGE,INK="#006D77","#305B93","#A64B2A","#202D36"

def csv_rows(path):
    with path.open(encoding="utf-8",newline="") as f:return list(csv.DictReader(f))

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def boot(a,b,seed=20261002):
    rng=np.random.default_rng(seed)
    av=np.median(rng.choice(a,(20000,len(a))),axis=1)
    bv=np.median(rng.choice(b,(20000,len(b))),axis=1)
    lo,hi=np.quantile(av/bv,[.025,.975])
    return float(np.median(a)/np.median(b)),float(lo),float(hi)

def analyze():
    for directory,name in ((PRIMARY,"evidence_audit.json"),(PUBLIC,"evidence_audit.json"),(THRESH,"audit.json"),(EXTENSION,"audit.json")):
        assert json.loads((directory/name).read_text())["passed"],directory
    old=csv_rows(PRIMARY/"summary.csv"); public=csv_rows(PUBLIC/"summary.csv")
    lookup={(r["scenario"],r["model"]):r for r in old if r["phase"]=="evaluation"}
    ratios={r["scenario"]:r for r in csv_rows(PRIMARY/"paired_ratio_uncertainty.csv") if r["numerator_model"]=="Revon-M (forced Merkle)" and r["metric"]=="diff_ms"}
    labels=[("small-sparse","Small sparse"),("medium-sparse","Medium sparse"),("medium-dense","Medium dense"),("large-sparse","Large sparse"),("large-application-key-local","Application-key local"),("large-range-local","Range local"),("large-repeated-key","Repeated key"),("large-hash-route-local","Hash-route local")]
    primary=[]
    for scenario,label in labels:
        m,h=lookup[(scenario,"Revon-M (forced Merkle)")],lookup[(scenario,"Revon-H")]
        ci=ratios[scenario]
        primary.append(dict(label=label,scenario=scenario,n=7,M=m,H=h,ratio=float(ci["median_paired_ratio"]),lo=float(ci["ci95_low"]),hi=float(ci["ci95_high"]),interval="paired fixed-workload",latency_reduction_percent=100*(1-float(h["diff_ms_median"])/float(m["diff_ms_median"])),storage_change_percent=100*(float(h["storage_bytes_median"])/float(m["storage_bytes_median"])-1)))
    raw=csv_rows(PUBLIC/"raw_results.csv")
    pub=[]
    for scenario,label in (("tlc-n10000-h10-c10","10k / H10 / 0.1%"),("tlc-n100000-h10-c100","100k / H10 / 0.1%"),("tlc-n1000000-h10-c1000","1M / H10 / 0.1%"),("tlc-n100000-h50-c100","100k / H50 / 0.1%"),("tlc-n100000-h10-c10","100k / H10 / 0.01%"),("tlc-n100000-h10-c1000","100k / H10 / 1%")):
        m=next(r for r in public if r["scenario"]==scenario and r["model"]=="Revon-M (forced Merkle)")
        h=next(r for r in public if r["scenario"]==scenario and r["model"]=="Revon-H")
        a=[float(r["diff_ms"]) for r in raw if r["scenario"]==scenario and r["model"]==m["model"] and r["status"]=="ok" and r["trial_kind"]=="measured"]
        b=[float(r["diff_ms"]) for r in raw if r["scenario"]==scenario and r["model"]==h["model"] and r["status"]=="ok" and r["trial_kind"]=="measured"]
        assert len(a)==len(b)==7
        ratio,lo,hi=boot(a,b)
        pub.append(dict(label=label,scenario=scenario,n=7,M=m,H=h,ratio=ratio,lo=lo,hi=hi,interval="exploratory unpaired fixed-workload",latency_reduction_percent=100*(1-float(h["diff_ms_median"])/float(m["diff_ms_median"])),storage_change_percent=100*(float(h["storage_bytes_median"])/float(m["storage_bytes_median"])-1)))
    threshold=csv_rows(THRESH/"summary.csv"); diagnostic=csv_rows(EXTENSION/"summary.csv")
    scores={split:{t:statistics.mean(float(r["median_ms"])/min(float(x["median_ms"]) for x in threshold if x["scenario"]==r["scenario"] and x["strategy"] in ("log","merkle")) for r in threshold if r["split"]==split and r["strategy"]=="hybrid" and int(r["threshold"])==t) for t in (1024,2048,4096,8192,16384)} for split in ("calibration","validation")}
    frozen=json.loads((THRESH/"frozen_selection.json").read_text())
    assert min(scores["calibration"],key=lambda t:(scores["calibration"][t],t))==frozen["threshold"]
    for t,value in scores["calibration"].items():assert abs(value-frozen["calibration_scores"][str(t)])<1e-10
    best=[]
    for scenario in sorted({r["scenario"] for r in threshold}):
        cases=[r for r in threshold if r["scenario"]==scenario and r["strategy"]=="hybrid"]
        winner=min(cases,key=lambda r:(float(r["median_ms"]),int(r["threshold"])))
        best.append({k:winner[k] for k in ("scenario","split","rows","commits","locality","operations","threshold","median_ms")})
    best.sort(key=lambda r:(r['split'],int(r['rows']),int(r['commits']),r['locality'],int(r['operations'])))
    measured=[r for r in diagnostic if r["strategy"]=="log"]
    crossover=[int(r["operations"]) for r in measured if float(r["median_ms"])>float(next(x for x in diagnostic if x["scenario"]==r["scenario"] and x["strategy"]=="merkle")["median_ms"])]
    result=dict(primary=primary,public=pub,scores=scores,frozen_threshold=frozen["threshold"],threshold_best=best,diagnostic_first_merkle_faster=min(crossover) if crossover else None,
                diagnostic_ratios=[dict(operations=int(r["operations"]),log_over_merkle=float(r["median_ms"])/float(next(x for x in diagnostic if x["scenario"]==r["scenario"] and x["strategy"]=="merkle")["median_ms"])) for r in measured])
    assert json.loads((PILOT/'independent_audit.json').read_text())['passed']
    result['pilot_ratios']=json.loads((PILOT/'paired_ratios.json').read_text())
    AUDIT.mkdir(parents=True,exist_ok=True);(AUDIT/"analysis.json").write_text(json.dumps(result,indent=2))
    return result,lookup,threshold,diagnostic

def theme():
    plt.rcParams.update({"font.family":"Arial","font.size":9,"axes.labelsize":9,"axes.titlesize":9,"xtick.labelsize":8,"ytick.labelsize":8,"pdf.fonttype":42,"svg.fonttype":"none","axes.spines.top":False,"axes.spines.right":False,"axes.axisbelow":True})

def save(fig,name):
    for ax in fig.axes:
        if ax.get_yscale()=="log":
            ax.yaxis.set_minor_formatter(NullFormatter())
            ax.yaxis.set_major_formatter(FuncFormatter(lambda x,pos:f"{x:g}"))
    path=WORK/"figures"/f"{name}.png";path.parent.mkdir(parents=True,exist_ok=True)
    fig.savefig(path,dpi=600,facecolor="white")
    fig.savefig(path.with_suffix(".svg"),facecolor="white")
    fig.savefig(path.with_suffix(".pdf"),facecolor="white")
    plt.close(fig);return str(path)

def box(ax,x,y,w,h,text,color=BLUE,fs=9):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.006,rounding_size=0.01",fc="white",ec=color,lw=1))
    ax.text(x+w/2,y+h/2,text,ha="center",va="center",fontsize=fs,color=INK)

def arrow(ax,a,b,label=None):
    ax.annotate("",xy=b,xytext=a,arrowprops=dict(arrowstyle="->",color=INK,lw=1))
    if label:ax.text((a[0]+b[0])/2,(a[1]+b[1])/2+.025,label,ha="center",fontsize=8)

def figures(data,lookup,threshold,diagnostic):
    theme();assets={}
    fig,ax=plt.subplots(figsize=(6.7,2.9));ax.set(xlim=(0,1),ylim=(0,1));ax.axis("off")
    for x,t in ((.01,"Frontend\nmetrics"),(.265,"REST API"),(.52,"Revon core"),(.775,"SQLite\nobjects / commits\nHEAD")):box(ax,x,.65,.215,.26,t,fs=8)
    for x in (.225,.48,.735):arrow(ax,(x,.78),(x+.035,.78))
    ax.text(.5,.43,"Core operations and representation",ha="center",fontsize=10,weight="bold")
    for x,t in ((.01,"Fixed-depth trie\nb = 8, d = 4"),(.35,"Incremental commit\nshared immutable paths"),(.69,"Hybrid diff\nlog or Merkle")):box(ax,x,.08,.29,.26,t,TEAL)
    arrow(ax,(.625,.64),(.625,.49));assets[1]=save(fig,"Fig1")
    fig,ax=plt.subplots(figsize=(3.4,4.7));ax.set(xlim=(0,1),ylim=(0,1));ax.axis("off")
    box(ax,.10,.87,.80,.10,"Requested version pair",fs=8)
    box(ax,.10,.68,.80,.13,"Validate versions and ancestry\nCount accumulated operations",fs=8)
    arrow(ax,(.5,.86),(.5,.82))
    box(ax,.12,.46,.76,.13,"Ancestor path exists\nand operation count ≤ T?",TEAL,8)
    arrow(ax,(.5,.67),(.5,.60))
    box(ax,.02,.23,.42,.14,"Log aggregation\nReverse if needed",TEAL,8)
    box(ax,.56,.23,.42,.14,"Merkle comparison\nPrune equal hashes",ORANGE,8)
    arrow(ax,(.30,.45),(.22,.38));arrow(ax,(.70,.45),(.78,.38))
    ax.text(.17,.41,"Yes",ha="center",fontsize=8);ax.text(.84,.41,"No",ha="center",fontsize=8)
    box(ax,.10,.035,.80,.11,"Sorted semantic diff output",fs=8)
    arrow(ax,(.22,.22),(.36,.15));arrow(ax,(.78,.22),(.64,.15))
    ax.text(.5,.005,"Both paths use the same persistent trie",ha="center",fontsize=8)
    fig.subplots_adjust(left=.04,right=.96,top=.99,bottom=.04);assets[2]=save(fig,"Fig2")
    fig,axes=plt.subplots(2,3,figsize=(6.7,4.8))
    groups=sorted({(r["split"],r["rows"],r["commits"],r["locality"]) for r in threshold})
    for index,(ax,group) in enumerate(zip(axes.flat,groups)):
        for strategy,color,marker in (("log",TEAL,"o"),("merkle",ORANGE,"s")):
            rows=sorted([r for r in threshold if (r["split"],r["rows"],r["commits"],r["locality"])==group and r["strategy"]==strategy],key=lambda r:int(r["operations"]))
            x=[int(r["operations"]) for r in rows];y=[float(r["median_ms"]) for r in rows]
            ax.plot(x,y,color=color,marker=marker,ms=3,label=strategy.capitalize())
            ax.fill_between(x,[float(r["p25_ms"]) for r in rows],[float(r["p75_ms"]) for r in rows],color=color,alpha=.15)
        ax.set(xscale="log",yscale="log");ax.set_title(f"({chr(97+index)}) {int(group[1])//1000}k, H{group[2]}\n{group[3]}, {group[0]}",fontsize=8)
        ax.set_xticks([512,4096,32768],["512","4,096","32,768"]);ax.grid(lw=.4,color="#D9E0E4")
        if index%3==0:ax.set_ylabel("Diff (ms)")
        if index>=3:ax.set_xlabel("Operations")
    fig.legend(*axes[0,0].get_legend_handles_labels(),loc="lower center",ncol=2,frameon=False)
    fig.subplots_adjust(left=.10,right=.98,top=.87,bottom=.17,hspace=.85,wspace=.42);assets[3]=save(fig,"Fig3")
    fig,axes=plt.subplots(1,3,figsize=(6.7,2.85))
    groups=sorted({(r["rows"],r["commits"],r["locality"]) for r in threshold if r["split"]=="validation"})
    for index,(ax,group) in enumerate(zip(axes,groups)):
        for operations,color,marker in ((8192,BLUE,"o"),(16384,TEAL,"s"),(32768,ORANGE,"^")):
            rows=sorted([r for r in threshold if (r["rows"],r["commits"],r["locality"])==group and r["split"]=="validation" and r["strategy"]=="hybrid" and int(r["operations"])==operations],key=lambda r:int(r["threshold"]))
            y=[float(r["median_ms"]) for r in rows]
            ax.errorbar([int(r["threshold"]) for r in rows],y,yerr=[[y[i]-float(r["p25_ms"]) for i,r in enumerate(rows)],[float(r["p75_ms"])-y[i] for i,r in enumerate(rows)]],color=color,marker=marker,ms=3,lw=.9,capsize=2,label=f"{operations:,} operations")
        ax.set(xscale="log",yscale="log");ax.set_title(f"({chr(97+index)}) {int(group[0])//1000}k / H{group[1]}\n{group[2]}",fontsize=8)
        ax.set_xticks([1024,4096,16384],["1,024","4,096","16,384"]);ax.set_xlabel("Threshold T");ax.grid(lw=.4,color="#D9E0E4")
        if index==0:ax.set_ylabel("Diff (ms)")
    fig.legend(*axes[0].get_legend_handles_labels(),loc="lower center",ncol=3,fontsize=8,frameon=False)
    fig.subplots_adjust(left=.10,right=.98,top=.80,bottom=.30,wspace=.42);assets[4]=save(fig,"Fig4")
    fig,axes=plt.subplots(1,2,figsize=(6.7,3.45))
    for ax,rows,panel in zip(axes,(data["primary"],data["public"]),("(a) Primary synthetic","(b) Public TLC records")):
        for offset,model,color,marker in ((-.12,"M",ORANGE,"s"),(.12,"H",TEAL,"o")):
            y=[i+offset for i in range(len(rows))];med=[float(r[model]["diff_ms_median"]) for r in rows]
            err=[[med[i]-float(r[model]["diff_ms_p25"]) for i,r in enumerate(rows)],[float(r[model]["diff_ms_p75"])-med[i] for i,r in enumerate(rows)]]
            ax.errorbar(med,y,xerr=err,fmt=marker,color=color,ms=3,capsize=2,label="Revon-"+model)
        ax.set(xscale="log",ylim=(len(rows)-.5,-.5));ax.set_yticks(range(len(rows)),[r["label"] for r in rows]);ax.set_xlabel("Diff latency (ms)");ax.set_title(panel,fontsize=9);ax.grid(axis="x",lw=.4,color="#D9E0E4")
    fig.legend(*axes[0].get_legend_handles_labels(),loc="lower center",ncol=2,frameon=False)
    fig.subplots_adjust(left=.22,right=.98,top=.87,bottom=.22,wspace=1.08);assets[5]=save(fig,"Fig5")
    fig,ax=plt.subplots(figsize=(6.7,3.65))
    for offset,model,color,marker in ((-.13,"Revon-H",TEAL,"o"),(.13,"Dolt",ORANGE,"s")):
        rows=[lookup[(r["scenario"],model)] for r in data["primary"]];med=[float(r["diff_ms_median"]) for r in rows]
        ax.errorbar(med,[i+offset for i in range(8)],xerr=[[med[i]-float(r["diff_ms_p25"]) for i,r in enumerate(rows)],[float(r["diff_ms_p75"])-med[i] for i,r in enumerate(rows)]],fmt=marker,color=color,capsize=2,label=model)
    ax.set(xscale="log",ylim=(7.5,-.5),xlabel="Workflow diff latency (ms)");ax.set_yticks(range(8),[r["label"] for r in data["primary"]]);ax.grid(axis="x",lw=.4,color="#D9E0E4")
    fig.legend(*ax.get_legend_handles_labels(),loc="lower center",ncol=2,frameon=False)
    fig.subplots_adjust(left=.25,right=.97,top=.96,bottom=.24);assets[6]=save(fig,"Fig6")
    fig,axes=plt.subplots(1,2,figsize=(6.7,3.5))
    for ax,rows,panel in zip(axes,(data["primary"],data["public"]),("(a) Synthetic","(b) TLC records")):
        for model,color,marker in (("M",ORANGE,"s"),("H",TEAL,"o")):
            ax.scatter([float(r[model]["storage_bytes_median"])/2**20 for r in rows],[float(r[model]["diff_ms_median"]) for r in rows],color=color,marker=marker,s=22,label="Revon-"+model,facecolors="none" if model=="M" else color)
        for i,r in enumerate(rows,1):
            x=float(r["H"]["storage_bytes_median"])/2**20;y=float(r["H"]["diff_ms_median"])
            offset={4:(25,15),5:(-15,-4),6:(25,-15)}.get(i,(3,3)) if panel.startswith("(a)") else (3,3)
            ax.annotate(str(i),(x,y),xytext=offset,textcoords="offset points",fontsize=8,
                arrowprops=dict(arrowstyle="-",lw=.5,color=INK) if offset!=(3,3) else None)
        ax.set(xscale="log",yscale="log",xlabel="Directory footprint (MiB)",ylabel="Diff (ms)");ax.set_title(panel);ax.grid(lw=.4,color="#D9E0E4")
    fig.legend(*axes[0].get_legend_handles_labels(),loc="lower center",ncol=2,frameon=False)
    fig.subplots_adjust(left=.10,right=.98,top=.87,bottom=.24,wspace=.45);assets[7]=save(fig,"Fig7")
    fig,axes=plt.subplots(1,2,figsize=(6.7,3.7))
    for ax,metric,title in zip(axes,("incremental_commit_ms","checkout_ms"),("(a) Incremental commit","(b) Historical checkout")):
        for offset,model,color,marker in ((-.12,"Revon-H",TEAL,"o"),(.12,"Dolt",ORANGE,"s")):
            rows=[lookup[(r["scenario"],model)] for r in data["primary"]];med=[float(r[metric+"_median"]) for r in rows]
            ax.errorbar(med,[i+offset for i in range(8)],xerr=[[med[i]-float(r[metric+"_p25"]) for i,r in enumerate(rows)],[float(r[metric+"_p75"])-med[i] for i,r in enumerate(rows)]],fmt=marker,color=color,ms=3,capsize=2,label=model)
        ax.set(xscale="log",ylim=(7.5,-.5),xlabel="Workflow latency (ms)");ax.set_title(title,fontsize=9);ax.set_yticks(range(8),[r["label"] for r in data["primary"]] if ax==axes[0] else []);ax.grid(axis="x",lw=.4,color="#D9E0E4")
    fig.legend(*axes[0].get_legend_handles_labels(),loc="lower center",ncol=2,frameon=False)
    fig.subplots_adjust(left=.25,right=.98,top=.87,bottom=.22,wspace=.25);assets[8]=save(fig,"Fig8")
    rows=csv_rows(ROOT/"evidence/trie-sensitivity-issues8-20260923/summary.csv")
    fig,ax=plt.subplots(figsize=(6.7,3.35))
    for i,r in enumerate(rows):
        x=float(r["storage_bytes_median"])/2**20;y=float(r["diff_ms_median"]);c=TEAL if r["config"]=="b8-d4" else BLUE
        ax.errorbar(x,y,yerr=[[y-float(r["diff_ms_p25"])],[float(r["diff_ms_p75"])-y]],fmt="o",color=c,capsize=3)
        ax.annotate(r["config"],(x,y),xytext=(8,5 if i%2 else -12),textcoords="offset points",fontsize=9)
    ax.set(xlabel="Repository footprint (MiB)",ylabel="Diff (ms)",xlim=(3.8,9),ylim=(12,max(float(r["diff_ms_p75"]) for r in rows)+4));ax.grid(lw=.4,color="#D9E0E4")
    fig.subplots_adjust(left=.12,right=.97,top=.96,bottom=.19);assets[9]=save(fig,"Fig9")
    rows=csv_rows(PUBLIC/"summary.csv");fig,axes=plt.subplots(1,3,figsize=(6.7,3.0))
    models=[("Snapshot","#697986","o"),("Log-only","#926C16","s"),("Revon-M (forced Merkle)",ORANGE,"^"),("Revon-H",TEAL,"D"),("Dolt",BLUE,"v"),("Dolt (bulk import)","#745785","P")]
    for ax,metric,title,unit in zip(axes,("incremental_commit_ms","diff_ms","storage_bytes"),("(a) Commit (ms)","(b) Diff (ms)","(c) Storage (MiB)"),(1,1,2**20)):
        for model,color,marker in models:
            values=[next(r for r in rows if r["scenario"]==f"tlc-n{n}-h10-c{n//1000}" and r["model"]==model) for n in (10000,100000,1000000)]
            y=[float(r[metric+"_median"])/unit for r in values]
            ax.errorbar([10000,100000,1000000],y,yerr=[[y[i]-float(r[metric+"_p25"])/unit for i,r in enumerate(values)],[float(r[metric+"_p75"])/unit-y[i] for i,r in enumerate(values)]],color=color,marker=marker,ms=3,lw=.9,capsize=2,label=model.replace(" (forced Merkle)","").replace("Dolt (bulk import)","Dolt CSV").replace("Dolt","Dolt SQL") if model=="Dolt" else model.replace(" (forced Merkle)","").replace("Dolt (bulk import)","Dolt CSV"))
        ax.set(xscale="log",yscale="log",xlabel="Initial records");ax.set_title(title);ax.set_xticks([10000,100000,1000000],["10k","100k","1M"]);ax.grid(lw=.4,color="#D9E0E4")
    fig.legend(*axes[0].get_legend_handles_labels(),loc="lower center",ncol=3,fontsize=8,frameon=False)
    fig.subplots_adjust(left=.07,right=.98,top=.89,bottom=.34,wspace=.38);assets[10]=save(fig,"Fig10")
    fig,ax=plt.subplots(figsize=(6.7,2.65))
    for strategy,color,marker in (("log",TEAL,"o"),("merkle",ORANGE,"s")):
        rs=sorted([r for r in diagnostic if r["strategy"]==strategy],key=lambda r:int(r["operations"]))
        ax.plot([int(r["operations"]) for r in rs],[float(r["median_ms"]) for r in rs],color=color,marker=marker,label=strategy.capitalize())
        ax.fill_between([int(r["operations"]) for r in rs],[float(r["p25_ms"]) for r in rs],[float(r["p75_ms"]) for r in rs],color=color,alpha=.15)
    ax.set(xscale="log",yscale="log",xlabel="Accumulated operations (256 generated commits)",ylabel="Diff (ms)");ax.set_xticks([8192,32768,131072],["8,192","32,768","131,072"]);ax.grid(lw=.4,color="#D9E0E4");ax.legend(frameon=False)
    fig.subplots_adjust(left=.12,right=.98,top=.96,bottom=.25);assets[11]=save(fig,"Fig11")
    return assets

def ivalue(row,key="diff_ms",digits=2):
    return f"{float(row[key+'_median']):.{digits}f} [{float(row[key+'_p25']):.{digits}f}, {float(row[key+'_p75']):.{digits}f}]"

def blocks(source,data):
    ps=[p.text for p in source.paragraphs];out=[]
    def p(text,kind="body"):out.append(dict(kind=kind,text=text))
    def span(a,b):
        for i in range(a,b+1):
            if ps[i].strip():p(ps[i],"h1" if re.match(r"^[IVX]+\. ",ps[i]) else "h2" if re.match(r"^[A-Z]\. ",ps[i]) else "body")
    def fig(n,caption):out.append(dict(kind="figure",number=n,caption=caption))
    def tab(n,caption,headers,rows,widths=None):out.append(dict(kind="table",number=n,caption=caption,headers=headers,rows=rows,widths=widths))
    def original_table(index,n,caption):
        rows=[[c.text for c in r.cells] for r in source.tables[index].rows];tab(n,caption,rows[0],rows[1:])
    p(TITLE,"title")
    for i in range(1,6):p(ps[i],"author")
    p("Abstract: Revon is a Python prototype for durable, linear-history key/value versioning using a fixed-depth Merkle hash trie. Its hybrid variant chooses between changeset aggregation and hash-pruned comparison; trie geometry does not adapt. We make forced-Merkle Revon-M the primary ablation because both variants retain the same persistent representation and changesets. On eight synthetic workloads, five paired diff-latency intervals favour Revon-H and three include parity. A separate three-seed study also leaves the 50-commit result inconclusive. Public NYC TLC records extend evaluation to one million records with generated corrections; Revon-H has no consistent advantage when both variants choose Merkle comparison. A supplemental study tests five thresholds on disjoint calibration and validation workloads, with seven measured cached queries per configuration. Calibration selects 16,384 operations, while the original workflow results retain their frozen 4,096 threshold. No forced-path crossover appears through 32,768 operations in the main supplement; a separate long-history diagnostic investigates larger counts. Revon-H and Revon-M have nearly equal repository footprints, whereas both use more storage than Dolt in the measured workflows. All measurements use one physical computer, including WSL2. The evidence supports workload-dependent differencing choices, not general efficiency or production-system superiority.","abstract")
    out[-1]['text']=out[-1]['text'].replace("a separate long-history diagnostic investigates larger counts.","separate exploratory diagnostics investigate larger counts. A three-seed pilot finds different preferred paths at the same operation count across two workload profiles.")
    p(ps[8].replace("Index Terms:","Keywords:"),"keywords")
    span(9,10)
    p("Revon asks whether a fixed-depth Merkle hash trie can combine incremental structural sharing with an operation index. Revon-H adapts only its diff path: it selects log aggregation for eligible intervals within an operation threshold and Merkle comparison otherwise. The trie geometry remains fixed. SQLite supplies atomic, durable storage. The title's efficiency claim is restricted to the measured operations and workloads; it does not assert that all structured-data management becomes cheaper.")
    span(12,35)
    # Place the existing related-work table beside its first caption.
    out=[b for b in out if not b["text"].startswith("TABLE I.")]
    location=next(i for i,b in enumerate(out) if b.get("text","")==ps[33])
    table_rows=[[c.text for c in r.cells] for r in source.tables[0].rows]
    out.insert(location,dict(kind="table",number=1,caption="Related work and the scope of Revon",headers=table_rows[0],rows=table_rows[1:],widths=[.22,.25,.53]))
    p(ps[36].replace("The measured crossover depends on the workload and configuration.","A crossover must be measured rather than inferred from the switching rule."))
    p(ps[37],"h1");fig(1,"Revon architecture and the implemented research scope")
    span(41,46);p(ps[50],"h2")
    p("The log path walks ancestry and aggregates key operations; the Merkle path compares roots recursively, skips equal hashes, and examines entries under unequal paths. Revon-H validates both version identifiers, checks ancestry in either direction, and counts changeset operations. It chooses log aggregation when an ancestor path exists and the count is at most T; otherwise it uses Merkle comparison. Reverse log requests swap old/new values, and a same-version request returns an empty diff. Forced log mode rejects unrelated versions. Unknown versions raise an error. Revon-M forces Merkle comparison, although its public version-pair API also checks ancestry. Both paths return the same sorted semantic diff. Commit and checkout are separate operations.")
    p("The original synthetic and public workflow matrices retain T=4,096. The supplemental calibration evaluates other thresholds without changing those archived results. The 4,096 possible leaf buckets arise from b=8,d=4 and are independent of the operation threshold.")
    fig(2,"Actual hybrid selection for a version pair; T denotes the operation threshold. Invalid versions raise an error before selection. Both paths share the persistent trie")
    p(ps[52],"h1");p("A. Environment and evidence bundles","h2")
    tab(2,"Environment and the distinct evidence bundles",["Bundle","Environment and configuration","Measurement scope"],[
        ["Primary synthetic","Windows 11; Python 3.14.3; Dolt 2.3.1; b8-d4, T4096","540 executions; 2 warm-ups and 7 measured trials per configuration"],
        ["Robustness","Same Windows host; three seeds; 10k rows; H10/H50","Seed-cluster intervals, reported separately"],
        ["WSL sensitivity","Same computer; Linux 6.6.87.2-microsoft-standard-WSL2; Python 3.14.4","Operating-environment comparison, not independent hardware"],
        ["Public TLC","Same Windows host; Python 3.12.14; pyarrow 25.0.1; psutil 7.2.2; Dolt 2.3.1","325 attempts; 324 successful; 1 interrupted attempt retained; recovery outside block"],
        ["Threshold supplement","Same Windows host; Python 3.14.3; 10k/100k rows; b8-d4","42 repositories; 2 warm-up and 7 measured queries per path/threshold"],
        ["Threshold robustness pilot","Same Windows host; Python 3.14.3; psutil 7.1.0; three new seeds","12 repositories; 714 queries, including 546 measured; exploratory, separate from held-out selection"],
        ["Physical host","Intel Core i7-1255U; 10 cores/12 logical CPUs; 15.64 GiB RAM; C: NTFS","CPU model/filesystem verified for this revision; historical manifests contain processor family"],
        ["Independent Linux host","Not available to this revision","Independent replication remains uncompleted"]],widths=[.20,.40,.40])
    p("These bundles share a physical host but differ in runtime, workloads and protocol. Results are not pooled across them. The public-data run started with 4.14 GiB available RAM; available memory was not controlled. The current CPU model and filesystem were inspected during revision, not retrospectively recorded for each historical trial. Recorded versions in each manifest remain authoritative.")
    p("B. Variants and synthetic workloads","h2");p(ps[54]);original_table(1,3,"Synthetic workloads in the primary evaluation")
    p("C. Timing correctness and uncertainty","h2");p(ps[58]);p(ps[60])
    p("The counterbalanced primary design pairs same-trial ratios within a scenario. Seven trials estimate repeated-run uncertainty on a fixed workload and host. The public-data ablation instead uses exploratory unpaired percentile bootstrap intervals for ratios of medians, with 20,000 resamples and seed 20261002. This avoids treating the isolated recovery as part of the original block. IQRs describe dispersion and are not confidence intervals. Repeated cached-query intervals in the threshold supplement do not estimate uncertainty across independent repository builds.")
    p("The prior calibration sampled through 4,096 operations and did not locate a crossover. Its 540 successful primary executions and all 127 Tukey outlier candidates are retained. New threshold results are supplemental measurements, not replacements for those trials.")
    p("D. Public dataset and scalability evaluation","h2");span(66,68)
    p("The public run retains 72 warm-ups and 252 successful measured executions across 36 configurations. One additional measured attempt was interrupted. Recovery contributes one of the 252 successful measurements and ran outside the original counterbalanced block. The evidence includes 1,221 successful historical state-hash checks; pilots do not enter reported performance summaries.")
    original_table(2,4,"Public TLC dataset [31], January 2024; all histories are generated corrections")
    p("E. Threshold calibration and held-out validation","h2")
    p("Before measuring, we specify five candidates: 1,024, 2,048, 4,096, 8,192 and 16,384 operations. Three calibration profiles use 10k/H4/spread, 100k/H16/spread and 10k/H64/repeated-key workloads. Three validation profiles use distinct seeds and 100k/H4/repeated-key, 10k/H16/spread and 100k/H64/spread workloads. Each profile is tested at 512, 1,024, 2,048, 4,096, 8,192, 16,384 and 32,768 accumulated operations, with 32-byte synthetic values. Changes per commit equal accumulated operations divided by history length, so sizes, histories and densities vary; this is not a full factorial population study.")
    p("One durable repository is constructed per workload, then forced log, forced Merkle and all five selectors are measured in deterministically shuffled query blocks with two warm-ups and seven measured queries. Timers include diff_versions and canonical semantic output materialization. Correctness and separate decision probes are outside the diff timer. The model caches stored objects in memory; these are repeated in-process queries, not independent repository trials or cold I/O measurements. Endpoint and midpoint states, reverse and identity diffs, integrity hashes, and reopen checks pass for all 42 repositories and 2,646 queries, including 2,058 measured queries.")
    p("Threshold selection minimizes the equal-weight mean of each calibration-case selector median divided by the faster forced-path median; an exact score tie selects the smaller threshold. The choice is recorded before validation starts. Validation profiles do not influence that choice. Both requested and actual selected paths are retained. A standalone ancestry/count probe has configuration medians of 0.0038-0.1186 ms in the main supplement. It measures decision work separately and is not subtracted from total latency as a causal estimate of selector overhead.")
    fig(3,"Forced-path latency across calibration and held-out profiles; medians and IQRs from seven measured cached queries per configuration")
    fig(4,"Held-out selector latency versus switching threshold for 8,192, 16,384 and 32,768 accumulated operations; medians and IQRs")
    tab(5,"Candidate scores; equal-weight mean latency relative to the faster forced-path median, lower is better",["Threshold","Calibration score","Held-out score","Selection"],[[f"{t:,}",f"{data['scores']['calibration'][t]:.3f}",f"{data['scores']['validation'][t]:.3f}","Frozen choice" if t==data["frozen_threshold"] else "Candidate"] for t in (1024,2048,4096,8192,16384)],widths=[.20,.26,.26,.28])
    p(f"Calibration selects T={data['frozen_threshold']:,}, the largest tested candidate. Its held-out normalized score is {data['scores']['validation'][data['frozen_threshold']]:.3f}, versus {data['scores']['validation'][4096]:.3f} for T=4,096. These scores are descriptive aggregates across the predefined cases, not population confidence intervals. Forced log remains faster at every largest sampled count through 32,768 in the main supplement. Consequently, the candidate winner is bounded by the search range and does not identify an optimal crossover. Per-case lowest-median candidates are retained in the supporting audit; held-out minima are diagnostic and never used to reselect T.")
    p("A separate exploratory extension investigates a 10k-row repeated-key workload with 256 commits and 8,192 through 131,072 operations. Its seed and five operation counts are fixed before extension measurements. The 315 queries, including 245 measured queries, pass the same checks. This extension is excluded from calibration and validation scores.")
    fig(11,"Exploratory long-history diagnostic, 10k rows and 256 repeated-key commits; medians and IQRs. This extension does not alter the frozen threshold")
    first=data["diagnostic_first_merkle_faster"]
    p("The diagnostic median log/Merkle ratios are "+", ".join(f"{r['log_over_merkle']:.2f} at {r['operations']:,}" for r in data["diagnostic_ratios"])+" operations. Merkle first has the lower sampled median at 32,768, log is lower again at 65,536, and Merkle is lower at 131,072. This non-monotonic ordering is not a resolved single crossover or an optimal switching rule; IQRs and raw queries are retained, and only one seed was used." if first else "The diagnostic still does not demonstrate a forced-path crossover within its tested range.")
    p("F. Exploratory threshold robustness pilot","h2")
    p("A separate pilot freezes thresholds 4,096, 16,384, 32,768, 65,536 and 131,072, with forced log (always-log) and forced Merkle baselines. Three new seeds, 20261021-20261023, generate independent histories for each condition. Profiles are 10k rows/256 commits/repeated-key and 100k rows/64 commits/spread, sampled at 32,768, 65,536 and 131,072 operations. A predeclared wall-time rule reduces the initial 18 cases to all nine repeated-key cases and three spread cases at 65,536, before latency trends are inspected. Six spread cases at the other counts are not run; no failed or unfavourable result is discarded. The retained 12 repositories take 28.3 minutes and reach 1,214 MiB sampled worker RSS.")
    p("Each repository uses two warm-up and five measured randomized query blocks for seven configurations. Both profiles at 65,536 also use three blocks with a separate reopened process for each configuration. The 588 warmed queries include 420 measured queries; all 126 fresh-process queries are measured. All 714 timed queries match the common canonical semantic oracle, and reverse/identity, representative checkout, persistent-integrity and reopen checks pass. Timers include semantic output materialization; separate decision probes are not subtracted. Each commit has one insert and one delete, with the remaining mutations being updates. Actual distinct keys, output sizes, work counters and mutation counts are retained. Query repeats are nested within three independent histories per condition; no population confidence interval is claimed.")
    p("Reopening eagerly loads and verifies the persistent object model before diff timing; open latency is separate and the OS filesystem cache is not cleared. Warmed workers use the workload seed as Python hash seed, while fresh workers use workload seed plus block. Both paths share a hash seed within each fresh block, but warmed/fresh comparisons do not isolate caching alone. This pilot neither recalibrates T nor reuses held-out results to select a new threshold. Its five warmed and three fresh query blocks differ from the earlier protocols; the bundles remain separate.")
    p("V. RESULTS","h1");p("A. Primary hybrid ablation","h2")
    p("Revon-M versus Revon-H is the central ablation: both use the same persistent trie and addressed changesets, and the requested diff strategy changes. Table VI reports all eight primary workloads; Table VII reports the separate public-record sweep. Figure 5 displays medians and IQRs. These measurements isolate diff-path selection more directly than cross-system workflows, although run order, runtime effects and the common output contract still affect observed latency.")
    fig(5,"Revon-M and Revon-H ablation; primary synthetic workloads and public TLC workloads are separate panels. Points show medians and bars show IQRs")
    headers=["Workload","M ms [IQR]","H ms [IQR]","M/H [95% CI]","n"]
    tab(6,"Primary ablation; paired ratio intervals on fixed synthetic workloads",headers,[[r["label"],ivalue(r["M"]),ivalue(r["H"]),f"{r['ratio']:.2f} [{r['lo']:.2f}, {r['hi']:.2f}]","7"] for r in data["primary"]],widths=[.22,.23,.23,.25,.07])
    p("A ratio above one favours Revon-H; intervals containing one are inconclusive. Three primary intervals contain parity, including the dense case in which both variants use Merkle comparison. The separate three-seed 50-commit result is 0.99 [0.90, 1.17] across 21 paired trials and also remains inconclusive. Four other robustness factors favour Revon-H on both Windows and WSL2, as documented in their separate seed-cluster analyses.")
    tab(7,"Public-record ablation; exploratory unpaired ratio-of-medians intervals on fixed workloads",headers,[[r["label"],ivalue(r["M"]),ivalue(r["H"]),f"{r['ratio']:.2f} [{r['lo']:.2f}, {r['hi']:.2f}]","7"] for r in data["public"]],widths=[.22,.23,.23,.25,.07])
    long=next(r for r in data["public"] if "h50" in r["scenario"])
    p(f"The public 50-commit case has an M/H ratio of {long['ratio']:.2f} [{long['lo']:.2f}, {long['hi']:.2f}] and median latencies {float(long['M']['diff_ms_median']):.2f} ms for M and {float(long['H']['diff_ms_median']):.2f} ms for H. Its point estimate favours M. This is a separate update-only workload and must not be presented as resolving the earlier mixed-workload 50-commit result. The million-record recovery was outside the original block, which further limits causal interpretation of that configuration.")
    p("B. Complete workflow comparison with Dolt","h2");fig(6,"Revon-H API and Dolt SQL/CLI diff workflows; medians and IQRs across eight synthetic workloads")
    original_table(3,8,"Revon-H and SQL-path Dolt diff latency; median paired ratios and 95% bootstrap intervals")
    p(ps[80].replace("Table IV","Table VIII"))
    p("C. Latency and storage attribution","h2")
    p("Diff-path selection does not imply a second persistent representation. Both M and H retain addressed changesets and trie objects. Their measured footprints are nearly equal. The large Revon-versus-Dolt footprint difference therefore cannot be assigned to hybrid selection. Table IX reports the actual M/H storage difference and point-estimate latency reduction; a negative reduction means H is slower. Percentage reductions are computed from scenario medians and are distinct from the paired median-ratio statistic in Table VI.")
    tab(9,"Latency and storage differences in primary synthetic workloads; directory footprints before compaction",["Workload","M MiB","H MiB","Storage change %","H latency reduction %"],[[r["label"],f"{float(r['M']['storage_bytes_median'])/2**20:.2f}",f"{float(r['H']['storage_bytes_median'])/2**20:.2f}",f"{r['storage_change_percent']:+.3f}",f"{r['latency_reduction_percent']:+.1f}"] for r in data["primary"]],widths=[.28,.15,.15,.20,.22])
    fig(7,"Latency versus directory footprint for M and H; panels use the row order of Tables VI and VII. Numerals identify H points, and overlapping footprints do not imply missing M measurements")
    p(ps[88])
    p("The supplemental repositories measure canonical object payload bytes by kind before deletion of the reproducible working database. Node and changeset payload totals are exact SQL sums, not complete allocated storage costs. For the held-out 100k-row, 64-commit spread workload at 32,768 operations, the closed file is 188.59 MiB: canonical node payloads total 60.34 MiB, changesets 15.21 MiB and commits 0.02 MiB. The remaining 113.02 MiB includes hashes, indexes, metadata, page structure and free space; it is not pure metadata overhead. This is a supplemental repository measurement, separate from the primary workflow directory footprints. These measurements identify retained object categories but do not measure the footprint of a changeset-free variant. Such a variant is not needed to compare existing M and H, and no equivalent variant is claimed here.")
    p("D. Commit checkout and trie geometry","h2");fig(8,"Incremental-commit and historical-checkout workflow latency; medians and IQRs for Revon-H and Dolt")
    p(ps[87]);fig(9,"Fixed-trie geometry sensitivity; points show median diff latency and bars show IQRs for seven measured trials")
    p(ps[93])
    p("E. Public dataset scale history and update density","h2")
    fig(10,"Public-record scalability for ten generated commits and 0.1% updates per commit; medians and IQRs from seven successful measured trials. Both axes are logarithmic. H uses log diff at 10k/100k and Merkle diff at 1M")
    span(97,100);original_table(4,10,"Public history and density results at 100k records; median commit and diff times in ms and footprint in MiB")
    p(ps[102])
    p("F. Exploratory pilot results","h2")
    pilot=data['pilot_ratios']
    pilot_rows=[]
    for locality,operations in [('repeated-key',32768),('repeated-key',65536),('repeated-key',131072),('spread',65536)]:
        rows=sorted((r for r in pilot if r['cache_condition']=='warmed-built' and r['locality']==locality and r['operations']==operations),key=lambda r:r['seed'])
        assert [r['seed'] for r in rows]==[20261021,20261022,20261023]
        profile='10k / H256 / repeated-key' if locality=='repeated-key' else '100k / H64 / spread'
        pilot_rows.append([profile,f'{operations:,}']+[f"{r['log_over_merkle']:.3f}" for r in rows])
    tab(11,"Exploratory pilot log/Merkle ratios of warmed-query medians; below one favours log. Seed suffixes 21-23 denote 20261021-20261023",['Rows / history / locality','Operations','Seed 21','Seed 22','Seed 23'],pilot_rows,widths=[.38,.17,.15,.15,.15])
    p("All three seeds favour log at 32,768 repeated-key operations and Merkle at 65,536 and 131,072. The earlier single-seed diagnostic's opposite preferences at 32,768 and 65,536 are not reproduced; those original measurements remain reported separately. At 65,536 operations, all three spread histories favour log instead. Reopened measurements retain the profile contrast for all three seeds: log/Merkle median ratios are 1.198-1.410 for repeated-key and 0.709-0.734 for spread. These profiles differ in row count and history length as well as locality, so the contrast does not isolate a causal locality effect or establish an exact crossover.")
    p("Relative to T=16,384, T=131,072 lowers warmed spread median latency by 13.1-22.2%, but raises repeated-key latency by 12.3-19.7% at 65,536 and 36.0-53.5% at 131,072 operations. These ranges describe three seed-specific differences, not confidence intervals. Same-path selector/forced median ratios across retained conditions range from 0.614 to 1.213, exposing substantial timing variability. Small differences among thresholds selecting the same algorithm cannot therefore be treated as algorithmic gains. Neither always-log nor simply increasing T is consistently best across these profiles. No new threshold is selected and no production default is changed.")
    p("VI. DISCUSSION","h1")
    p("For RQ1, the public-record study shows a material scale cost: at one million records, H's 4,764.26 ms commit median exceeds the other measured workflows, and its 1,446.69 MiB footprint is 10.90 times Dolt SQL's. Lower diff latency in selected workflows does not establish general efficiency. Interface-level Dolt ratios do not establish an intrinsically faster trie.")
    p("For RQ2, adaptive differencing helps selected sparse and repeated-key workloads without changing their version semantics. M/H footprints are nearly equal because the representations are shared. Dense and long histories do not consistently favour H. The new threshold supplement shows that 4,096 is a conservative frozen rule for earlier results, not a measured universal crossover; 16,384 is the best sampled candidate in a distinct cached-query protocol.")
    p("The separate pilot strengthens the case for investigating workload-sensitive selection: one operation count corresponds to different preferred paths across the retained profiles. It also shows why a larger threshold alone is insufficient. The single-seed reversal does not repeat under the new histories, but three histories and confounded profiles cannot resolve its cause or validate a replacement selector. The frozen 16,384 candidate choice and original 4,096 workflow settings remain unchanged.")
    p(ps[106].replace("the results rule out one trie geometry as best for every workload", "the sampled geometries trade latency against storage; no universal optimum is established"));p(ps[107]);p(ps[108].replace("Figure 7 and Table V","Figure 10 and Table X"))
    p("VII. LIMITATIONS AND THREATS TO VALIDITY","h1")
    p("Runtime and interfaces: Revon runs in-process in Python, while Dolt is a Go executable reached through CLI/SQL. The measured workflows differ in schema and feature costs. Revon supports one writer and linear history; SQL, branching, merging, remotes, schema evolution, ordered range scans and concurrent workloads are not evaluated.")
    p(ps[112])
    p("Environment and replication: Windows and WSL2 measurements use the same physical computer. No independent Linux machine was available to this revision, and no external researcher or clean clone has reproduced the timings. The internal audit checks evidence consistency; it does not reproduce wall-clock performance. Broader hardware claims remain unsupported.")
    p("Calibration: the five-candidate, predefined held-out supplement studies cached queries on synthetic workloads with one seed per profile and one repository per operation count. The largest tested candidate wins, so the search does not establish an optimum. The long-history extension is exploratory and excluded from selection. Query repetitions do not supply independent dataset or repository replications. The fixed-depth trie does not adapt its geometry.")
    p("Pilot scope: only three independent histories per retained condition are measured. Six planned spread cases are omitted by the predeclared resource rule. Row count, history length and locality vary together; background resource conditions are not controlled. Fresh-process hash seeds differ from warmed seeds, and eager reopen plus an uncleared OS cache is not cold storage. Large same-path timing variation prevents causal claims about selector overhead or small candidate differences. The pilot is exploratory and supplies no new calibration/held-out selection or independent Linux replication.")
    p(ps[115]);p(ps[116]);p(ps[117])
    p("Storage and scope: both M and H retain changesets. Canonical payload accounting excludes allocated index and page costs, and comparing roots does not itself explain the repository footprint. Here, adaptive refers solely to diff-path selection; the geometry is fixed. The results cover selected versioning operations on a prototype and do not establish general efficiency across structured-data management.")
    p("VIII. FUTURE WORK","h1")
    p("Priority work is independent hardware replication and additional public datasets with observed histories. Threshold calibration should test a wider candidate range, multiple seeds per profile, different output sizes, and cold-storage behaviour before deployment. A future selector could use measured prefix density and I/O as well as operation count; none of that richer classification exists in the present implementation.")
    p("A focused follow-up should hold row count and history length constant while varying locality and update density, fix Python hash seeds across process conditions, and record GC and resource behaviour. Independent histories should be the uncertainty unit. More seeds, fresh calibration and untouched validation workloads are needed before adopting a richer selector; the present pilot does not justify changing T.")
    p(ps[120]);p("Named branches, merge commits, conflict reporting, schema-aware records, concurrency and ordered queries require implementation and separate correctness/performance studies. They are future capabilities rather than demonstrated properties.")
    p("IX. CONCLUSION","h1")
    p("Revon combines immutable objects, incremental trie updates, SQLite persistence, historical checkout and adaptive log/Merkle differencing. A prominent same-representation ablation supports advantages for selected sparse and repeated-key intervals, with dense and long-history qualifications retained. Public records reach one million rows using generated corrections and expose substantial commit and storage costs. New calibration and held-out measurements improve threshold evidence without rewriting earlier results or establishing a universal optimum. The contribution is a reproducible comparison of diff choices in a single-writer prototype; independent hardware validation remains uncompleted.")
    out[-1]['text']=out[-1]['text'].replace("The contribution is", "A separate three-seed pilot finds profile-dependent path preferences and opposing effects from increasing the threshold, without validating a new optimum. The contribution is")
    p("DECLARATIONS","h1")
    p("Funding: the funding statement must be confirmed by the authors before submission.")
    p("Competing interests and author contributions: author-confirmed declarations have not yet been supplied for this revision. Names, affiliations and contact details are preserved from the reviewed manuscript.")
    p("Ethics: this revision uses public TLC trip records and generated modifications. Any required ethics or exemption statement must be confirmed by the authors for the journal's policy; no approval is invented.")
    p("AI assistance: OpenAI Codex assisted code development, analysis scripting, manuscript restructuring and editing, and plotting code during this revision. Graphs are generated from retained measured outputs or the inspected implementation, not invented empirical data. The authors must review the revised arguments, interpretations, references and artifacts and confirm the final disclosure and accountability before submission.")
    p("Data and code availability: the existing project repository is https://github.com/atharvasheersh/Revon. TLC acquisition and checksum-pinned preparation are documented in docs/experiments/PUBLIC_DATASET_EVALUATION.md [31]. The new public-data and threshold evidence is retained locally with this revision; public release of these new bundles is pending. Appendix A identifies the exact bundles and reproduction commands. Source trip files are downloaded separately and excluded from Git.")
    p("ACKNOWLEDGMENT","h1");p(ps[125]);p("REFERENCES","h1")
    for i in range(127,158):
        text=ps[i].replace("A. G. Parameswaran","A. Parameswaran")
        number=int(re.match(r"\[(\d+)\]",text).group(1))
        if number==17:text=text.replace("https://www.researchgate.net/publication/241633210_Efficient_Versioning_for_Scientific_Array_Databases","https://dspace.mit.edu/entities/publication/a5c2c050-6abd-47c1-8bee-4f8c51b6a15f")
        if number==2:text=text.replace("Dolt Documentation, 2024.","Dolt Documentation, accessed Oct. 2, 2026.")
        if number==22:text=text.replace("https://www.research.ed.ac.uk/files/","https://www.pure.ed.ac.uk/ws/files/")
        if number==26:text=text.replace("2024, doi:","2024, pp. 479-493, doi:")
        # Expand every DOI to a stable, human-readable DOI link.
        text=re.sub(r"doi:\s*([^\s]+)",lambda m:"https://doi.org/"+m.group(1).rstrip(".")+".",text)
        if number==6:text+=" [Online]. Available: https://arxiv.org/abs/1407.3561."
        if number in (4,5,7,8,9,10,11,12,13,15):
            from audit_revision_sources import URLS
            text+=" Open copy: "+URLS[number]+"."
        p(text,"reference")
    p("APPENDIX A. REPRODUCIBILITY AND DATA AVAILABILITY","h1")
    p("The primary synthetic evidence is evidence/paper-corrected-issues13-17-counterbalanced-final-20260923, with paired_ratio_uncertainty.csv. Separate robustness, WSL and telemetry bundles are evidence/paper-issues18-20-robustness-20260923, evidence/paper-issues18-20-linux-wsl-20260923 and evidence/paper-issues14-windows-telemetry-20260924. Trie sensitivity is evidence/trie-sensitivity-issues8-20260923. Public TLC analysis is evidence/public-tlc-combined-20261002, referring to its original execution and isolated recovery parents. TLC preparation pilots are excluded. Local revision evidence is not yet a public release.")
    p("New held-out threshold evidence is evidence/threshold-heldout-20261002. Reproduce with Python 3.14: python -m experiments.threshold_validation --output <new-directory>. The diagnostic extension is evidence/crossover-diagnostic-20261002, reproduced with python -m experiments.crossover_extension --output <new-directory>. Each bundle archives exact source files and hashes, workloads, raw query records, selection, summaries and correctness checks. The pilot at evidence/threshold-pilot-20261002 is excluded. Working databases are reproducible and removed after verified reopen checks; payload accounting and integrity counts are retained.")
    p("The separate robustness pilot is evidence/threshold-robustness-pilot-20261002. It retains the frozen manifest and source hashes, raw timings, seed-specific ratios, resource exclusions and semantic checks. It is not the earlier excluded threshold-pilot bundle. Reproduction commands and process/cache qualifications are in docs/experiments/THRESHOLD_ROBUSTNESS_PILOT.md. Public release of this new bundle is pending; it is reported only as exploratory evidence, separate from original and held-out results.")
    # Put reproducibility before declarations and references so an internal
    # report paragraph does not become a sparse final manuscript page.
    appendix_index=next(i for i,b in enumerate(out) if b.get('text','').startswith('APPENDIX A.'))
    appendix=out[appendix_index:];out=out[:appendix_index]
    declaration_index=next(i for i,b in enumerate(out) if b.get('text','').lower()=='declarations')
    out[declaration_index:declaration_index]=appendix
    # Correct cross-references after moving tables and figures.
    for b in out:
        if b["kind"] in ("body","abstract"):
            b["text"]=b["text"].replace("Table III specifies","Table IV specifies")
            b["text"]=b["text"].replace("Table IV gives","Table VIII gives")
        if "text" in b:b["text"]=b["text"].replace("\u2014",", ").replace("\u2013","-").replace("LBFS introduced content-defined chunking", "LBFS uses content-defined chunking")
    remap={11:5,5:6,6:7,7:8,8:9,9:10,10:11}
    for b in out:
        if b["kind"]=="figure":b["number"]=remap.get(b["number"],b["number"])
        if "text" in b:b["text"]=re.sub(r"Figure (\d+)",lambda m:"Figure "+str(remap.get(int(m[1]),int(m[1]))),b["text"])
    table_mentions={1:"Table I compares the related systems and the scope addressed by Revon.",2:"Table II separates the environments and measurement scopes of the retained evidence bundles.",3:"Table III defines the primary synthetic workload matrix.",4:"Table IV summarizes the public dataset mapping and generated history sweeps.",5:"Table V compares the predefined threshold candidates using calibration and held-out scores.",8:"Table VIII reports the paired workflow comparison with Dolt.",10:"Table X reports public-record history and density results at 100,000 records."}
    figure_mentions={1:"Figure 1 summarizes the architecture and implemented research scope.",2:"Figure 2 shows the implemented version-pair decision, including eligibility and reverse handling.",3:"Figure 3 compares forced-path latency across the predefined calibration and validation profiles.",4:"Figure 4 shows how changing the threshold affects selector latency in held-out workloads.",5:"Figure 5 examines the separate long-history diagnostic without changing the calibrated choice.",6:"Figure 6 compares both variants across synthetic and public-record workloads.",7:"Figure 7 compares the complete Revon-H and Dolt diff workflows.",8:"Figure 8 places measured diff latency against repository footprint for each variant.",9:"Figure 9 reports commit and checkout workflow latency separately from diff latency.",10:"Figure 10 examines the measured latency and footprint of alternative fixed geometries.",11:"Figure 11 reports the separate public-record scalability sweep."}
    table_mentions[11]="Table XI reports all three seed-specific forced-path ratios for each retained warmed pilot condition."
    expanded=[]
    for b in out:
        mention=table_mentions.get(b.get("number")) if b["kind"]=="table" else figure_mentions.get(b.get("number")) if b["kind"]=="figure" else None
        if mention:expanded.append(dict(kind="body",text=mention))
        expanded.append(b)
    return expanded

def columns(section,n):
    cols=section._sectPr.find(qn("w:cols"))
    if cols is None:cols=OxmlElement("w:cols");section._sectPr.append(cols)
    cols.set(qn("w:num"),str(n));cols.set(qn("w:space"),"288")

def new_section(doc,n):
    section=doc.add_section(WD_SECTION_START.CONTINUOUS);columns(section,n)
    p=doc.paragraphs[-1];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
    for r in p.runs:r.font.size=Pt(1)
    return section

def setup_document(ieee=False):
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5 if ieee else 8.2677);s.page_height=Inches(11 if ieee else 11.6929)
    s.top_margin=Inches(.75);s.bottom_margin=Inches(.75);s.left_margin=Inches(.75);s.right_margin=Inches(.75)
    s.header_distance=Inches(.3);s.footer_distance=Inches(.3)
    font="Times New Roman" if ieee else "Arial";size=10 if ieee else 12
    for style in ("Normal","Title","Heading 1","Heading 2","Caption"):
        st=doc.styles[style];st.font.name=font;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.line_spacing=1.04 if ieee else 1.12
        st.paragraph_format.space_after=Pt(4 if ieee else 7)
    doc.styles["Title"].font.size=Pt(20 if ieee else 18)
    doc.styles["Caption"].font.bold=False
    doc.styles["Caption"].font.size=Pt(9 if ieee else 12)
    doc.styles["Title"].paragraph_format.line_spacing=1.08
    for style in ("Heading 1","Heading 2"):
        doc.styles[style].font.bold=True;doc.styles[style].paragraph_format.keep_with_next=True
        doc.styles[style].paragraph_format.space_before=Pt(9 if ieee else 12)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.style=doc.styles["Normal"]
    run=footer.add_run();field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");run._r.addnext(field)
    doc.core_properties.title=TITLE;doc.core_properties.author="Atharva Sheersh Pandey; Adhyan Jain; Poornima Nedunchezhian"
    return doc

def sanitize(doc,font):
    # Remove the default Word theme's title rule and theme-font overrides.
    for root in (doc._element,doc.styles.element):
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for fonts in root.xpath(".//w:rFonts"):
            for key in list(fonts.attrib):
                if "Theme" in key or "theme" in key:del fonts.attrib[key]
            fonts.set(qn("w:ascii"),font);fonts.set(qn("w:hAnsi"),font);fonts.set(qn("w:eastAsia"),font);fonts.set(qn("w:cs"),font)

ROMAN={1:"I",2:"II",3:"III",4:"IV",5:"V",6:"VI",7:"VII",8:"VIII",9:"IX",10:"X",11:"XI"}

def add_table(doc,block,ieee):
    size=8.5 if ieee else 12;full=7 if ieee else 6.76
    cap=doc.add_paragraph(("TABLE "+ROMAN[block["number"]]+". " if ieee else "Table "+str(block["number"])+" ")+block["caption"],style="Caption")
    cap.paragraph_format.keep_with_next=True;cap.alignment=WD_ALIGN_PARAGRAPH.CENTER if ieee else WD_ALIGN_PARAGRAPH.LEFT
    cols=len(block["headers"]);t=doc.add_table(rows=1,cols=cols);t.autofit=False
    fractions=block.get("widths") or ({2:[.26,.74],3:[.22,.26,.52],5:[.24,.15,.14,.20,.27]}.get(cols,[1/cols]*cols))
    for col,f in zip(t.columns,fractions):col.width=Inches(full*f)
    for i,row in enumerate([block["headers"]]+block["rows"]):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        tr=cells[0]._tc.getparent();prop=tr.get_or_add_trPr();prop.append(OxmlElement("w:cantSplit"))
        if i==0:prop.append(OxmlElement("w:tblHeader"))
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(full*fractions[j]);cell.text=str(text)
            tc=cell._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ("top","bottom","left","right"):
                e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
            tc.append(borders)
            if i==0:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"E8EDF1");tc.append(shade)
            margins=OxmlElement("w:tcMar")
            for edge in ("top","bottom","left","right"):
                e=OxmlElement("w:"+edge);e.set(qn("w:w"),"60");e.set(qn("w:type"),"dxa");margins.append(e)
            tc.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.04
                # A short table stays intact; long journal tables may continue
                # with repeated headers, but never strand the last data row.
                short_table=(ieee and len(block['rows'])<=18) or (not ieee and block['number'] in (1,2,3,4,5,10,11) and len(block['rows'])<=24)
                p.paragraph_format.keep_with_next=i<len(block['rows']) if short_table else (i==0 or i==len(block['rows'])-1)
                p.paragraph_format.widow_control=True
                for r in p.runs:r.font.name="Times New Roman" if ieee else "Arial";r.font.size=Pt(size);r.bold=i==0
    # Keep caption/header with the first data row, but permit deliberate long-table continuation.
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    return t

def build_doc(blocks,assets,path,ieee):
    doc=setup_document(ieee);body_started=False
    for b in blocks:
        kind=b["kind"]
        if kind=="h1" and not body_started and ieee:
            new_section(doc,2);body_started=True
        if kind=="table":
            if ieee:new_section(doc,1)
            add_table(doc,b,ieee)
            if ieee:new_section(doc,2)
            continue
        if kind=="figure":
            number=b["number"];full=number!=2
            if ieee:new_section(doc,1)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
            p.add_run().add_picture(assets[number],width=Inches(7 if ieee and full else 3.4 if ieee else 6.4 if full else 3.8))
            caption=f"Fig. {number}"+(". " if ieee else " ")+b["caption"]+". Plotting/diagram code assisted by OpenAI Codex; content derives from measured data or inspected code."
            if not ieee:caption=caption.replace('Tables VI and VII','Tables 6 and 7').rstrip(".")
            p=doc.add_paragraph(caption,style="Caption");p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True;p.paragraph_format.widow_control=True
            if ieee:new_section(doc,2)
            continue
        text=b["text"]
        if not ieee:
            for n,r in sorted(ROMAN.items(),reverse=True):text=text.replace("Table "+r,"Table "+str(n))
            if kind=="h1":text=re.sub(r"^([IVX]+)\. ",lambda m:str(next(k for k,v in ROMAN.items() if v==m[1]))+" ",text).title()
            if kind=="h2":text=re.sub(r"^[A-Z]\. ","",text)
        style={"title":"Title","h1":"Heading 1","h2":"Heading 2"}.get(kind,"Normal")
        p=doc.add_paragraph(text,style=style);p.paragraph_format.widow_control=True
        if kind in ("title","author"):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            if kind=="author":
                if "Atharva" in text:
                    p.clear()
                    for part in re.split(r"([12](?=,|$))",text):
                        r=p.add_run(part);r.font.superscript=part in ("1","2")
                for r in p.runs:r.font.size=Pt(10 if ieee else 12)
                p.paragraph_format.space_after=Pt(2)
        elif kind in ("body","abstract"):p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind=="h1" and ieee:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif kind=="reference":
            p.paragraph_format.keep_together=True
            p.paragraph_format.space_after=Pt(2 if ieee else 4)
            p.paragraph_format.left_indent=Inches(.27);p.paragraph_format.first_line_indent=Inches(-.27)
            number=int(re.match(r"\[(\d+)\]",text)[1]);bookmark=OxmlElement("w:bookmarkStart");bookmark.set(qn("w:id"),str(number+1000));bookmark.set(qn("w:name"),f"ref_{number}")
            end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(number+1000));p._p.insert(0,bookmark);p._p.append(end)
            if ieee:
                for r in p.runs:r.font.size=Pt(9)
        if kind=="title":
            p.paragraph_format.space_after=Pt(10);p.paragraph_format.line_spacing=1.08
    sanitize(doc,"Times New Roman" if ieee else "Arial");doc.save(path)
    return doc

def report(data):
    doc=Document();s=doc.sections[0];s.top_margin=s.bottom_margin=Inches(.55);s.left_margin=s.right_margin=Inches(.65)
    for st in ("Normal","Title","Heading 1"):
        doc.styles[st].font.name="Arial";doc.styles[st].font.size=Pt(10.5);doc.styles[st].font.color.rgb=RGBColor(0,0,0);doc.styles[st].paragraph_format.space_after=Pt(5);doc.styles[st].paragraph_format.line_spacing=1.0
    doc.styles["Title"].font.size=Pt(15);doc.styles["Title"].font.bold=True
    doc.add_paragraph("Discover Computing Internal Manuscript Review",style="Title")
    doc.add_paragraph(TITLE).runs[0].bold=True
    doc.add_paragraph("Internal simulated review | 2 October 2026 | Not an official journal decision")
    p=doc.add_paragraph("Recommendation: further revision required before journal submission. Suitable for supervisor review after the authors inspect the revised analysis.");p.runs[0].bold=True
    statuses=[("1","Completed","Dedicated M/H tables and figure retain inconclusive and adverse findings."),("2","Completed",f"Five candidates; disjoint held-out workloads; frozen T={data['frozen_threshold']:,}; diagnostic extension separate."),("3","Completed","Measured M/H storage differences replace unsupported hybrid-storage attribution."),("4","Blocked","Independent Linux hardware unavailable; WSL is the same physical host."),("5","Completed","Environment, protocols, bundle boundaries and recovery limitations consolidated."),("6","Completed","Diagram follows real ancestry/count selection, reverse handling and shared trie."),("7","Completed","Scope and limitations tightened; broad title risk explicitly retained.")]
    t=doc.add_table(rows=1,cols=3);t.autofit=False
    for cell,label in zip(t.rows[0].cells,("Point","Status","Finding")):cell.text=label
    for row in statuses:
        cells=t.add_row().cells
        for c,text in zip(cells,row):c.text=text
    for i,c in enumerate(t.columns):c.width=Inches([.55,1.50,4.55][i])
    for index,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j,c in enumerate(row.cells):
            c.width=Inches([.55,1.50,4.55][j])
            borders=OxmlElement("w:tcBorders")
            for edge in ("top","bottom","left","right"):
                e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
            c._tc.get_or_add_tcPr().append(borders)
            if index==0:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"E8EDF1");c._tc.get_or_add_tcPr().append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.size=Pt(10);r.bold=index==0
    doc.add_paragraph("Strongest contribution",style="Heading 1")
    doc.add_paragraph("Two semantically verified diff paths share one durable representation. The ablation supports selected sparse/repeated-key benefits and documents cases without a resolved benefit.")
    doc.add_paragraph("Remaining faults and submission requirements",style="Heading 1")
    doc.add_paragraph("Title and Section 7 (Limitations): adaptation is diff selection, not trie geometry; general efficiency remains unsupported. Section 4 (Threshold calibration): the largest candidate wins a bounded cached-query study, not a universal optimum. Section 5 (Public dataset scale): million-record commit and storage costs are substantial. Declarations/Appendix A: funding, contributions, competing interests, ethics/AI accountability and public release of new evidence require author confirmation or action.")
    doc.add_paragraph("Originality and AI review",style="Heading 1")
    doc.add_paragraph("31 real references checked against public sources/metadata; inaccessible full text disclosed. Local exact-sequence screening and manual source/prose review found no flagged unattributed copying in the inspected material. This is not an institutional similarity certificate or AI-authorship determination; proprietary corpora were unavailable. No detector score was invented and no manuscript was uploaded to a checking service.")
    doc.add_paragraph("Checks: semantic/evidence audit and a separate exported-artifact review. Full results, source-verification records and unresolved issues accompany the detailed revision audit.")
    path=ROOT/"output/docs/Revon_Discover_Computing_One_Page_Review.docx";sanitize(doc,"Arial");doc.save(path)

def audit_document(data):
    doc=setup_document(False)
    doc.add_paragraph("Revon Seven Revision Change and Evidence Audit",style="Title")
    doc.add_paragraph("2 October 2026. This report records the seven requested revisions, their evidence and the remaining submission requirements. Independent hardware replication and author-confirmed declarations remain outstanding.")
    for i,title,text in (
        (1,"Primary ablation","Added dedicated synthetic and public M/H tables plus an IQR figure. Primary paired intervals, public exploratory unpaired intervals and three-seed history results remain separate. Negative and inconclusive cases are visible."),
        (2,"Threshold evidence",f"42 repositories, 2,646 checked queries, 2,058 measured queries. Calibration and validation profiles and seeds were defined in advance. T={data['frozen_threshold']:,} was frozen before validation. The original 4,096 workflow evidence is unchanged. The separate 256-commit diagnostic contributes 315 queries and is excluded from threshold selection."),
        (3,"Storage attribution","Added M/H footprint deltas and latency reductions. Both existing variants store changesets. Eight post-compaction samples per system reproduce means of 22.08 MiB for M/H, 2.17 MiB for Dolt SQL and 2.19 MiB for Dolt bulk import. A held-out 100k-row/H64/32,768-operation repository has a 188.59 MiB closed file, 60.34 MiB of canonical node payloads, 15.21 MiB of changesets and 0.02 MiB of commits. Its 113.02 MiB residual does not isolate allocated SQLite index overhead. No changeset-free variant is represented as an equivalent system."),
        (4,"Independent replication","Only the local Windows computer and its WSL environment were available. No remote Linux machine or authorization was provided. Independent replication is blocked, not completed by limitations prose."),
        (5,"Protocol consolidation","Added the environment/bundle table and a coherent timing, correctness, repetition, memory, failure/recovery and statistical protocol. CPU model and NTFS were inspected during revision; historical manifests remain authoritative."),
        (6,"Decision mechanism","Replaced the mixed commit/diff diagram with the actual version-pair decision. Added ancestry eligibility, operation counting, reverse handling, shared output and separate checkout semantics. Forced Merkle's public API still checks ancestry."),
        (7,"Bounded claims","Retained the requested title and explained its ambiguity. The fixed-depth geometry does not adapt. Consolidated limitations and prioritised independent hardware, real histories and broader threshold testing. General efficiency is not established.")):
        doc.add_paragraph(f"{i} {title}",style="Heading 1");doc.add_paragraph(text)
    doc.add_paragraph("Threshold scores and exploratory minima",style="Heading 1")
    doc.add_paragraph("Scores weight the 21 calibration or 21 validation cases equally relative to each case's faster forced-path median. Validation is not used for selection. Per-case candidates with the lowest observed medians are exploratory descriptions and may differ only through timing noise when they select the same path. In Table 2, C denotes calibration and V denotes held-out validation.")
    add_table(doc,dict(number=1,caption="Candidate scores",headers=["T","Calibration","Validation"],rows=[[str(t),f"{data['scores']['calibration'][t]:.6f}",f"{data['scores']['validation'][t]:.6f}"] for t in (1024,2048,4096,8192,16384)],widths=[.2,.4,.4]),False)
    add_table(doc,dict(number=2,caption="Observed case minima in the bounded candidate set",headers=["Case","Operations","T","ms"],rows=[[r["scenario"].replace("calibration-","C ").replace("validation-","V ").replace("-o"+r["operations"],""),r["operations"],r["threshold"],f"{float(r['median_ms']):.3f}"] for r in data["threshold_best"]],widths=[.48,.18,.16,.18]),False)
    doc.add_paragraph("Originality AI and reference verification",style="Heading 1")
    doc.add_paragraph("Fetched public sources for every bibliography entry and compared identifiers, titles, authors, venues and dates. Corrected the OrpheusDB author's initials to A. Parameswaran, added IPDPS pages 479-493 to reference 26, expanded DOI links, replaced the ResearchGate link for reference 17 with the MIT repository record, and repaired the blocked provenance-paper link using the university's alternate host. The MIT full-text bitstream returned HTTP 405; its abstract and Crossref metadata were checked. One Crossref lookup was rate-limited, but the author PDF and IEEE record verified reference 20. Crossref's malformed Merkle Search Trees page range was not blindly copied; the conference citation and author PDF establish the pagination.")
    doc.add_paragraph("Local screening compares twelve-word exact sequences in non-bibliography paragraphs with accessible extracted public source texts. Ref. 4 is a scanned paper inspected visually; ref. 17 has an accessible MIT abstract/metadata record but its full PDF could not be retrieved. Both are excluded from automatic sequence matching. The accessible related-work abstracts and relevant descriptions were manually checked: persistence, content addressing, ordered reconciliation, relational branching, scientific arrays, provenance, semantic differencing and reproducibility are not treated as equivalent contributions. No suspicious copied or close-paraphrase passage was identified in the inspected material. Preprints remain labelled as preprints. Screening is repeated after final edits and its JSON records the findings and coverage.")
    doc.add_paragraph("This does not cover proprietary databases, unpublished documents, every possible source, paraphrase plagiarism or authorship attribution. No institutional similarity checker was available. No AI score or plagiarism-clearance certificate is claimed. Only public sources were downloaded; the private manuscript was not uploaded. Codex assistance in code, prose and plotting is disclosed for author review rather than hidden through detector-evasion edits.")
    refs=json.loads((ROOT/'docs/experiments/REVISION_REFERENCE_VERIFICATION.json').read_text())
    def reference_title(record):
        match=re.search(r'[“"]([^”"]+)[”"]',record['reference'])
        return match[1] if match else record.get('page_title','Public source record')
    add_table(doc,dict(number=3,caption="Reference checks",headers=["Ref","Source record","Verification and access"],rows=[[str(r["number"]),reference_title(r),"Public source/metadata checked; full text limited" if r["number"]==17 else "Public source checked; DOI lookup rate-limited" if r["number"]==20 else "Public source and metadata checked"] for r in refs],widths=[.08,.53,.39]),False)
    doc.add_paragraph("Submission requirements and double check",style="Heading 1")
    doc.add_paragraph("Discover Computing's official guidelines were checked on 2 October 2026: Word figures/tables in the text, a consistent font of at least 12 pt, numbered square-bracket citations, consecutive table/figure numbering and an abstract below 250 words. The separate IEEE document uses two-column body text and full-width displays where required. Both derive from one semantic manuscript model. Funding, contributions, competing interests, ethics and complete AI-use/accountability confirmation are still required from the authors. New evidence remains local until an authorized public release.")
    doc.add_paragraph("Pass 1 checks semantics, frozen source hashes, raw/summary agreement, threshold leakage, statistics, source attribution and unsupported claims. Twenty-seven focused versioning, SQLite and methodological regression tests passed. Pass 2 checks exported PDFs, content agreement, every caption/image placement, table headers and continuation, bounds, citation numbering and rendered pages. The one-page review states the final recommendation without claiming journal approval.")
    doc.add_paragraph("Sources",style="Heading 1")
    doc.add_paragraph("Discover Computing submission guidelines: https://link.springer.com/journal/10791/submission-guidelines. Springer Nature AI preparation policy: https://www.springernature.com/gp/policies/editorial-policies/ai-manuscript-preparation. Reference URLs, retrieved text/PDF hashes, and metadata are in output/docs/seven_revision_audit/reference_verification.json. Threshold analysis, minima and diagnostic ratios are in analysis.json; raw frozen study sources remain under their evidence bundles.")
    path=ROOT/"output/docs/Revon_Seven_Revision_Detailed_Audit.docx";sanitize(doc,"Arial");doc.save(path)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--manuscripts-only',action='store_true');args=parser.parse_args()
    WORK.mkdir(parents=True,exist_ok=True)
    backup=WORK/"originals";backup.mkdir(exist_ok=True)
    original=ROOT/"paper/older versions/sources/Revon_Reviewed_Public_Dataset_Manuscript.docx"
    for path in list((ROOT/"paper").glob("Revon*.docx"))+list((ROOT/"paper").glob("Revon*.pdf")):
        if not (backup/path.name).exists():shutil.copy2(path,backup/path.name)
    assert sha(original)=='22b8c68e42f2f305c1c43d095dbe719556ababa95c38722fed184c85d871b8d3','authoritative reviewed source changed'
    source=Document(original)
    assert len(source.tables)==5 and len(source.inline_shapes)==7,"unexpected authoritative source"
    AUDIT.mkdir(parents=True,exist_ok=True)
    (AUDIT/"authoritative_source.json").write_text(json.dumps(dict(source=str(original),sha256=sha(original),reason="archived approved pagination-fixed source includes completed TLC evaluation; legacy IEEE DOCX excluded"),indent=2))
    data,lookup,threshold,diagnostic=analyze();assets=figures(data,lookup,threshold,diagnostic)
    remap={11:5,5:6,6:7,7:8,8:9,9:10,10:11}
    assets={remap.get(k,k):v for k,v in assets.items()}
    content=blocks(source,data);(AUDIT/"manuscript_model.json").write_text(json.dumps(content,indent=2),encoding="utf-8")
    build_doc(content,assets,ROOT/"paper/Revon_Research_Paper_Discover_Computing.docx",False)
    build_doc(content,assets,ROOT/"paper/Revon_Research_Paper_IEEE.docx",True)
    if not args.manuscripts_only:report(data);audit_document(data)
    print(json.dumps(dict(title=TITLE,blocks=len(content),figures=len(assets),tables=sum(b['kind']=='table' for b in content),frozen_threshold=data['frozen_threshold']),indent=2))

if __name__=="__main__":main()
