"""Remove the unnecessary page break before the conclusion.

The final manuscript already fits comfortably except for a manual page break
that leaves the lower half of page 9 empty and pushes two appendix lines onto
page 11.  This targeted edit preserves all manuscript text and formatting while
allowing Word to paginate the conclusion and end matter naturally.
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def normalized(text: str) -> str:
    return " ".join(text.split())


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    parser.add_argument("--backup", type=Path, required=True)
    args = parser.parse_args()

    source = args.document.resolve()
    backup = args.backup.resolve()
    if not source.is_file():
        raise FileNotFoundError(source)

    backup.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        shutil.copy2(source, backup)

    document = Document(source)
    before_text = [paragraph.text for paragraph in document.paragraphs]
    before_tables = [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]
    before_images = len(document.inline_shapes)

    conclusion_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if normalized(paragraph.text) == "9 Conclusion"
    )
    if conclusion_index == 0:
        raise RuntimeError("Conclusion has no preceding paragraph")

    break_paragraph = document.paragraphs[conclusion_index - 1]
    page_breaks = break_paragraph._p.xpath('.//w:br[@w:type="page"]')
    if break_paragraph.text.strip() or len(page_breaks) != 1:
        raise RuntimeError("Expected one empty manual page-break paragraph before the conclusion")

    break_paragraph._p.getparent().remove(break_paragraph._p)

    # Remove stale pagination hints adjacent to the deleted manual break. Word
    # will regenerate these during PDF export; they do not carry manuscript text.
    conclusion = next(
        paragraph
        for paragraph in document.paragraphs
        if normalized(paragraph.text) == "9 Conclusion"
    )
    for marker in conclusion._p.xpath('.//w:lastRenderedPageBreak'):
        marker.getparent().remove(marker)

    after_text = [paragraph.text for paragraph in document.paragraphs]
    expected_text = before_text[: conclusion_index - 1] + before_text[conclusion_index:]
    if after_text != expected_text:
        raise RuntimeError("Unexpected manuscript text change")
    if before_tables != [[cell.text for row in table.rows for cell in row.cells] for table in document.tables]:
        raise RuntimeError("Unexpected table change")
    if before_images != len(document.inline_shapes):
        raise RuntimeError("Unexpected figure change")

    document.save(source)
    print(f"Removed the manual page break before the conclusion in {source}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
