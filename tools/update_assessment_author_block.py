from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"C:\Users\admin\Downloads\Revon_Title_Abstract_Literature_Review.docx")
OUTPUT = Path(r"C:\Users\admin\Desktop\Revon\output\docs\Revon_Title_Abstract_Literature_Review.docx")


def set_times(run, size_pt: float, *, bold: bool = False, superscript: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.superscript = superscript


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)


def configure_paragraph(paragraph, *, after: float, line_spacing: float = 1.0) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing


def main() -> None:
    document = Document(SOURCE)
    if len(document.paragraphs) < 4:
        raise RuntimeError("Unexpected document structure: author block is missing")

    authors = document.paragraphs[1]
    clear_paragraph(authors)
    configure_paragraph(authors, after=2)

    author_parts = [
        ("Atharva Sheersh Pandey", False),
        ("1", True),
        (", ", False),
        ("Adhyan Jain", False),
        ("1", True),
        (", ", False),
        ("Poornima Nedunchezhian", False),
        ("2", True),
    ]
    for text, superscript in author_parts:
        run = authors.add_run(text)
        set_times(run, 8 if superscript else 12, superscript=superscript)

    emails = document.paragraphs[2]
    clear_paragraph(emails)
    configure_paragraph(emails, after=2)
    email_run = emails.add_run(
        "atharva.sheersh2024@vitstudent.ac.in; adhyan.jain2024@vitstudent.ac.in;\n"
        "poornima.n@vit.ac.in"
    )
    set_times(email_run, 10)

    affiliation = document.paragraphs[3]
    clear_paragraph(affiliation)
    configure_paragraph(affiliation, after=8)
    affiliation_run = affiliation.add_run(
        "School of Computer Science and Engineering\nVIT Vellore"
    )
    set_times(affiliation_run, 11)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
