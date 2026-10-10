"""Negative controls for the current manuscript's numerical and link integrity."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from docx import Document
from tools.check_final_paper import ROOT, check_package
from tools.check_springer_paper import check_springer, check_springer_doc, REPOSITORY


class SpringerPublicationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = ROOT / "paper/Revon Research Paper Springer.docx"
        cls.model, _, _ = check_package(ROOT / "paper/Revon_Linux_Reproducibility_Supplement.zip", ROOT)

    def test_current_pair_passes(self):
        self.assertEqual(check_springer(self.source, self.model)["references"], 31)

    def mutate(self, edit, error):
        with TemporaryDirectory() as directory:
            doc = Document(self.source)
            edit(doc)
            path = Path(directory) / "modified.docx"
            doc.save(path)
            with self.assertRaisesRegex(ValueError, error):
                check_springer_doc(path, self.model)

    def test_modified_result_is_rejected(self):
        self.mutate(lambda d: setattr(d.tables[4].cell(1, 1), "text", "999.999"), "Springer table 5 differs")

    def test_missing_figure_is_rejected(self):
        def remove(d):
            node = d.inline_shapes[0]._inline
            node.getparent().remove(node)
        self.mutate(remove, "Springer figure count changed")

    def test_stale_commit_link_is_rejected(self):
        def pin(d):
            for rel in d.part.rels.values():
                if rel.target_ref == REPOSITORY:
                    rel._target = REPOSITORY + "/tree/" + "a" * 40
        self.mutate(pin, "Springer current GitHub links missing")
