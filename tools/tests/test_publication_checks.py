"""Ensure the publication gate rejects broken artifacts, not just counts them."""
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import zipfile

from docx import Document
import pymupdf

from tools.check_final_paper import (
    ROOT, check_doc, check_hashes, check_package, check_pdf, normalize, sha,
)


class PublicationCheckTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.package = ROOT / "paper/Revon_Linux_Reproducibility_Supplement.zip"
        cls.model, cls.assets, _ = check_package(cls.package, ROOT)
        cls.source = ROOT / "paper/Revon_Research_Paper_Discover_Computing.docx"

    def mutate_document(self, mutation, error):
        with TemporaryDirectory() as directory:
            doc = Document(self.source)
            mutation(doc)
            path = Path(directory) / "modified.docx"
            doc.save(path)
            with self.assertRaisesRegex(ValueError, error):
                check_doc(path, self.model, self.assets, True)

    def test_both_committed_formats_match_the_packaged_model(self):
        for form in ("IEEE", "Discover_Computing"):
            with self.subTest(form=form):
                result = check_doc(ROOT / f"paper/Revon_Research_Paper_{form}.docx",
                                   self.model, self.assets, form == "Discover_Computing")
                self.assertEqual((result["tables"], result["figures"], result["references"]), (11, 7, 31))

    def test_changed_result_cell_is_rejected(self):
        self.mutate_document(lambda doc: setattr(doc.tables[4].cell(1, 1), "text", "999.999"),
                             "table 5 differs")

    def test_missing_limitation_is_rejected(self):
        target = next(b["text"] for b in self.model if b["kind"] == "body" and "32,768" in b["text"])
        def remove(doc):
            paragraph = next(p for p in doc.paragraphs if normalize(p.text) == normalize(target))
            paragraph._p.getparent().remove(paragraph._p)
        self.mutate_document(remove, "missing or reordered manuscript content")

    def test_missing_figure_is_rejected(self):
        def remove(doc):
            shape = doc.inline_shapes[2]._inline
            shape.getparent().remove(shape)
        self.mutate_document(remove, "figure count changed")

    def test_newline_conversion_breaks_collection_checksum(self):
        with TemporaryDirectory() as directory:
            path = Path(directory)
            data = b"collected\nbytes\n"
            (path / "record.txt").write_bytes(data)
            (path / "SHA256SUMS.json").write_text(json.dumps({"record.txt": sha(data)}))
            self.assertEqual(check_hashes(path), 1)
            (path / "record.txt").write_bytes(data.replace(b"\n", b"\r\n"))
            with self.assertRaisesRegex(ValueError, "collection checksum mismatch"):
                check_hashes(path)

    def test_tampered_supplementary_figure_is_rejected(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "tampered.zip"
            with zipfile.ZipFile(self.package) as source, zipfile.ZipFile(path, "w") as target:
                for item in source.infolist():
                    data = b"changed" if item.filename == "figures/Fig3.png" else source.read(item)
                    target.writestr(item, data)
            with self.assertRaisesRegex(ValueError, "package checksum mismatch: figures/Fig3.png"):
                check_package(path, ROOT)

    def test_blank_pdf_is_rejected(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "blank.pdf"
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(path)
            with self.assertRaisesRegex(ValueError, "near-empty PDF page"):
                check_pdf(path)


if __name__ == "__main__":
    unittest.main()
