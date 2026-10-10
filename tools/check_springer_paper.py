"""Check the current reference-format manuscript without rewriting any artifacts.

The archived formats retain their exact packaged-model checks. This companion
checks the current paper's numerical tables, structure, citations, link targets,
and PDF/DOCX agreement after prose and reference-layout revisions.
"""
from pathlib import Path
import re
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pymupdf

REPOSITORY = "https://github.com/atharvasheersh/Revon"
ARCHIVE = REPOSITORY + "/blob/main/paper/Revon_Linux_Reproducibility_Supplement.zip"
TITLE = "Adaptive Merkle Trie Versioning for Efficient Structured Data Management"
# Original packaged-model IDs -> alphabetized bibliography IDs in this format.
REFERENCE_IDS = {24: 1, 13: 2, 14: 3, 7: 4, 6: 5, 9: 6, 12: 7, 19: 8,
                 22: 9, 1: 10, 2: 11, 5: 12, 30: 13, 27: 14, 29: 15,
                 26: 16, 11: 17, 10: 18, 23: 19, 4: 20, 8: 21, 31: 22,
                 20: 23, 21: 24, 17: 25, 3: 26, 25: 27, 18: 28, 16: 29,
                 15: 30, 28: 31}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def compact(text):
    return re.sub(r"\s+", "", text).casefold()


def remap_model_citations(text):
    def replace(match):
        return "[" + ", ".join(str(REFERENCE_IDS[int(n)]) for n in re.findall(r"\d+", match[1])) + "]"
    return re.sub(r"\[([\d, ]+)\]", replace, text)


def check_springer_doc(path, model):
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    require(paragraphs[0].casefold() == TITLE.casefold(), "Springer title changed")
    tables = [b for b in model if b["kind"] == "table"]
    require(len(doc.tables) == len(tables) == 11, "Springer table count changed")
    for observed, expected in zip(doc.tables, tables):
        expected_cells = [[remap_model_citations(c) for c in row] for row in [expected["headers"]] + expected["rows"]]
        require([[c.text for c in row.cells] for row in observed.rows] == expected_cells,
                f"Springer table {expected['number']} differs from recorded model")
    require(len(doc.inline_shapes) == 7, "Springer figure count changed")
    captions = [int(m[1]) for p in paragraphs if (m := re.match(r"^Fig\.?\s+(\d+)\.", p))]
    require(captions == list(range(1, 8)), "Springer figure numbering changed")
    numbers = [int(m[1]) for p in paragraphs if (m := re.match(r"^\[(\d+)\]", p))]
    require(numbers == list(range(1, 32)), "Springer references changed")
    references = {int(re.match(r"^\[(\d+)\]", p)[1]): p for p in paragraphs if re.match(r"^\[\d+\]", p)}
    for block in (b for b in model if b["kind"] == "reference"):
        original = int(re.match(r"^\[(\d+)\]", block["text"])[1])
        # Reformatting moves years and changes author case; source URLs retain identity.
        urls = [url.rstrip(".,") for url in re.findall(r"https?://[^\s]+", block["text"])]
        require(bool(urls) and all(url in references[REFERENCE_IDS[original]] for url in urls),
                f"Springer reference identity changed: original [{original}]")
    body = "\n".join(p for p in paragraphs if not re.match(r"^\[\d+\]", p))
    body += "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    cited = {int(n) for group in re.findall(r"\[([\d, ]+)\]", body) for n in re.findall(r"\d+", group)}
    require(cited == set(numbers), "Springer uncited or undefined reference")
    for marker in ("516", "72", "32,768", "16,384", "poornima.n@vit.ac.in"):
        require(marker in body, f"Springer missing evidence marker: {marker}")
    require("Online Resource 1" not in body, "Springer still claims journal supplementary hosting")
    availability = next(p for p in paragraphs if p.startswith("Data and code availability:"))
    start = next(i for i, p in enumerate(paragraphs) if p.startswith("APPENDIX A:"))
    appendix = paragraphs[start + 1:]
    require(not re.search(r"\bcommit\b", availability + "\n" + "\n".join(appendix)),
            "Springer availability statement still pins a repository commit")
    targets = {rel.target_ref for rel in doc.part.rels.values() if rel.reltype == RT.HYPERLINK}
    require(REPOSITORY in targets and ARCHIVE in targets, "Springer current GitHub links missing")
    require(not any(re.search(r"github\.com/atharvasheersh/Revon/(?:tree|blob)/[0-9a-f]{40}(?:/|$)", t)
                    for t in targets), "Springer GitHub link still points to a fixed commit")
    return dict(tables=11, figures=7, references=31), availability, appendix


def check_springer(path, model):
    result, availability, appendix = check_springer_doc(path, model)
    with pymupdf.open(Path(path).with_suffix(".pdf")) as pdf:
        require(len(pdf) > 0, "Springer PDF empty")
        text = "".join(page.get_text() for page in pdf)
        for i, page in enumerate(pdf, 1):
            require(len(page.get_text().split()) >= 8, f"Springer near-empty PDF page {i}")
            for block in page.get_text("dict")["blocks"]:
                require(page.rect.contains(pymupdf.Rect(block["bbox"])), f"Springer PDF content outside page {i}")
        numbers = [int(n) for n in re.findall(r"\bFig\.?\s+(\d+)\.", text)]
        require(sorted(numbers) == list(range(1, 8)), "Springer PDF figure captions changed")
        require("Online Resource 1" not in text, "Springer PDF supplementary label is stale")
        require(compact(availability.split(":", 1)[1]) in compact(text), "Springer PDF/DOCX availability differs")
        for p in appendix:
            require(compact(p) in compact(text), "Springer PDF/DOCX appendix differs")
        links = {link.get("uri", "") for page in pdf for link in page.get_links()}
        require(REPOSITORY in links and ARCHIVE in links, "Springer PDF GitHub links missing")
        require(not any(re.search(r"github\.com/atharvasheersh/Revon/(?:tree|blob)/[0-9a-f]{40}(?:/|$)", t)
                        for t in links), "Springer PDF GitHub link still pins a commit")
        result.update(pages=len(pdf), github_links="current repository and main archive", pdf_docx_availability_agrees=True)
    return result
