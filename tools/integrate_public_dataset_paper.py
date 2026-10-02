"""Integrate audited Comment 1 results into the four approved manuscript sources."""
from __future__ import annotations
import argparse
import copy
import csv
import hashlib
import json
import shutil
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".codex_tmp/public_dataset_paper"
PAPERS = (
    ("paper/Revon_Research_Paper_IEEE_Pagination_Fixed.docx", "paper/Revon_Research_Paper_IEEE.pdf"),
    ("paper/Revon_Final_Research_Paper.docx", "paper/Revon_Final_Research_Paper.pdf"),
    ("paper/Revon_Research_Paper_Single_Column.docx", "paper/Revon_Research_Paper_Single_Column.pdf"),
    ("output/docs/Revon_Research_Paper_BERT.docx", "output/docs/Revon_Research_Paper_BERT.pdf"),
)
MODELS = ("Snapshot", "Log-only", "Revon-M (forced Merkle)", "Revon-H", "Dolt", "Dolt (bulk import)")
SHORT = ("Snapshot", "Log-only", "Revon-M", "Revon-H", "Dolt SQL", "Dolt CSV")
COLORS = ("#687989", "#C98A13", "#8767A2", "#007F87", "#C75632", "#296C9E")


def evidence(directory):
    audit = json.loads((directory / "evidence_audit.json").read_text())
    if not audit["passed"]:
        raise ValueError(f"Evidence audit failed: {directory}")
    with (directory / "summary.csv").open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    manifest = json.loads((directory / "manifest.json").read_text())
    return rows, manifest, audit


def plot_scale(rows, path, width=3.45):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8,
                         "axes.labelsize": 8, "xtick.labelsize": 8, "ytick.labelsize": 8,
                         "pdf.fonttype": 42, "svg.fonttype": "none"})
    fig, axes = plt.subplots(3, 1, figsize=(width, 5.85))
    metrics = (("incremental_commit_ms", "(a) Incremental commit (ms)", 1),
               ("diff_ms", "(b) Initial-to-final diff (ms)", 1),
               ("storage_bytes", "(c) Repository footprint (MiB)", 2**20))
    values = []
    for ax, (metric, label, unit) in zip(axes, metrics):
        for model, short, color, marker, style in zip(MODELS, SHORT, COLORS, ("o", "s", "^", "D", "v", "P"), ("-", "-", "--", "-", "-.", ":")):
            group = []
            for n in (10000, 100000, 1000000):
                scenario = f"tlc-n{n}-h10-c{n//1000}"
                match = [r for r in rows if r["scenario"] == scenario and r["model"] == model]
                if match:
                    r = match[0]
                    group.append((n, r))
                    values.append({"panel": metric, "scenario": scenario, "model": model,
                        "median": float(r[f"{metric}_median"]), "p25": float(r[f"{metric}_p25"]),
                        "p75": float(r[f"{metric}_p75"]), "trials": int(r["trials"])})
            if not group:
                continue
            x = [n for n, _ in group]
            y = [float(r[f"{metric}_median"]) / unit for _, r in group]
            lo = [y[i] - float(r[f"{metric}_p25"]) / unit for i, (_, r) in enumerate(group)]
            hi = [float(r[f"{metric}_p75"]) / unit - y[i] for i, (_, r) in enumerate(group)]
            ax.errorbar(x, y, yerr=[lo, hi], color=color, marker=marker, ls=style,
                        lw=1.1, ms=4, capsize=2, elinewidth=.7, label=short, alpha=.92)
        ax.set_xscale("log"); ax.set_yscale("log")
        ax.set_xticks((10000, 100000, 1000000), ("10k", "100k", "1M"))
        ax.set_title(label.split(" (")[0], loc="left", fontsize=8, pad=4)
        ax.set_ylabel("Footprint (MiB)" if metric == "storage_bytes" else "Time (ms)")
        ax.grid(which="major", color="#E0E6EA", lw=.5)
        ax.set_axisbelow(True)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    axes[-1].set_xlabel("Initial records (log scale)")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=2, fontsize=8, frameon=False,
               bbox_to_anchor=(.52, .005), columnspacing=.9, handlelength=1.6)
    fig.subplots_adjust(left=.25, right=.96, top=.935, bottom=.22, hspace=.48)
    fig.savefig(path, dpi=600, facecolor="white")
    fig.savefig(path.with_suffix(".svg"), facecolor="white")
    plt.close(fig)
    (path.parent / "plot_values.json").write_text(json.dumps(values, indent=2))


def find(doc, prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def insert_before(doc, anchor, text, template, keep=False):
    p = doc.add_paragraph()
    p._p.clear()
    if template._p.pPr is not None:
        properties = copy.deepcopy(template._p.pPr)
        for section in properties.findall(qn("w:sectPr")):
            properties.remove(section)
        p._p.append(properties)
    r = p.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        r._r.insert(0, copy.deepcopy(template.runs[0]._r.rPr))
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.page_break_before = False
    anchor._p.addprevious(p._p)
    return p


def replace_text(p, text):
    # Preserve paragraph geometry and the existing leading run's appearance.
    template = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    is_abstract = text.startswith("Abstract:") and len(p.runs) > 1 and p.runs[0].text == "Abstract:"
    body_template = copy.deepcopy(max(p.runs[1:], key=lambda r: len(r.text))._r.rPr) if is_abstract else None
    for node in list(p._p):
        if node.tag != qn("w:pPr"):
            p._p.remove(node)
    run = p.add_run("Abstract:" if is_abstract else text)
    if template is not None:
        run._r.insert(0, template)
    if is_abstract:
        body = p.add_run(text[len("Abstract:"):])
        if body_template is not None:
            body._r.insert(0, body_template)


def table(doc, anchor, headers, rows, width, fractions):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, fraction in zip(t.columns, fractions):
        col.width = Inches(width * fraction)
    for index, (row_data, row) in enumerate([(headers, t.rows[0])] + [(data, t.add_row()) for data in rows]):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit"); trpr.append(no_split)
        for column_index, (cell, value, fraction) in enumerate(zip(row.cells, row_data, fractions)):
            cell.width = Inches(width * fraction)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("left", "right"):
                node = OxmlElement(f"w:{side}"); node.set(qn("w:w"), "58"); node.set(qn("w:type"), "dxa"); margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            cell.text = str(value)
            for p in cell.paragraphs:
                p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if len(headers) == 5 and column_index >= 2 and index > 0 else WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.right_indent = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.keep_with_next = index == 0
                p.paragraph_format.line_spacing = 1
                for r in p.runs:
                    r.font.name = "Times New Roman"; r.font.size = Pt(8)
                    r.bold = index == 0
    header = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(header)
    # Both additions fit in a manuscript column; keep each table with its header.
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "bottom", "insideH"):
        b = OxmlElement(f"w:{name}"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "B6BEC5"); borders.append(b)
    t._tbl.tblPr.append(borders)
    anchor._p.addprevious(t._tbl)
    return t


def reference(doc, anchor, template, author_year=False):
    text = ('New York City Taxi and Limousine Commission. 2024. TLC Trip Record Data: Yellow Taxi, January 2024. Public dataset. Retrieved October 1, 2026 from https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page.' if author_year else '[31] New York City Taxi and Limousine Commission, "TLC Trip Record Data: Yellow Taxi, January 2024," public dataset, accessed Oct. 1, 2026. [Online]. Available: https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page.')
    p = insert_before(doc, anchor, text, template)
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), "31001"); start.set(qn("w:name"), "ref_31")
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), "31001")
    p._p.insert(1, start); p._p.append(end)


def link_dataset_citations(doc, author_year=False):
    paragraphs = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
    for p in paragraphs:
        if p.text.startswith("[31]") or "[31]" not in p.text:
            continue
        for run in list(p.runs):
            if "[31]" not in run.text:
                continue
            before, after = run.text.split("[31]", 1)
            run.text = before
            link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), "ref_31")
            r = OxmlElement("w:r")
            if run._r.rPr is not None:
                r.append(copy.deepcopy(run._r.rPr))
            text = OxmlElement("w:t"); text.text = "(NYC TLC, 2024)" if author_year else "[31]"; r.append(text); link.append(r)
            run._r.addnext(link)
            tail = copy.deepcopy(run._r)
            for node in tail.findall(qn("w:t")):
                node.text = after
            link.addnext(tail)


def integrate(directories):
    WORK.mkdir(parents=True, exist_ok=True)
    rows, manifests, audits = [], [], []
    for directory in directories:
        r, m, a = evidence(directory); rows += r; manifests.append(m); audits.append(a)
    lookup = {(r["scenario"], r["model"]): r for r in rows}
    if len(lookup) != len(rows):
        raise ValueError("Duplicate configuration summaries across bundles")
    if any(r["all_correct"] != "True" for r in rows):
        raise ValueError("Incorrect results")
    plot_scale(rows, WORK / "public_dataset_scale.png")
    successes = sum(a["ok"] for a in audits)
    total = sum(a["executions"] for a in audits)
    max_n = max(int(s["rows"]) for m in manifests for s in m["workloads"]
                if any(r["scenario"] == s["name"] for r in rows))
    baseline = lookup[("tlc-n100000-h10-c100", "Revon-H")]
    large = lookup.get(("tlc-n1000000-h10-c1000", "Revon-H"))
    large_m = lookup.get(("tlc-n1000000-h10-c1000", "Revon-M (forced Merkle)"))
    repetitions = {(m["warmups"], m["measured_trials"]) for m in manifests}
    if len(repetitions) == 1:
        warmups, trials = next(iter(repetitions))
        protocols = f"Each configuration scheduled {warmups} warm-ups and {trials} measured trials."
    else:
        protocols = "Repetition counts are recorded per configuration in the evidence manifests: " + "; ".join(f"{m['warmups']} warm-ups and {m['measured_trials']} measured trials for the {max(s['rows'] for s in m['workloads']):,}-record bundle" for m in manifests) + "."
    fmt = lambda r, field: f"{float(r[field+'_median']):.2f} [{float(r[field+'_p25']):.2f}, {float(r[field+'_p75']):.2f}]"
    scale_claim = (f"At one million records, Revon-H's median incremental commit was {fmt(large, 'incremental_commit_ms')} ms and its initial-to-final diff was {fmt(large, 'diff_ms')} ms. Its median repository footprint was {float(large['storage_bytes_median'])/2**20:.2f} MiB. " if large else f"The largest completed Revon-H configuration contained {max_n:,} records. ")
    if large and large_m:
        scale_claim += "Both Revon variants selected Merkle differencing for the 10,000 accumulated operations at one million records; this configuration does not test a log-path benefit. "
    if large:
        comparison = [(model, short, lookup.get(("tlc-n1000000-h10-c1000", model))) for model, short in zip(MODELS, SHORT)]
        slower_than = [short for model, short, r in comparison if r and model not in {"Revon-H", "Revon-M (forced Merkle)"}
            and float(r["incremental_commit_ms_median"]) < float(large["incremental_commit_ms_median"])]
        if slower_than:
            scale_claim += "Revon-H's incremental-commit median exceeded those of " + ", ".join(slower_than) + " at this scale. "
        dolt = lookup.get(("tlc-n1000000-h10-c1000", "Dolt"))
        bulk = lookup.get(("tlc-n1000000-h10-c1000", "Dolt (bulk import)"))
        def seconds(r, field):
            return f"{float(r[field+'_median'])/1000:.2f} [{float(r[field+'_p25'])/1000:.2f}, {float(r[field+'_p75'])/1000:.2f}]"
        scale_claim += f"Initial import for Revon-H took {seconds(large, 'initial_import_ms')} s"
        if dolt:
            scale_claim += f", compared with {seconds(dolt, 'initial_import_ms')} s for Dolt SQL"
        if bulk:
            scale_claim += f" and {seconds(bulk, 'initial_import_ms')} s for Dolt CSV import"
        scale_claim += ". These initial loads are separate from incremental commits. "
        if dolt:
            ratio = float(large["storage_bytes_median"]) / float(dolt["storage_bytes_median"])
            scale_claim += f"The ratio of median Revon-H/Dolt SQL repository footprints was {ratio:.2f}; this includes serialization, indexes and compression effects and is not a pure metadata comparison. "
    scale_claim += "Brackets give the interquartile range across measured trial summaries, not confidence intervals or variation across datasets. Full import, checkout, storage and failure records remain in the evidence bundle."
    failures = [f for a in audits for f in a["failures"]]
    successful_repetitions = sorted({int(r["trials"]) for r in rows})
    protocols += (f" Successful measured-trial counts range from {successful_repetitions[0]} to {successful_repetitions[-1]} per configuration."
                  if len(successful_repetitions) > 1 else f" All configurations have {successful_repetitions[0]} successful measured trials.")
    if failures:
        scale_claim += " Failed executions are retained and excluded from performance summaries."
        if manifests[0].get("recovery_policy"):
            scale_claim += " One interrupted million-record Revon-M attempt was followed by a fresh recovery trial after the matrix, under identical frozen source and workload. Seven successful measured trials enter each summary; the failed attempt remains in the original bundle. The recovery was outside the counterbalanced block and host memory conditions were not held constant."
        if any((d / "interruption.json").exists() for d in directories):
            scale_claim += " One measured Revon-M trial at one million records lost its worker during an interrupted run; six successful trials remain for that configuration. Its retained worker report includes abnormally long timing intervals and is not used to recover a performance point. The interruption does not establish an algorithmic scale limit."
    report = []
    for source, pdf in PAPERS:
        path = ROOT / source
        backup = WORK / "backups" / source
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            shutil.copy2(path, backup)
        doc = Document(backup)
        if "BERT" in source:
            # The supplied BERT source contains a second, stale abstract.
            duplicates = [p for p in doc.paragraphs if p.text.startswith("Versioning structured datasets requires durable history")]
            if len(duplicates) != 1 or duplicates[0]._p.xpath(".//w:sectPr"):
                raise ValueError("Expected one removable stale BERT abstract")
            duplicates[0]._p.getparent().remove(duplicates[0]._p)
            find(doc, "VIT Vellore").alignment = WD_ALIGN_PARAGRAPH.LEFT
            # BERT places existing captions after tables: keep each complete
            # table and its following caption together after text reflow.
            for existing_table in doc.tables:
                for row in existing_table.rows:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
        obsolete = "Our results do not establish performance on public datasets, independent hardware, concurrent workloads, or production systems."
        updated = "The public-record extension covers one release with generated updates. It does not establish performance across diverse datasets, independent hardware, concurrent workloads, or production systems."
        replaced = 0
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if obsolete in run.text:
                    run.text = run.text.replace(obsolete, updated)
                    replaced += 1
        if replaced != 1:
            raise ValueError(f"Expected one obsolete public-data exclusion in {source}, found {replaced}")
        body = find(doc, "Structured-data version control")
        numbered = "Single_Column" in source or "BERT" in source
        heading = find(doc, "4.3 Repetition" if numbered else "C. Repetition")
        caption = find(doc, "Fig. 6.")
        wide = "Single_Column" in source
        section = doc.sections[1]
        cols = section._sectPr.find(qn("w:cols"))
        count = int(cols.get(qn("w:num"), "1"))
        gap = int(cols.get(qn("w:space"), "0")) / 1440
        column_width = (section.page_width.inches-section.left_margin.inches-section.right_margin.inches-gap*(count-1))/count
        width = min(6.25 if wide else 3.45, column_width)
        figure_width = min(3.45, column_width)
        if "BERT" in source:
            for shape in doc.inline_shapes:
                if shape.width.inches > column_width:
                    ratio = column_width / shape.width.inches
                    shape.height = int(shape.height * ratio)
                    shape.width = Inches(column_width)
        asset = WORK / f"public_dataset_scale_{figure_width:.3f}.png"
        plot_scale(rows, asset, figure_width)
        retained_figures = directories[0] / "figures"
        retained_figures.mkdir(parents=True, exist_ok=True)
        for extension in (".png", ".svg"):
            shutil.copy2(asset.with_suffix(extension), retained_figures / asset.with_suffix(extension).name)
        shutil.copy2(WORK / "plot_values.json", retained_figures / "plot_values.json")
        method_anchor = find(doc, "5. RESULTS" if "Single_Column" in source else "5 RESULTS" if "BERT" in source else "V. RESULTS")
        insert_before(doc, method_anchor, ("4.4 " if numbered else "D. ") + "Public Dataset and Scalability Evaluation", heading, True)
        insert_before(doc, method_anchor,
            "We extend the synthetic evaluation with January 2024 yellow-taxi records from the New York City Taxi and Limousine Commission (NYC TLC) [31]. The checksum-pinned Parquet file contains 2,964,624 rows; a dispersed affine permutation selects nested subsets of 10,000, 100,000 and 1,000,000 rows using seed 20261001. We retain eleven fields: vendor, pickup/drop-off times, passenger count, distance, pickup/drop-off locations, payment type, fare, tip and total. Keys combine the release, file-checksum prefix and physical row ordinal, and therefore identify rows within this frozen file rather than trips across releases. Compact sorted-key UTF-8 JSON preserves values; nulls remain null. The million-row subset has 47,277 missing passenger counts and no serialization failures. Identical field values remain separate rows. The affine sample is not a uniformly random permutation.", body)
        insert_before(doc, method_anchor,
            "Monthly trip files contain different trips, so we generate synthetic correction histories over these public records. Each commit samples distinct keys with a fixed seed, increases fare and total by 0.01, and preserves the other fields; keys may repeat across commits. Table IV specifies the separate scale, history and density sweeps. All systems receive identical states and histories. We retain b=8, d=4 and the 4,096-operation selector. Diff compares version 1 with version H+1. Import is separate from incremental version creation; each trial reports the median of H commit intervals. Canonical output materialization is included in diff and checkout. Every trial checks both endpoint states and semantic diff; the first measured trial for each system/configuration checks every historical version outside timing. Directory footprint is measured before close and compaction. These are whole-system workflow measurements.", body)
        insert_before(doc, method_anchor,
            f"The public-data run used Python {manifests[0]['python']} on the same Windows host, with {manifests[0]['physical_memory_bytes']/2**30:.2f} GiB visible RAM and {manifests[0]['available_memory_at_start_bytes']/2**30:.2f} GiB available at launch. Available memory was not held constant. {protocols} Clean workers and counterbalanced order follow the existing protocol; pilot measurements are excluded. Each trial has a 900-second wall-clock limit including setup and verification. We report medians and linearly interpolated interquartile ranges. {successes} of {total} retained executions completed successfully. The source checksum, workload digests, raw trials, environment metadata, verification hashes and reproduction commands accompany the new evidence; prior bundles are preserved.", body)
        insert_before(doc, method_anchor, "TABLE IV. Public TLC dataset [31], January 2024; all histories are generated corrections.", caption, True)
        table(doc, method_anchor, ("Parameter", "Dataset and sweep specification"), [
            ("Dataset", "NYC TLC yellow taxi, January 2024 [31]"),
            ("Records", "2,964,624 source rows; 1,000,000 retained; nested subsets of 10k, 100k and 1M"),
            ("Fields", "Vendor; pickup/drop-off times; passenger count; distance; pickup/drop-off locations; payment type; fare; tip; total"),
            ("Key", "Month + source-checksum prefix + physical row ordinal; identity within the frozen file"),
            ("Histories", "Generated fare/total corrections; 10 or 50 incremental commits; no observed release histories"),
            ("Scale", "N=10k/100k/1M; H=10; C=10/100/1,000 per commit (0.1%)"),
            ("History", "N=100k; H=10/50; C=100 per commit (0.1%)"),
            ("Density", "N=100k; H=10; C=10/100/1,000 per commit (0.01%/0.1%/1%)"),
            ("Sampling", "Seed 20261001; nested prefixes of a dispersed affine permutation; not a uniform random permutation")], width, (.25, .75))
        # Existing Table III occurs after the new methodology table; renumber by reading order.
        replace_text(find(doc, "TABLE III."), find(doc, "TABLE III.").text.replace("TABLE III.", "TABLE IV.", 1))
        replace_text(find(doc, "TABLE IV. Public"), find(doc, "TABLE IV. Public").text.replace("TABLE IV.", "TABLE III.", 1))
        for p in doc.paragraphs:
            if "Table IV specifies" in p.text:
                replace_text(p, p.text.replace("Table IV specifies", "Table III specifies"))
            elif "Table III gives" in p.text:
                replace_text(p, p.text.replace("Table III gives", "Table IV gives"))
        result_anchor = find(doc, "6. DISCUSSION" if "Single_Column" in source else "6 DISCUSSION" if "BERT" in source else "VI. DISCUSSION")
        insert_before(doc, result_anchor, ("5.5 " if numbered else "E. ") + "Public Dataset Scale History and Update Density", heading, True)
        figp = insert_before(doc, result_anchor, "", body, True)
        figp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figp.paragraph_format.first_line_indent = Pt(0)
        figp.paragraph_format.left_indent = Pt(0)
        figp.paragraph_format.right_indent = Pt(0)
        figp.paragraph_format.line_spacing = 1
        figp.paragraph_format.keep_together = True
        figp.paragraph_format.space_before = Pt(0)
        figp.paragraph_format.space_after = Pt(0)
        figp.add_run().add_picture(str(asset), width=Inches(figure_width))
        sample_note = "seven successful measured trials per configuration" if successful_repetitions == [7] else "six successful trials for Revon-M at 1M, seven elsewhere"
        insert_before(doc, result_anchor, "Fig. 7. Public-record scalability with ten generated commits and 0.1% updates per commit; medians and interquartile ranges (" + sample_note + "). Both axes use logarithmic scales. Revon-H uses log diff at 10k/100k and Merkle diff at 1M.", caption)
        result_paragraphs = [scale_claim]
        for boundary in ("Initial import for", "Brackets give", "Failed executions"):
            last = result_paragraphs.pop()
            index = last.find(boundary)
            result_paragraphs.extend([last[:index].strip(), last[index:].strip()] if index > 0 else [last])
        for paragraph in result_paragraphs:
            insert_before(doc, result_anchor, paragraph, body)
        insert_before(doc, result_anchor, "TABLE V. Public-data history and density results at 100,000 records; median commit and diff times in ms and footprint in MiB. Full interquartile ranges are retained in summary.csv.", caption, True)
        result_rows = []
        for scenario, label in (("tlc-n100000-h10-c100", "10 / 0.1%"), ("tlc-n100000-h50-c100", "50 / 0.1%"), ("tlc-n100000-h10-c10", "10 / 0.01%"), ("tlc-n100000-h10-c1000", "10 / 1%")):
            for model, short in zip(MODELS, SHORT):
                r = lookup.get((scenario, model))
                result_rows.append((label, short, f"{float(r['incremental_commit_ms_median']):.2f}" if r else "Failed",
                    f"{float(r['diff_ms_median']):.2f}" if r else "--", f"{float(r['storage_bytes_median'])/2**20:.2f}" if r else "--"))
        table(doc, result_anchor, ("H / density", "System", "Commit", "Diff", "MiB"), result_rows, width, (.24, .25, .18, .16, .17))
        insert_before(doc, result_anchor,
            "At 100,000 records, the 10-commit 0.1% case accumulates 1,000 operations, while the 50-commit case accumulates 5,000 and the 1% case accumulates 10,000. Revon-H therefore uses the log path for the sparse short history and the Merkle path for the latter cases. This sweep evaluates the frozen selector under public record contents; it does not establish a new threshold or universal advantage. The update-only histories differ from the earlier mixed synthetic workloads.", body)
        discussion_anchor = find(doc, "7. LIMITATIONS" if "Single_Column" in source else "7 LIMITATIONS" if "BERT" in source else "VII. LIMITATIONS")
        insert_before(doc, discussion_anchor,
            "The public-data sweeps extend the evidence to real record contents and larger states, but the histories remain generated. At fixed depth, increasing N increases expected leaf occupancy, so rewriting touched leaves can become more expensive even with an unchanged update percentage. The scale sweep also crosses the frozen selector threshold: Revon-H uses log diff at 10k/100k records and Merkle diff at 1M. Its diff curve therefore combines larger states with a strategy change. Longer histories likewise accumulate operations and can move Revon-H to the Merkle path. Figure 7 and Table V describe this host and record mapping; they do not isolate a tree algorithm, test relational queries, or measure naturally occurring corrections.", body)
        replace_text(find(doc, "Dataset scope:"),
            f"Dataset scope: The earlier synthetic study reaches 100,000 rows; the new public-record evaluation completes configurations up to {max_n:,} rows from one TLC monthly file with eleven retained fields. All public-data updates are generated, the scale sample uses one seed and an affine permutation, and key identities are valid only within the frozen release. We did not evaluate observed histories, multiple public datasets, multi-table workloads, relational queries or scale beyond the completed configurations.")
        memory = find(doc, "Memory:")
        replace_text(memory, memory.text + " Public-data trials also include workload/oracle memory. The host had limited available RAM at launch; memory pressure and paging costs were not isolated from the workflow, so the measured scaling should not be read as an algorithm-only complexity result.")
        future = find(doc, "Moving beyond a linear prototype")
        replace_text(future, future.text.replace("Evaluation also needs public structured datasets, workloads with one million rows, and replications by independent researchers on separate hardware.", "Evaluation also needs additional public datasets, observed update histories, larger feasible states, and replications by independent researchers on separate hardware."))
        conclusion = find(doc, "Revon demonstrates one way")
        replace_text(conclusion, conclusion.text.replace("The experiments do not establish performance on real schemas or public datasets, or across independent hardware, concurrent workloads, and ordered range queries.", f"The public-record extension reaches {max_n:,} records from one TLC release using generated corrections. It does not establish performance across datasets, observed histories, independent hardware, concurrent workloads or ordered range queries."))
        abstract = find(doc, "Abstract:")
        abstract_text = abstract.text.replace("on deterministic workloads of up to", "on synthetic workloads of up to")
        abstract_text = abstract_text.replace("For the tested workflows, median paired", "In that synthetic study, median paired")
        abstract_text = abstract_text.replace("We ran the experiments on one Windows host and under WSL2 on that same computer.", "We used one Windows host; the earlier synthetic study also includes WSL2 on that computer.")
        replace_text(abstract, abstract_text.replace("We used one Windows host;", f"A separate public TLC record evaluation reaches {max_n:,} records with generated update histories. We used one Windows host;"))
        appendix = find(doc, "APPENDIX A.")
        reference_template = next(p for p in reversed(doc.paragraphs) if any(b.get(qn("w:name"), "").startswith("ref_") for b in p._p.xpath(".//w:bookmarkStart")))
        reference(doc, appendix, reference_template, "BERT" in source)
        availability = find(doc, "The repository includes")
        availability.add_run(" Public dataset acquisition and preparation are documented in docs/experiments/PUBLIC_DATASET_EVALUATION.md; experiments/public_dataset.py and experiments/public_dataset_study.py regenerate the checksum-pinned workload and run its sweeps. The TLC source is cited in [31]. New bundles are " + ", ".join(str(p.relative_to(ROOT)).replace('\\', '/') for p in directories) + ". Raw and prepared trip records are downloaded locally and excluded from Git; retained timings, manifests, verification hashes and audited summaries are available with the code.")
        link_dataset_citations(doc, "BERT" in source)
        doc.save(path)
        report.append({"docx": str(path), "pdf": str(ROOT / pdf), "before_sha256": hashlib.sha256(backup.read_bytes()).hexdigest(), "after_sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    (WORK / "integration.json").write_text(json.dumps({"sources": [str(d) for d in directories], "outputs": report, "max_completed_records": max_n}, indent=2))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--evidence", nargs="+", type=Path, required=True)
    a = p.parse_args()
    integrate([d.resolve() for d in a.evidence])
