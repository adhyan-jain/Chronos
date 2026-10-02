from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = Path(r"C:\Users\admin\Desktop\Revon\paper\Revon_Cover_Letter_Computing.docx")


def set_run_font(run, name="Times New Roman", size=Pt(11.5), bold=False, italic=False):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = None
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_text_paragraph(doc, text="", *, before=0, after=7, line=1.08, keep=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_together = keep
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11.5)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08

props = doc.core_properties
props.title = "Cover Letter for Revon Manuscript Submission"
props.subject = "Submission to Computing"
props.author = "Atharva Sheersh Pandey"
props.keywords = "Revon, cover letter, Computing, Springer Nature"

sender = add_text_paragraph(doc, after=2, line=1.0)
sender.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = sender.add_run("Atharva Sheersh Pandey")
set_run_font(run, size=Pt(12), bold=True)

for value in (
    "School of Computer Science and Engineering",
    "VIT Vellore",
    "atharva.sheersh2024@vitstudent.ac.in",
):
    p = add_text_paragraph(doc, after=1, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run(value), size=Pt(10.5))

date = add_text_paragraph(doc, "11 September 2026", before=9, after=12, line=1.0)

for value in ("The Editors", "Computing", "Springer Nature"):
    p = add_text_paragraph(doc, after=1, line=1.0)
    run = p.add_run(value)
    set_run_font(run, italic=(value == "Computing"))

subject = add_text_paragraph(doc, before=10, after=10, line=1.05, keep=True)
set_run_font(subject.add_run("Subject: "), bold=True)
set_run_font(subject.add_run("Submission of the Revon research manuscript"))

add_text_paragraph(doc, "Dear Editors,", after=8, keep=True)

p = add_text_paragraph(doc)
set_run_font(p.add_run("On behalf of my co-authors, Adhyan Jain and Poornima Nedunchezhian, I am submitting our manuscript, "))
set_run_font(
    p.add_run(
        '"Revon: A Content-Addressed Versioned Data Store: A Git-Inspired Approach to Structured Data Versioning Using Merkle Hash Tries,"'
    ),
    italic=True,
)
set_run_font(p.add_run(" for consideration as a Research article in "))
set_run_font(p.add_run("Computing"), italic=True)
set_run_font(p.add_run("."))

add_text_paragraph(
    doc,
    "The manuscript presents Revon, a Python prototype for versioning structured key-value data with an immutable fixed-depth Merkle hash trie, content-addressed objects, atomic SQLite persistence, and a hybrid diff policy. The hybrid design selects between addressed change-log aggregation and hash-pruned Merkle comparison using a threshold calibrated separately from the reported evaluation.",
)

add_text_paragraph(
    doc,
    "We evaluate Revon against snapshot and log-only baselines, a forced-Merkle ablation, and Dolt 2.3.1 on deterministic workloads of up to 100,000 rows. The corrected run contains 333 executions: 74 warm-ups and 259 measured runs, including 175 in the primary five-system evaluation. Dolt's structured diff output is checked against the workload oracle. Under the tested in-process Revon versus Dolt CLI workflow, Revon-H had lower median diff, commit, and checkout latency in the reported scenarios, while using more repository storage and showing slower initial import at the largest scale. The manuscript reports these measured trade-offs without claiming universal superiority over a production SQL system.",
)

add_text_paragraph(
    doc,
    "We believe the work fits Computing because it connects persistent data structures, content addressing, database versioning, and reproducible experimental systems evaluation. The manuscript also isolates the effect of hybrid diff selection through an ablation and reports the sensitivity of trie geometry to latency and storage objectives.",
)

add_text_paragraph(
    doc,
    "All authors have approved the manuscript and its submission. We confirm that the manuscript is original, has not been published previously, and is not under consideration by another journal. The source code, benchmark harness, and audited evidence are identified in the manuscript to support reproducibility.",
)

add_text_paragraph(doc, "Thank you for considering our manuscript.", after=11)
add_text_paragraph(doc, "Sincerely,", after=11, keep=True)

signature = add_text_paragraph(doc, after=1, line=1.0, keep=True)
set_run_font(signature.add_run("Atharva Sheersh Pandey"), bold=True)

for value in (
    "on behalf of Adhyan Jain and Poornima Nedunchezhian",
    "School of Computer Science and Engineering, VIT Vellore",
    "atharva.sheersh2024@vitstudent.ac.in",
):
    p = add_text_paragraph(doc, after=1, line=1.0, keep=True)
    set_run_font(p.add_run(value), size=Pt(10.5))

doc.save(OUTPUT)
print(OUTPUT)
