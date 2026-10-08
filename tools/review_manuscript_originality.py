"""Bounded, reproducible phrase screening; no detector or Turnitin score."""
from pathlib import Path
from collections import defaultdict, Counter
import hashlib,json,re,unicodedata
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/pdf/revon-integrity-review-20261008'
CORPUS=ROOT/'output/docs/seven_revision_audit'

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def tokens(text):
    text=unicodedata.normalize('NFKC',text).lower()
    text=re.sub(r'https?://\S+',' ',text)
    text=re.sub(r'\[[\d, -]+\]',' ',text)
    return re.findall(r'[a-z0-9]+',text)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    docpath=ROOT/'paper/Revon_Research_Paper_Discover_Computing.docx'
    pdfpath=docpath.with_suffix('.pdf'); doc=Document(docpath)
    model=json.loads((ROOT/'output/docs/linux-final-20261008/manuscript_model.json').read_text(encoding='utf-8'))
    metadata=json.loads((CORPUS/'reference_verification.json').read_text(encoding='utf-8'))
    pdf=PdfReader(pdfpath); pages=[p.extract_text() for p in pdf.pages]
    paragraphs=[];section='Front matter'
    for b in model:
        if b['kind'].startswith('h'):section=b['text']
        if b['kind'] not in ('abstract','body'):continue
        text=b['text']
        if not any(p.text==text for p in doc.paragraphs):
            # Journal replaces Roman table labels with Arabic numbers.
            for n,roman in reversed(list(enumerate(['I','II','III','IV','V','VI','VII','VIII','IX','X','XI'],1))):
                text=re.sub(r'Table '+roman+r'\b','Table '+str(n),text)
            assert any(p.text==text for p in doc.paragraphs),text
        prefix=' '.join(tokens(text)[:14])
        page=next((i+1 for i,p in enumerate(pages) if prefix in ' '.join(tokens(p))),None)
        paragraphs.append(dict(id=len(paragraphs)+1,section=section,text=text,page=page,tokens=tokens(text)))
    sources=[]
    for m in metadata:
        path=CORPUS/f"source-{m['number']}.txt"
        text=path.read_text(encoding='utf-8',errors='replace'); tok=tokens(text)
        state='searchable retained source text'
        if m['number']==4:state='image-only PDF; no usable extracted text'
        if m['number']==17:state='metadata/abstract only; full-text access failed'
        sources.append(dict(number=m['number'],reference=m['reference'],url=m.get('source_url'),
            resolved_url=m.get('resolved_url'),source_checked_date=m.get('checked_utc'),
            extracted_tokens=len(tok),text_sha256=sha(path),coverage=state,
            included=m['number'] not in [4,17],tokens=tok))
    matches={}
    for width in (6,8):
        index=defaultdict(list)
        for s in sources:
            if not s['included']:continue
            tok=s['tokens']
            for i in range(len(tok)-width+1):index[tuple(tok[i:i+width])].append((s['number'],i))
        found=[]
        for p in paragraphs:
            tok=p['tokens'];covered=set()
            for i in range(len(tok)-width+1):
                for sid,j in index.get(tuple(tok[i:i+width]),[]):
                    st=sources[sid-1]['tokens']
                    # Keep maximal runs, not overlapping sliding-window hits.
                    if i and j and tok[i-1]==st[j-1]:continue
                    k=width
                    while i+k<len(tok) and j+k<len(st) and tok[i+k]==st[j+k]:k+=1
                    key=(sid,i,k)
                    if key in covered:continue
                    covered.add(key)
                    found.append(dict(paragraph=p['id'],section=p['section'],page=p['page'],
                        source=sid,words=k,normalized_match=' '.join(tok[i:i+k]),
                        manuscript_text=p['text']))
        matches[str(width)]=found
    reference_blocks=[b['text'] for b in model if b['kind']=='reference']
    prose='\n'.join(p['text'] for p in paragraphs)
    citations=set(int(n) for group in re.findall(r'\[([0-9, -]+)\]',prose) for n in re.findall(r'\d+',group))
    report=dict(date='2026-10-08',document=str(docpath),docx_sha256=sha(docpath),pdf_sha256=sha(pdfpath),
        pages=len(pdf.pages),scope='Final Discover Computing manuscript only; not a Turnitin or institutional report',
        method='Lowercase NFKC ASCII alphanumeric tokens; URLs and bracketed numeric citations removed; paragraph-bounded exact 6/8-token maximal runs. No stemming, fuzzy search or semantic plagiarism model.',
        screened_paragraphs=len(paragraphs),screened_tokens=sum(len(p['tokens']) for p in paragraphs),
        excluded='Title, author block, keywords, headings, bibliography, table cells, figure captions and figure text. No quotation exclusion within prose.',
        references=len(reference_blocks),cited_reference_numbers=sorted(citations),
        searchable_sources=sum(s['included'] for s in sources),source_tokens=sum(s['extracted_tokens'] for s in sources if s['included']),
        sources=[{k:v for k,v in s.items() if k!='tokens'} for s in sources],matches=matches,
        ai_assistance='Confirmed from this task history: experimental scripting, evidence analysis, manuscript drafting and figure generation code. Percentage not measured; no AI classifier run.',
        limitations=['Only retained cited-source texts screened; proprietary student repositories and paywalled/unseen sources unavailable.',
            'Source texts were acquired 2026-10-02; corpus is a snapshot, not a live crawl.',
            'PDF/OCR artifacts can miss matches. Sources 4 and 17 excluded from exact-text screening.',
            'No cross-source paraphrase detection, idea ownership adjudication, image plagiarism check or external certification.',
            'No similarity percentage or AI-authorship percentage calculated.'])
    (OUT/'screening.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (OUT/'paragraphs.json').write_text(json.dumps([{k:v for k,v in p.items() if k!='tokens'} for p in paragraphs],ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['pages','screened_paragraphs','screened_tokens','searchable_sources','source_tokens','references','matches']},ensure_ascii=True,indent=2))
if __name__=='__main__':main()
