"""Fetch public reference sources and metadata; locally inspect manuscript overlap.

No manuscript is uploaded. This is not an institutional plagiarism certificate.
"""
from __future__ import annotations
import concurrent.futures
import hashlib
import json
import re
from pathlib import Path
from docx import Document
import pymupdf

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"output/docs/seven_revision_audit"
SOURCE=ROOT/"paper/Revon_Research_Paper_IEEE.docx"
URLS={
1:"https://git-scm.com/book/en/v2/Git-Internals-Git-Objects",
2:"https://www.dolthub.com/docs/architecture/storage-engine/block-store/",
3:"https://www.dolthub.com/blog/2024-02-29-storage-engine/",
4:"https://people.eecs.berkeley.edu/~raluca/cs261-f15/readings/merkle.pdf",
5:"https://www.cs.cmu.edu/~sleator/papers/making-data-structures-persistent.pdf",
6:"https://arxiv.org/abs/1407.3561",
7:"https://zenodo.org/records/3935655/files/article.pdf",
8:"https://pdos.csail.mit.edu/papers/lbfs:sosp01/lbfs.pdf",
9:"https://www.cidrdb.org/cidr2015/Papers/CIDR15_Paper18.pdf",
10:"https://www.vldb.org/pvldb/vol9/p624-maddox.pdf",
11:"https://www.vldb.org/pvldb/vol10/p1130-huang.pdf",
12:"https://www.vldb.org/pvldb/vol8/p1346-bhattacherjee.pdf",
13:"https://www.vldb.org/pvldb/vol13/p3411-armbrust.pdf",
14:"https://raw.githubusercontent.com/attic-labs/noms/master/doc/intro.md",
15:"https://www.vldb.org/pvldb/vol11/p1137-wang.pdf",
16:"https://arxiv.org/abs/1103.4282",
17:"https://dspace.mit.edu/bitstreams/0fc426dd-d55a-4e90-98e2-298297d92d99/download",
18:"https://scidb.cs.washington.edu/paper/ICDE13_conf_full_422.pdf",
19:"https://cs.au.dk/~gerth/papers/tcs20.pdf",
20:"https://racelab.cs.ucsb.edu/papers/tpds22-versioned.pdf",
21:"https://www-db.cs.tum.edu/~schuele/data/tardis.pdf",
22:"https://www.pure.ed.ac.uk/ws/files/16509989/Why_and_Where_A_Characterization_of_Data_Provenance.pdf",
23:"https://www.cidrdb.org/cidr2013/Papers/CIDR13_Paper111.pdf",
24:"https://arxiv.org/abs/2006.03018",
25:"https://arxiv.org/abs/2301.13095",
26:"https://arxiv.org/abs/2402.11741",
27:"https://www.nature.com/articles/s41597-024-03153-y",
28:"https://arxiv.org/abs/2405.17701",
29:"https://arxiv.org/abs/2609.02106",
30:"https://arxiv.org/abs/2512.09762",
31:"https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page"}

def fetch(index,text):
    import requests
    from bs4 import BeautifulSoup
    record=dict(number=index,reference=text,source_url=URLS[index],checked_utc="2026-10-02")
    try:
        response=requests.get(URLS[index],timeout=35,headers={"User-Agent":"Revon-reference-verification/1.0"})
        record.update(http_status=response.status_code,resolved_url=response.url)
        if response.status_code==200:
            pdf=response.content.startswith(b"%PDF")
            file=OUT/f"source-{index}.{'pdf' if pdf else 'html'}"
            file.write_bytes(response.content)
            record["source_sha256"]=hashlib.sha256(response.content).hexdigest()
            if pdf:
                doc=pymupdf.open(stream=response.content,filetype="pdf")
                source_text="\n".join(page.get_text() for page in doc)
            else:
                soup=BeautifulSoup(response.content,"html.parser")
                source_text=soup.get_text(" ",strip=True)
                record["page_title"]=soup.title.get_text() if soup.title else ""
                tags=soup.select('meta[name^="citation_"]')
                record["citation_metadata"]={tag.get("name"):tag.get("content") for tag in tags if tag.get("name")!="citation_author"}
                record["citation_authors"]=[tag.get("content") for tag in tags if tag.get("name")=="citation_author"]
            (OUT/f"source-{index}.txt").write_text(source_text,encoding="utf-8")
            record["extracted_words"]=len(source_text.split())
            if "/abs/" in URLS[index]:
                pdf_response=requests.get(URLS[index].replace("/abs/","/pdf/"),timeout=35)
                if pdf_response.status_code==200 and pdf_response.content.startswith(b"%PDF"):
                    (OUT/f"source-{index}.pdf").write_bytes(pdf_response.content)
                    doc=pymupdf.open(stream=pdf_response.content,filetype="pdf")
                    (OUT/f"source-{index}.txt").write_text("\n".join(p.get_text() for p in doc),encoding="utf-8")
                    record["full_text_pdf_sha256"]=hashlib.sha256(pdf_response.content).hexdigest()
        match=re.search(r"doi:\s*([^\s]+)",text)
        if match:
            doi=match.group(1).rstrip(".")
            metadata=requests.get("https://api.crossref.org/works/"+doi,timeout=35)
            record["doi"]=doi; record["crossref_status"]=metadata.status_code
            if metadata.status_code==200:
                m=metadata.json()["message"]
                record["crossref"]={k:m.get(k) for k in ("title","author","published","container-title","volume","issue","page","article-number","URL")}
    except Exception as exc:
        record["error"]=f"{type(exc).__name__}: {exc}"
    return record

def tokens(text): return re.findall(r"[a-z0-9]+",text.lower())

def overlap(document, output_name="overlap_triage.json"):
    sources={int(p.stem.split('-')[1]):tokens(p.read_text(encoding="utf-8")) for p in OUT.glob("source-*.txt") if int(p.stem.split('-')[1]) not in (4,17)}
    # Twelve-word exact sequences are a triage signal, not a plagiarism verdict.
    windows={i:{tuple(t[j:j+12]) for j in range(max(0,len(t)-11))} for i,t in sources.items()}
    findings=[]
    paragraphs=[]
    for number,p in enumerate(document.paragraphs):
        if p.text.strip().lower()=="references": break
        if p.text.strip(): paragraphs.append((number,p.text))
    for number,text in paragraphs:
        words=tokens(text)
        hits=[]
        for i,source in windows.items():
            matches={tuple(words[j:j+12]) for j in range(max(0,len(words)-11))} & source
            if matches: hits.append(dict(reference=i,exact_12_word_matches=len(matches),examples=[" ".join(m) for m in sorted(matches)[:2]]))
        if hits: findings.append(dict(paragraph=number,excerpt=text[:180],matches=hits))
    report=dict(method="local exact 12-word sequence triage against downloaded public reference sources, excluding bibliography; manual source and paraphrase review required",
                source_count=len(sources),paragraphs_checked=len(paragraphs),findings=findings,
                automatic_screening_exclusions={4:"scanned paper reviewed visually; no extracted text",17:"MIT abstract and metadata reviewed; full text unavailable"},
                limitations="Not an iThenticate/Turnitin search. Does not cover inaccessible sources, unpublished works, proprietary corpora, semantic plagiarism, or certify authorship.")
    (OUT/output_name).write_text(json.dumps(report,indent=2),encoding="utf-8")
    return report

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    document=Document(SOURCE)
    refs={int(m.group(1)):p.text for p in document.paragraphs if (m:=re.match(r"^\[(\d+)\]",p.text))}
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
        records=list(pool.map(lambda pair:fetch(*pair),sorted(refs.items())))
    (OUT/"reference_verification.json").write_text(json.dumps(records,indent=2),encoding="utf-8")
    print(json.dumps([dict(number=r["number"],http=r.get("http_status"),title=r.get("page_title",r.get("crossref",{}).get("title")),crossref=r.get("crossref_status"),error=r.get("error")) for r in records],indent=2))
    print(json.dumps(overlap(document),indent=2))

if __name__=="__main__":main()
