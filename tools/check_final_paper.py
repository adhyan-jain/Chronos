"""Read-only CI checks of current/archived manuscripts and recorded evidence.

These are integrity/content checks, not timing reproduction or visual review.
No ignored local output, private cloud archive, Word export or benchmark needed.
"""
from __future__ import annotations
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import statistics
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TITLE = "Adaptive Merkle Trie Versioning for Efficient Structured Data Management"
BUNDLES = ("linux-validation-20261008", "linux-final-supplement-20261008")
ROMANS = dict(zip(("XI", "X", "IX", "VIII", "VII", "VI", "V", "IV", "III", "II", "I"), range(11, 0, -1)))


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def normalize(text):
    text = re.sub(r"^(Fig\. \d+)\. ", r"\1 ", text)
    for roman, number in ROMANS.items():
        text = re.sub(r"\b(?:TABLE|Table) " + roman + r"\b", "Table " + str(number), text)
    return " ".join(text.split()).rstrip(".")


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def check_hashes(directory):
    hashes = json.loads((directory / "SHA256SUMS.json").read_text(encoding="utf-8"))
    require(bool(hashes), f"empty checksum manifest: {directory}")
    for name, digest in hashes.items():
        path = directory / name
        require(path.is_file() and sha(path.read_bytes()) == digest,
                f"collection checksum mismatch: {directory.name}/{name}")
    return len(hashes)


def check_package(path, root):
    with zipfile.ZipFile(path) as package:
        hashes = json.loads(package.read("PACKAGE_SHA256SUMS.json"))
        require(bool(hashes), "empty supplementary package checksum manifest")
        for name, digest in hashes.items():
            require(sha(package.read(name)) == digest, f"package checksum mismatch: {name}")
            if any(name.startswith("evidence/" + bundle + "/") for bundle in BUNDLES):
                require((root / name).read_bytes() == package.read(name),
                        f"packaged Linux evidence differs from repository: {name}")
        model = json.loads(package.read("support/manuscript_model.json"))
        assets = {n: package.read(f"figures/Fig{n}.png") for n in range(1, 8)}
    for kind, count in (("table", 11), ("figure", 7)):
        require([b["number"] for b in model if b["kind"] == kind] == list(range(1, count + 1)),
                f"model {kind} numbering/count changed")
    return model, assets, len(hashes)


def check_doc(path, model, assets, discover):
    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    require(normalize(paragraphs[0]) == TITLE, f"wrong manuscript title: {path.name}")
    tables = [b for b in model if b["kind"] == "table"]
    require(len(document.tables) == len(tables), f"table count changed: {path.name}")
    require(len(document.inline_shapes) == 7, f"figure count changed: {path.name}")
    for table, expected in zip(document.tables, tables):
        require([[c.text for c in row.cells] for row in table.rows] == [expected["headers"]] + expected["rows"],
                f"table {expected['number']} differs from packaged model: {path.name}")
    expected = [normalize(b["text"]) for b in model if b["kind"] in ("body", "reference")]
    observed = [normalize(p) for p in paragraphs]
    offset = 0
    for text in expected:
        require(text in observed[offset:], f"missing or reordered manuscript content: {text[:100]}")
        offset = observed.index(text, offset) + 1
    for kind in ("author", "abstract"):
        for block in (b for b in model if b["kind"] == kind):
            require(any(normalize(block["text"]) in p for p in observed), f"missing {kind}: {path.name}")
    abstract = next(p for p in paragraphs if p.startswith("Abstract:"))
    require(len(abstract.split()) - 1 < 250, f"abstract too long: {path.name}")
    numbers = [int(m[1]) for p in paragraphs if (m := re.match(r"^\[(\d+)\]", p))]
    require(numbers == list(range(1, 32)), f"bibliography numbering changed: {path.name}")
    body = "\n".join(p for p in paragraphs if not re.match(r"^\[\d+\]", p))
    body += "\n" + "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    citations = set()
    for group in re.findall(r"\[([\d, -]+)\]", body):
        for item in group.split(","):
            bounds = [int(n) for n in item.strip().split("-")]
            citations.update(range(bounds[0], bounds[-1] + 1))
    require(citations == set(numbers), f"uncited or undefined reference: {path.name}")
    captions = [normalize(p) for p in paragraphs if re.match(r"^Fig\. \d+", p)]
    for block in (b for b in model if b["kind"] == "figure"):
        n = block["number"]
        require(normalize(f"Fig. {n} {block['caption']}") in captions,
                f"figure {n} caption differs from model: {path.name}")
        blip = document.inline_shapes[n - 1]._inline.xpath(".//a:blip")[0]
        require(document.part.related_parts[blip.get(qn("r:embed"))].blob == assets[n],
                f"figure {n} image differs from package: {path.name}")
    if discover:
        sizes = [float(e.get(qn("w:val"))) / 2 for e in document._element.xpath(".//w:sz")]
        require(sizes and min(sizes) >= 12, "Discover Computing text smaller than 12 pt")
        require(all(s._sectPr.find(qn("w:cols")) is None or
                    s._sectPr.find(qn("w:cols")).get(qn("w:num"), "1") == "1"
                    for s in document.sections), "Discover Computing must be single column")
    else:
        require(any(s._sectPr.find(qn("w:cols")) is not None and
                    s._sectPr.find(qn("w:cols")).get(qn("w:num")) == "2"
                    for s in document.sections), "IEEE two-column section missing")
    return dict(tables=len(tables), figures=7, references=len(numbers), abstract_words=len(abstract.split()) - 1)


def check_pdf(path):
    import pymupdf
    with pymupdf.open(path) as pdf:
        require(len(pdf) > 0, f"empty PDF: {path.name}")
        captions, text = [], ""
        for index, page in enumerate(pdf, 1):
            page_text = page.get_text()
            text += page_text
            require(len(page_text.split()) >= 8, f"near-empty PDF page {index}: {path.name}")
            blocks = page.get_text("dict")["blocks"]
            images = [b for b in blocks if b["type"] == 1]
            for block in blocks:
                require(page.rect.contains(pymupdf.Rect(block["bbox"])),
                        f"PDF content outside page {index}: {path.name}")
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    match = re.match(r"^Fig\.\s*(\d+)", "".join(s["text"] for s in line["spans"]).strip())
                    if match:
                        captions.append(int(match[1]))
                        require(bool(images), f"figure caption without image on page {index}: {path.name}")
            require("\ufffd" not in page_text, f"replacement glyph on page {index}: {path.name}")
        require(sorted(captions) == list(range(1, 8)), f"PDF figure captions missing: {path.name}")
        for marker in ("516", "72", "32,768", "16,384", "[31]", "poornima.n@vit.ac.in"):
            require(marker in text, f"PDF is missing {marker}: {path.name}")
        return dict(pages=len(pdf), figure_captions=len(captions))


def check_evidence(root):
    sys.path.insert(0, str(ROOT))
    from experiments.evidence_audit import _check_summary, METRICS, EXTRA_TELEMETRY_METRICS
    from experiments.paired_uncertainty import analyze
    from tempfile import TemporaryDirectory
    report = {}
    for name, counts in zip(BUNDLES, ((516, 399, 114, 3), (27, 21, 6, 0))):
        directory = root / "evidence" / name
        hashes = check_hashes(directory)
        rows = read_csv(directory / "raw_results.csv")
        protocol = json.loads((directory / "protocol.json").read_text(encoding="utf-8"))
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        require(sha((directory / "protocol.json").read_bytes()) == manifest["protocol_sha256"], f"frozen protocol changed: {name}")
        require((len(rows), *(sum(r["trial_kind"] == kind for r in rows)
                             for kind in ("measured", "warmup", "diagnostic"))) == counts, f"trial counts changed: {name}")
        require(len(rows) == len(protocol["execution_plan"]), f"incomplete execution plan: {name}")
        for relative, digest in manifest["source_hashes"].items():
            require(sha((directory / "source" / relative).read_bytes()) == digest, f"executed source changed: {name}/{relative}")
        for index, (row, planned) in enumerate(zip(rows, protocol["execution_plan"])):
            require(row["status"] == "ok" and row["correctness"] == "True", f"failed execution {index}: {name}")
            for field in ("phase", "scenario", "model", "trial_kind", "trial", "execution_order"):
                require(row[field] == str(planned[field]), f"plan mismatch {field} at {index}: {name}")
            require(int(row["hybrid_threshold"]) == planned["threshold"], f"threshold changed at {index}: {name}")
            proof = json.loads((directory / "verification" / f"{index:05d}.json").read_text(encoding="utf-8"))
            require(proof["diff_correct"] and proof["all_history_verified"] and
                    proof["diff_output_sha256"] == proof["diff_oracle_sha256"] and
                    proof["workload_sha256"] == row["workload_sha256"] and
                    len(proof["historical_states"]) == int(row["commits"]) + 1 and
                    all(h["observed_sha256"] == h["oracle_sha256"] for h in proof["historical_states"]),
                    f"semantic/history proof mismatch at {index}: {name}")
        issues = []
        _check_summary(rows, read_csv(directory / "summary.csv"), issues, METRICS + EXTRA_TELEMETRY_METRICS)
        require(not issues, f"raw/summary disagreement: {name}: {issues}")
        if name == BUNDLES[0]:
            with TemporaryDirectory() as temp:
                calculated = analyze(directory / "raw_results.csv", Path(temp) / "paired.csv")
                require(calculated == read_csv(directory / "paired_uncertainty.csv"), "paired interval recomputation differs")
            adverse = [r for r in rows if r["scenario"] == "linux-threshold-h64-c512" and r["trial_kind"] == "measured"]
            for trial in range(1, 8):
                pair = {r["model"]: r for r in adverse if int(r["trial"]) == trial}
                require(pair["Revon-H"]["strategy_selected"] == "merkle" and
                        float(pair["Revon-log calibration"]["diff_ms"]) < float(pair["Revon-H"]["diff_ms"]),
                        "adverse frozen-threshold finding changed")
        else:
            require(all(r["process_tree_read_bytes"] != "" and float(r["process_tree_read_bytes"]) == 0 for r in rows),
                    "observed zero I/O was lost")
            geometry = read_csv(directory / "geometry/raw_results.csv")
            require(len(geometry) == 45 and sum(r["trial_kind"] == "measured" for r in geometry) == 35,
                    "geometry trial counts changed")
            require(all(r["status"] == "ok" and r["correctness"] == "True" for r in geometry), "geometry correctness failed")
            for row in read_csv(directory / "geometry/summary.csv"):
                measured = [r for r in geometry if r["trial_kind"] == "measured" and r["config"] == row["config"]]
                require(len(measured) == 7, "geometry repetitions changed")
                for metric in (k[:-7] for k in row if k.endswith("_median")):
                    if row[metric + "_median"]:
                        require(math.isclose(statistics.median(float(r[metric]) for r in measured),
                                             float(row[metric + "_median"]), rel_tol=1e-10, abs_tol=1e-10),
                                f"geometry summary changed: {metric}")
        report[name] = dict(collection_hashes=hashes, executions=len(rows))
    return report


def main(root=ROOT):
    model, assets, package_hashes = check_package(root / "paper/Revon_Linux_Reproducibility_Supplement.zip", root)
    manuscripts = {}
    for form in ("IEEE", "Discover_Computing"):
        path = root / f"paper/previous versions/Revon_Research_Paper_{form}.docx"
        manuscripts[f"archived_{form}"] = check_doc(path, model, assets, form == "Discover_Computing")
        manuscripts[f"archived_{form}"].update(check_pdf(path.with_suffix(".pdf")))
    sys.path.insert(0, str(ROOT))
    from tools.check_springer_paper import check_springer
    manuscripts["current_Springer"] = check_springer(root / "paper/Revon Research Paper Springer.docx", model)
    report = dict(passed=True, manuscripts=manuscripts, package_hashes=package_hashes, evidence=check_evidence(root),
                  scope="read-only content, provenance, semantic proofs and summary integrity; not timing reproduction or visual certification")
    print(json.dumps(report, indent=2))
    return report


if __name__ == "__main__":
    main()
