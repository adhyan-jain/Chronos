"""Create a reusable proceedings template and a content-preserving Revon copy.

Uses the two supplied local Word references. This is a conference variant,
not the Discover Computing journal submission template.
"""
from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'output/templates/springer_template_.doc_2.docx'
GUIDE = ROOT / 'output/templates/How to Write a Springer Conference Paper.docx'
SOURCE = ROOT / 'paper/Revon_Research_Paper_Discover_Computing.docx'
TEMPLATE = ROOT / 'output/templates/Springer_Conference_Template.docx'
PAPER = ROOT / 'output/docs/Revon_Research_Paper_Springer_Conference.docx'
QA = ROOT / '.codex_tmp/springer-template-20261007'
PRESERVE = ['word/theme/theme1.xml', 'word/fontTable.xml',
            'word/numbering.xml', 'word/webSettings.xml']
CITE = re.compile(r'\[(\d+(?:\s*,\s*\d+)*)\]')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def text(p):
    return ''.join(n.text or '' for n in p.iter(qn('w:t')))


def replace_text(p, start, end, replacement):
    """Replace a text span while retaining surrounding run formatting/links."""
    nodes = list(p.iter(qn('w:t')))
    pos = 0
    inserted = False
    for n in nodes:
        value = n.text or ''
        stop = pos + len(value)
        if stop > start and pos < end:
            a, b = max(0, start-pos), min(len(value), end-pos)
            n.text = value[:a] + (replacement if not inserted else '') + value[b:]
            n.set(qn('xml:space'), 'preserve')
            inserted = True
        pos = stop


def set_font(run, size=10, bold=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if key.endswith('Theme'):
            del rf.attrib[key]
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:'+key), 'Times New Roman')


def springer_reference(value):
    """Change bibliography punctuation and author order, preserving metadata."""
    m=re.fullmatch(r'\[(\d+)\] (.+?), [“"](.+?)[”"](.*)',value)
    assert m,value
    number,authors,title,tail=m.groups()
    def author(a):
        a=a.strip()
        etal=a.endswith(' et al.')
        if etal:a=a[:-7]
        parts=a.split()
        initials=[]
        while parts and re.fullmatch(r'[A-Z](?:\.-?[A-Z])*\.',parts[0]):initials.append(parts.pop(0))
        result=(' '.join(parts)+', '+''.join(initials)) if initials and parts else a
        return result+(' et al.' if etal else '')
    names=([author(a) for a in re.split(r',\s*(?:and\s+)?|\s+and\s+',authors)]
           if re.match(r'[A-Z]\.',authors) else [authors])
    title=title.rstrip(',.')
    tail=tail.lstrip(', .')
    if tail.startswith('in '):tail='In: '+tail[3:]
    tail=re.sub(r'vol\. (\d+), no\. (\d+), pp\. ([\d-]+), (\d{4})([,.])',r'\1(\2), \3 (\4).',tail)
    tail=re.sub(r'vol\. (\d+), pp\. ([\d-]+), (\d{4})([,.])',r'\1, \2 (\3).',tail)
    # Where the year is not already enclosed or part of an access/preprint date,
    # keep it at its original metadata position and use Springer parentheses.
    urls=[(m.start(),m.end()) for m in re.finditer(r'https?://\S+',tail)]
    def year(m):
        protected=any(a<=m.start()<b for a,b in urls)
        dated=any(x in tail[:m.start()] for x in ('accessed','arXiv:','Blog,'))
        return m[1] if protected or dated else '('+m[1]+')'
    tail=re.sub(r'(?<![\d(/:])\b((?:19|20)\d{2})\b(?![\d/])',year,tail)
    tail=tail.replace(').,',').').replace(')..',').')
    return f'[{number}] '+', '.join(names)+': '+title+'. '+tail


def style(doc, name, size, bold=False, before=0, after=0, center=False):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'Times New Roman'
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = RGBColor(0,0,0)
    f = s.paragraph_format
    f.left_indent = f.right_indent = f.first_line_indent = Cm(0)
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing = 1
    f.keep_with_next = name.startswith('Heading')
    f.page_break_before = False
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if name.startswith('Heading'):
        f.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return s


def base():
    d = Document(REFERENCE)
    body = d._element.body
    sec = deepcopy(body.sectPr)
    for child in list(body):
        body.remove(child)
    body.append(sec)
    for name in ['w:headerReference', 'w:footerReference', 'w:pgNumType', 'w:cols']:
        for n in list(sec.findall(qn(name))):
            sec.remove(n)
    for rid, rel in list(d.part.rels.items()):
        if rel.reltype in {RT.HEADER, RT.FOOTER, RT.IMAGE, RT.HYPERLINK, RT.FOOTNOTES, RT.ENDNOTES}:
            d.part.drop_rel(rid)
    for rid, rel in list(d.part.package.rels.items()):
        if rel.reltype in {RT.THUMBNAIL, RT.CUSTOM_PROPERTIES}:
            d.part.package.rels.pop(rid)
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(4.4)
    s.top_margin, s.bottom_margin = Cm(4.6), Cm(5.8)
    s.header_distance = s.footer_distance = Cm(1)
    style(d, 'Normal', 10)
    style(d, 'Title', 14, True, after=14, center=True)
    style(d, 'Heading 1', 12, True, before=12, after=6)
    style(d, 'Heading 2', 10, True, before=10, after=4)
    style(d, 'Caption', 9, before=5, after=5)
    style(d, 'Springer Authors', 12, after=8, center=True)
    style(d, 'Springer Affiliation', 10, after=2, center=True)
    style(d, 'Springer Reference', 9)
    rf = d.styles['Springer Reference'].paragraph_format
    rf.left_indent, rf.first_line_indent = Cm(.6), Cm(-.6)
    rf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.core_properties.title = ''
    d.core_properties.author = ''
    d.core_properties.subject = 'Springer conference proceedings format from supplied references'
    # The converted sample points to note separators in removed parts. Word
    # rejects that dangling settings reference even when LibreOffice accepts it.
    for tag in ['w:trackRevisions', 'w:documentProtection', 'w:footnotePr',
                'w:endnotePr', 'w:evenAndOddHeaders', 'w:savePreviewPicture']:
        for n in list(d.settings.element.findall(qn(tag))):
            d.settings.element.remove(n)
    return d


def para_format(p, role='body', first=False):
    f = p.paragraph_format
    f.left_indent = f.right_indent = f.first_line_indent = Cm(0)
    f.space_before = f.space_after = Pt(0)
    f.line_spacing = 1
    f.keep_with_next = False
    f.keep_together = False
    f.page_break_before = False
    f.widow_control = True
    f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for br in p._p.findall('.//' + qn('w:br')):
        if br.get(qn('w:type')) == 'page':
            br.getparent().remove(br)
    for n in p._p.findall('.//' + qn('w:lastRenderedPageBreak')):
        n.getparent().remove(n)
    size = 10
    if role == 'title':
        p.style = 'Title'; size = 14
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f.space_after = Pt(14); f.keep_with_next = True
    elif role == 'authors':
        p.style = 'Springer Authors'; size = 12
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f.space_after = Pt(8); f.keep_with_next = True
    elif role == 'affiliation':
        p.style = 'Springer Affiliation'
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f.space_after = Pt(2); f.keep_with_next = True
    elif role in ('h1', 'h2'):
        p.style = 'Heading 1' if role == 'h1' else 'Heading 2'
        size = 12 if role == 'h1' else 10
        f.alignment = WD_ALIGN_PARAGRAPH.LEFT
        f.space_before, f.space_after = Pt(12 if role=='h1' else 10), Pt(6 if role=='h1' else 4)
        f.keep_with_next = f.keep_together = True
    elif role == 'caption':
        p.style = 'Caption'; size = 9
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(p.text) < 95 else WD_ALIGN_PARAGRAPH.JUSTIFY
        f.space_before = f.space_after = Pt(5)
        f.keep_together = True
        f.keep_with_next = p.text.startswith('Table')
    elif role == 'reference':
        p.style = 'Springer Reference'; size = 9
        f.left_indent, f.first_line_indent = Cm(.6), Cm(-.6)
        f.alignment = WD_ALIGN_PARAGRAPH.LEFT
        f.space_after = Pt(0)
        f.keep_together = True
    elif role == 'figure':
        p.style = 'Normal'
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f.space_before = Pt(7)
        f.keep_with_next = f.keep_together = True
    else:
        p.style = 'Normal'
        f.first_line_indent = Cm(0 if first else .4)
    for r in p.runs:
        set_font(r, size, True if role in ('title','h1','h2') else None)


def tables(doc):
    for ti, t in enumerate(doc.tables,1):
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = [c.width or Cm(1) for c in t.rows[0].cells]
        new = [int(Cm(12.2)*w/sum(widths)) for w in widths]
        for col,w in zip(t.columns,new): col.width = w
        pr = t._tbl.tblPr
        for tag in ['w:tblBorders','w:tblCellMar','w:tblInd']:
            for n in list(pr.findall(qn(tag))): pr.remove(n)
        border=OxmlElement('w:tblBorders')
        for edge in ['top','bottom','left','right','insideH','insideV']:
            e=OxmlElement('w:'+edge)
            e.set(qn('w:val'),'single' if edge in ('top','bottom') else 'nil')
            e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');border.append(e)
        pr.append(border)
        mar=OxmlElement('w:tblCellMar')
        for edge in ['top','bottom','left','right']:
            n=OxmlElement('w:'+edge);n.set(qn('w:w'),'45' if edge in ('left','right') else '40');n.set(qn('w:type'),'dxa');mar.append(n)
        pr.append(mar)
        compact = ti not in (1,2,4,10,15)
        for ri,row in enumerate(t.rows):
            rp=row._tr.get_or_add_trPr()
            for n in list(rp.findall(qn('w:trHeight'))): rp.remove(n)
            if rp.find(qn('w:cantSplit')) is None: rp.append(OxmlElement('w:cantSplit'))
            if ri==0 and rp.find(qn('w:tblHeader')) is None: rp.append(OxmlElement('w:tblHeader'))
            for ci,c in enumerate(row.cells):
                c.width=new[ci]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=c._tc.get_or_add_tcPr()
                for n in list(cp.findall(qn('w:shd')))+list(cp.findall(qn('w:tcBorders'))):cp.remove(n)
                if ri==0:
                    cb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');cb.append(b);cp.append(cb)
                for p in c.paragraphs:
                    para_format(p,first=True)
                    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or ti in (1,2,4) else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_with_next=compact and ri<len(t.rows)-1
                    if ri==0:p.paragraph_format.keep_with_next=True
                    for r in p.runs:set_font(r,9,ri==0)


def save(doc,path):
    path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(path)
    # Retain opaque reference parts byte-for-byte, rather than reserialize.
    with ZipFile(REFERENCE) as z: originals={n:z.read(n) for n in PRESERVE}
    with ZipFile(path) as z: contents={n:z.read(n) for n in z.namelist()}
    contents.update(originals)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for n,value in contents.items():z.writestr(n,value)


def paper():
    src=Document(SOURCE); d=base()
    roles=[]
    for i,p in enumerate(src.paragraphs):
        roles.append('title' if i==0 else 'authors' if i==1 else 'affiliation' if i<6 else 'h1' if p.style.name=='Heading 1' else 'h2' if p.style.name=='Heading 2' else 'caption' if p.style.name=='Caption' else 'figure' if p._p.find('.//'+qn('w:drawing')) is not None else 'reference' if i>207 else 'body')
    relmap={}
    for rid,rel in src.part.rels.items():
        if rel.reltype in (RT.IMAGE,RT.HYPERLINK):
            relmap[rid]=d.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,is_external=rel.is_external)
    for child in src._element.body:
        if child.tag==qn('w:sectPr'):continue
        c=deepcopy(child)
        for n in c.iter():
            for attr in [qn('r:embed'),qn('r:link'),qn('r:id')]:
                if n.get(attr) in relmap:n.set(attr,relmap[n.get(attr)])
        d._element.body.insert(-1,c)
    # Keep facts; move affiliation ahead of email lines to match references.
    pp=d.paragraphs
    pp[1]._p.addnext(pp[5]._p);pp[1]._p.addnext(pp[4]._p)
    rolemap={p._p:role for p,role in zip(pp,roles)}
    # Split the inline abstract label into the guide's heading and text.
    abstract=pp[6]
    assert abstract.text.startswith('Abstract: ')
    replace_text(abstract._p,0,len('Abstract: '),'')
    ah=abstract.insert_paragraph_before('Abstract','Heading 1');rolemap[ah._p]='h1'
    seen=[]
    for child in d._element.body:
        if child.tag==qn('w:p') and rolemap.get(child)=='reference':continue
        for p in ([child] if child.tag==qn('w:p') else child.iter(qn('w:p'))):
            for m in CITE.finditer(text(p)):
                ids=[int(x) for x in m[1].split(',')]
                if all(1<=x<=31 for x in ids):
                    for x in ids:
                        if x not in seen:seen.append(x)
    assert set(seen)==set(range(1,32)),seen
    mapping={old:new for new,old in enumerate(seen,1)}
    for p in d._element.body.iter(qn('w:p')):
        if rolemap.get(p)=='reference':continue
        for m in reversed(list(CITE.finditer(text(p)))):
            ids=[int(x) for x in m[1].split(',')]
            if all(x in mapping for x in ids):replace_text(p,m.start(),m.end(),'['+', '.join(str(mapping[x]) for x in ids)+']')
    refs={}
    for p in pp[208:]:
        m=re.match(r'\[(\d+)\]',p.text);old=int(m[1]);refs[old]=p._p
        replace_text(p._p,0,len(m[0]),f'[{mapping[old]}]')
        formatted=springer_reference(p.text)
        replace_text(p._p,0,len(p.text),formatted)
    for old in seen:d._element.body.insert(-1,refs[old])
    for n in d._element.body.iter():
        for attr in [qn('w:anchor'),qn('w:name')]:
            v=n.get(attr)
            if v and re.fullmatch(r'ref_\d+',v):n.set(attr,'ref_'+str(mapping[int(v[4:])]))
    sec=0;sub=0;prev='h1'
    for p in d.paragraphs:
        role=rolemap.get(p._p,'body')
        if role=='h1':
            m=re.match(r'\d+\s+',p.text)
            if m:sec=int(p.text.split()[0]);sub=0
            elif p.text.startswith('Appendix'):sec=p.text[9];sub=0
        elif role=='h2':
            sub+=1
            p.runs[0].text=f'{sec}.{sub} '+p.runs[0].text
        para_format(p,role,first=prev in ('h1','h2') or p==abstract)
        if p==abstract:p.paragraph_format.space_after=Pt(7)
        if p.text.startswith('Keywords:'):
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(8)
            value=p.text.replace(', ', ' · ')
            p.clear();r=p.add_run('Keywords: ');set_font(r,10);r.italic=True
            set_font(p.add_run(value[len('Keywords: '):]),10)
        for label,num in [('A','1'),('B','2')]:
            for m in reversed(list(re.finditer(r'Appendix '+label+r'\b',p.text))):replace_text(p._p,m.start(),m.end(),'Appendix '+num)
        if p.text.startswith('B.1 '):replace_text(p._p,0,4,'2.1 ')
        if p._p.getprevious() is not None and p._p.getprevious().tag==qn('w:tbl'):
            p.paragraph_format.space_before=Pt(5)
        prev=role
    # Unnumbered run-in acknowledgment rather than another major section.
    ack=next(p for p in d.paragraphs if p.text=='Acknowledgment')
    following=ack._p.getnext()
    p=Paragraph(following,d._body)
    lead=OxmlElement('w:r');pr=OxmlElement('w:rPr');pr.append(OxmlElement('w:b'));lead.append(pr);tx=OxmlElement('w:t');tx.text='Acknowledgment. ';tx.set(qn('xml:space'),'preserve');lead.append(tx);following.insert(1 if following.find(qn('w:pPr')) is not None else 0,lead)
    ack._p.getparent().remove(ack._p)
    tables(d)
    for i,s in enumerate(d.inline_shapes,1):
        width=Cm(8.6 if i==2 else 12.2)
        ratio=width/s.width;s.width=width;s.height=int(s.height*ratio)
        for n in s._inline.iter(qn('a:ext')):
            n.set('cx',str(s.width));n.set('cy',str(s.height))
    d.core_properties.title=d.paragraphs[0].text
    d.core_properties.author=src.paragraphs[1].text
    save(d,PAPER)
    return mapping


def template():
    d=base()
    items=[('title','[Paper title]'),('authors','[First author], [Second author], [Third author]'),('affiliation','[Department, Institution, City, Country]'),('affiliation','[Author email addresses]'),('affiliation','[Corresponding author and email]'),('h1','Abstract'),('body','[Replace with a 150–250 word abstract describing the problem, method, principal results and limitations. Remove all bracketed instructions before submission.]'),('body','Keywords: [Term 1] · [Term 2] · [Term 3] · [Term 4]'),('h1','1 Introduction'),('body','[Explain the problem, motivation, scope and contributions. Cite genuine sources using numbered brackets, in order of first appearance.]'),('h1','2 Related Work'),('body','[Compare the most relevant published work and identify the specific gap addressed.]'),('h1','3 Method'),('h2','3.1 Design and assumptions'),('body','[Describe the method and assumptions precisely enough to reproduce the work.]'),('h2','3.2 Equations'),('body','[Replace the editable example equation with the equations required by your method.]')]
    for role,t in items:para_format(d.add_paragraph(t),role,True)
    p=d.add_paragraph();para_format(p,first=True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    m=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='y = a x + b';mr.append(mt);m.append(mr);p._p.append(m)
    for role,t in [('h1','4 Experimental Evaluation'),('body','[Describe datasets, versions, hardware, software, correctness checks, repetitions and uncertainty. Report actual measurements only.]'),('caption','Table 1 [Concise table caption]')]:para_format(d.add_paragraph(t),role,True)
    t=d.add_table(rows=3,cols=3)
    for row,vals in zip(t.rows,[['Metric','Method A','Method B'],['[Metric 1]','[Value]','[Value]'],['[Metric 2]','[Value]','[Value]']]):
        for c,v in zip(row.cells,vals):c.text=v
    tables(d)
    for role,t in [('body','[Insert a legible high-resolution or vector figure here, within the text width. Keep its paragraph with the following caption.]'),('caption','Fig. 1 [Concise figure caption]'),('h1','5 Results and Discussion'),('body','[Interpret results, compare methods and report threats to validity without overstating generality.]'),('h1','6 Conclusion'),('body','[Summarize the supported findings, limits and specific next steps.]'),('h1','Appendix'),('body','[Optional reproducibility details. Remove this appendix if unused. Number multiple appendices 1, 2 and so on.]'),('body','Acknowledgments. [Insert only applicable acknowledgments and required declarations.]'),('h1','References'),('reference','[1] [Authors]: [Actual title]. [Publication, volume, pages] ([Year]). [Verified DOI or URL].')]:
        p=d.add_paragraph(t);para_format(p,role,True)
        if t.startswith('[Insert a legible'):p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
        if t.startswith('Acknowledgments.'):p.runs[0].bold=True
    d.core_properties.title='Springer Conference Template'
    save(d,TEMPLATE)


def verify(mapping,source_hashes):
    src=Document(SOURCE);out=Document(PAPER)
    assert len(out.tables)==len(src.tables)==16
    assert len(out.inline_shapes)==len(src.inline_shapes)==11
    inverse={v:k for k,v in mapping.items()}
    def decite(t):
        def f(m):
            ids=[int(x) for x in m[1].split(',')]
            return '['+', '.join(str(inverse[x]) for x in ids)+']' if all(x in inverse for x in ids) else m[0]
        return CITE.sub(f,t)
    def cite_spacing(t):
        return CITE.sub(lambda m:'['+', '.join(str(int(x)) for x in m[1].split(','))+']',t)
    for a,b in zip(src.tables,out.tables):
        assert [[cite_spacing(c.text) for c in r.cells] for r in a.rows]==[[cite_spacing(decite(c.text)) for c in r.cells] for r in b.rows]
    def imgs(doc):return [hashlib.sha256(doc.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest() for s in doc.inline_shapes]
    assert imgs(src)==imgs(out)
    expected=[p.text for p in src.paragraphs if p.text]
    actual=[decite(p.text) for p in out.paragraphs if p.text and p.style.name!='Springer Reference']
    expected=expected[:-31]
    def norm(t):
        t=re.sub(r'^(?:\d+|[A-Z])\.\d+\s+','',t)
        if t=='Abstract':return ''
        if t.startswith('Abstract: '):t=t[len('Abstract: '):]
        if t=='Acknowledgment':return ''
        t=re.sub(r'^Acknowledgment\. ','',t)
        t=t.replace('Appendix 1','Appendix A').replace('Appendix 2','Appendix B')
        if t.startswith('Keywords:'):t=t.replace(' · ', ', ')
        return cite_spacing(t)
    assert sorted(filter(None,map(norm,expected)))==sorted(filter(None,map(norm,actual))), 'Body text preservation failed'
    entries=[p.text for p in out.paragraphs if p.style.name=='Springer Reference']
    assert [int(re.match(r'\[(\d+)\]',p)[1]) for p in entries]==list(range(1,32))
    for p in src.paragraphs[208:]:
        old=int(re.match(r'\[(\d+)\]',p.text)[1]);new=mapping[old]
        original=re.sub(r'^\[\d+\]',f'[{new}]',p.text)
        assert entries[new-1]==springer_reference(original)
        assert re.findall(r'https?://\S+',original)==re.findall(r'https?://\S+',entries[new-1]),'Reference link changed'
    assert len(out.sections)==1
    for path,h in source_hashes.items():assert digest(path)==h
    with ZipFile(REFERENCE) as z:
        for p in [TEMPLATE,PAPER]:
            with ZipFile(p) as o:
                assert all(z.read(n)==o.read(n) for n in PRESERVE)
    result={'passed':True,'source_hashes_unchanged':source_hashes,'citation_old_to_new':mapping,'tables':16,'figures':11,'references':31,'body_text_preserved':True,'table_cells_preserved':True,'image_bytes_preserved':True,'opaque_parts_preserved':PRESERVE,'outputs':{str(p):digest(p) for p in [TEMPLATE,PAPER]}}
    (QA/'verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps({'passed':True,'outputs':list(result['outputs']),'tables':16,'figures':11,'references':31}))


if __name__=='__main__':
    source_hashes={str(p):digest(p) for p in [REFERENCE,GUIDE,*sorted((ROOT/'paper').glob('*'))] if p.is_file()}
    mapping=paper()
    template()
    verify(mapping,source_hashes)
