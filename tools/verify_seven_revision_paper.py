"""Check scientific evidence and synchronized publication artifacts, not old geometry."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import re
import statistics
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
import pymupdf

ROOT=Path(__file__).resolve().parents[1]
TITLE="Adaptive Merkle Trie Versioning for Efficient Structured Data Management"
ROMANS={"I":1,"II":2,"III":3,"IV":4,"V":5,"VI":6,"VII":7,"VIII":8,"IX":9,"X":10}

def normalize(text):
    text=text.replace('Tables VI and VII','Tables 6 and 7')
    text=re.sub(r"\b(?:TABLE|Table) (VIII|VII|III|VI|IV|IX|II|V|X|I)\b",lambda m:"Table "+str(ROMANS[m[1]]),text)
    text=re.sub(r"^([IVX]+)\. ",lambda m:str(ROMANS[m[1]])+" ",text)
    text=re.sub(r"^[A-Z]\. ","",text)
    text=re.sub(r"(Fig\. \d+|Table \d+)\. ",r"\1 ",text)
    return " ".join(text.lower().split()).rstrip(".")

def content(doc):
    result=[]
    for element in doc.element.body:
        if element.tag==qn("w:p"):
            p=Paragraph(element,doc)
            if p.text.strip():result.append(("p",normalize(p.text)))
        elif element.tag==qn("w:tbl"):
            table=Table(element,doc)
            result.append(("table",[[normalize(c.text) for c in row.cells] for row in table.rows]))
    return result

def check_doc(path,discover):
    doc=Document(path)
    assert " ".join(doc.paragraphs[0].text.split())==TITLE
    abstract=next(p.text for p in doc.paragraphs if p.text.startswith("Abstract:"))
    assert len(abstract.split())-1<250
    assert len(doc.tables)==10 and len(doc.inline_shapes)==11
    numbers=[int(m[1]) for p in doc.paragraphs if (m:=re.match(r"^\[(\d+)\]",p.text))]
    assert numbers==list(range(1,32))
    ref_index=next(i for i,p in enumerate(doc.paragraphs) if p.text.lower()=="references")
    body="\n".join(p.text for p in doc.paragraphs[:ref_index])+"\n"+"\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    citations=set()
    for match in re.finditer(r"\[([\d, -]+)\]",body):
        for chunk in match[1].split(','):
            bounds=[int(x) for x in chunk.strip().split('-')]
            citations.update(range(bounds[0],bounds[-1]+1))
    assert citations==set(range(1,32)),citations
    figures=[int(m[1]) for p in doc.paragraphs if (m:=re.match(r"^Fig\. (\d+)",p.text))]
    assert figures==list(range(1,12)),figures
    tables=[p.text for p in doc.paragraphs if p.style.name=='Caption' and re.match(r"^(?:TABLE [IVX]+\.|Table \d+ )",p.text)]
    assert len(tables)==10
    assert not doc.element.xpath(".//w:pBdr") and not doc.styles.element.xpath(".//w:pBdr"),"title-rule residue"
    assert "\u2014" not in body,"em dash remains"
    for text in ("geometry remains fixed","same physical computer","independent","changesets","16,384","4,096","unpaired","recovery","author-confirmed declarations have not yet been supplied"):
        assert text.lower() in body.lower(),text
    for table in doc.tables:
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"),"missing repeated header"
        assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in table.rows),"split row enabled"
    if discover:
        for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
            for r in p.runs:
                size=r.font.size or p.style.font.size or doc.styles['Normal'].font.size
                assert size is None or size.pt>=12,(p.text[:60],size.pt)
        assert all(s._sectPr.find(qn("w:cols")) is None or s._sectPr.find(qn("w:cols")).get(qn("w:num"),"1")=="1" for s in doc.sections)
    else:
        assert any(s._sectPr.find(qn("w:cols")).get(qn("w:num"),"1")=="2" for s in doc.sections)
    return doc,dict(path=str(path),paragraphs=len(doc.paragraphs),tables=10,figures=11,references=31,abstract_words=len(abstract.split())-1)

def check_evidence():
    # Retain the existing publication gate for supplemental process telemetry.
    telemetry=ROOT/'evidence/paper-issues14-windows-telemetry-20260924'
    telemetry_audit=json.loads((telemetry/'audit.json').read_text())
    assert telemetry_audit['passed'] and not telemetry_audit['issues'] and telemetry_audit['all_correct']
    assert (telemetry_audit['rows'],telemetry_audit['measured_rows'])==(405,315)
    with (telemetry/'raw_results.csv').open(newline='',encoding='utf-8') as f:telemetry_rows=list(csv.DictReader(f))
    assert len(telemetry_rows)==405
    assert all(r['status']=='ok' and r['correctness']=='True' for r in telemetry_rows)
    for counter in ('process_tree_peak_rss_bytes','process_tree_cpu_seconds','process_tree_read_bytes','process_tree_write_bytes','commit_operations_per_second'):
        assert all(r[counter] for r in telemetry_rows),counter
    primary=ROOT/'evidence/paper-corrected-issues13-17-counterbalanced-final-20260923'
    public=ROOT/'evidence/public-tlc-combined-20261002'
    for directory,total,successful,measured in ((primary,540,540,420),(public,325,324,252)):
        audit=json.loads((directory/'evidence_audit.json').read_text())
        assert audit['passed'] and not audit['issues']
        with (directory/'raw_results.csv').open(newline='',encoding='utf-8') as f:raw=list(csv.DictReader(f))
        assert len(raw)==total
        ok=[r for r in raw if r['status']=='ok']
        assert len(ok)==successful and all(r['correctness']=='True' for r in ok)
        assert sum(r['trial_kind']=='measured' for r in ok)==measured
        for name,digest in audit.get('evidence_files_sha256',{}).items():
            assert hashlib.sha256((directory/name).read_bytes()).hexdigest()==digest,name
        with (directory/'summary.csv').open(newline='',encoding='utf-8') as f:summary=list(csv.DictReader(f))
        for row in summary:
            trials=[r for r in ok if r['trial_kind']=='measured' and r['scenario']==row['scenario'] and r['model']==row['model'] and r['phase']==row['phase']]
            assert len(trials)==7,(row['scenario'],row['model'])
            for metric in ('diff_ms','incremental_commit_ms','checkout_ms','storage_bytes'):
                assert abs(statistics.median(float(r[metric]) for r in trials)-float(row[metric+'_median']))<1e-7,(row['scenario'],metric)
    with (primary/'raw_results.csv').open(newline='',encoding='utf-8') as f:primary_rows=list(csv.DictReader(f))
    compacted={}
    for model,expected in (('Revon-M (forced Merkle)',22.08),('Revon-H',22.08),('Dolt',2.17),('Dolt (bulk import)',2.19)):
        samples=[float(r['storage_bytes_after_compaction'])/2**20 for r in primary_rows if r['phase']=='evaluation' and r['trial_kind']=='measured' and r['model']==model and r['storage_bytes_after_compaction']]
        assert len(samples)==8 and round(statistics.mean(samples),2)==expected,(model,samples)
        compacted[model]=dict(samples=len(samples),mean_mib=statistics.mean(samples))
    base=ROOT/"evidence/threshold-heldout-20261002";extension=ROOT/"evidence/crossover-diagnostic-20261002"
    for directory,cases,queries,measured in ((base,42,2646,2058),(extension,5,315,245)):
        manifest=json.loads((directory/"manifest.json").read_text());audit=json.loads((directory/"audit.json").read_text())
        assert audit['passed'] and (audit['cases'],audit['queries'],audit['measured'])==(cases,queries,measured)
        for name,digest in manifest['source_code_sha256'].items():
            assert hashlib.sha256((directory/'source'/name.replace('\\','/')).read_bytes()).hexdigest()==digest,name
        results=[json.loads(p.read_text()) for p in (directory/'cases').glob('*/result.json')]
        assert len(results)==cases
        for result in results:
            assert result['semantic_checks']
            assert len(result['records'])==63
            assert len({r['output_sha256'] for r in result['records']})==1,"paths produce differing canonical output"
            assert all(r['correctness'] for r in result['records'])
            assert all(r['selected']==('log' if r['operations']<=r['threshold'] else 'merkle') for r in result['records'] if r['strategy']=='hybrid')
            assert sum(r['payload_bytes'] for r in result['components'])<result['storage_bytes'],"payload accounting exceeds whole-file bytes"
        with (directory/'summary.csv').open(newline='') as f:summary=list(csv.DictReader(f))
        for row in summary:
            result=next(x for x in results if x['spec']['name']==row['scenario'])
            measurements=[r['diff_ms'] for r in result['records'] if r['trial_kind']=='measured' and r['strategy']==row['strategy'] and r['threshold']==int(row['threshold'])]
            assert len(measurements)==7 and abs(statistics.median(measurements)-float(row['median_ms']))<1e-8
            quartiles=statistics.quantiles(measurements,n=4,method='inclusive')
            assert abs(quartiles[0]-float(row['p25_ms']))<1e-8
            assert abs(quartiles[2]-float(row['p75_ms']))<1e-8
    representative=json.loads((base/'cases/validation-n100000-h64-spread-o32768/result.json').read_text())
    payloads={r['kind']:r['payload_bytes'] for r in representative['components']}
    assert round(representative['storage_bytes']/2**20,2)==188.59
    assert {kind:round(value/2**20,2) for kind,value in payloads.items()}=={'node':60.34,'changeset':15.21,'commit':0.02}
    assert round((representative['storage_bytes']-sum(payloads.values()))/2**20,2)==113.02
    frozen=json.loads((base/'frozen_selection.json').read_text())
    assert frozen['threshold']==16384
    assert not frozen['validation_started']
    archived=(base/'source/experiments/threshold_validation.py').read_text()
    assert archived.index('freeze_path.write_text')<archived.index('for profile in profiles',archived.index('for split in'))
    # The archived execution protocol writes the frozen choice before running
    # validation. Per-case UTC timestamps were not recorded in this supplement;
    # filesystem mtimes cannot establish chronology after a Git checkout.
    return dict(telemetry_rows=405,primary_executions=540,public_attempts=325,public_successful=324,main_cases=42,main_queries=2646,main_measured=2058,diagnostic_cases=5,diagnostic_queries=315,diagnostic_measured=245,frozen_threshold=16384,compacted=compacted)

def check_pdf(path):
    pdf=pymupdf.open(path);captions={};faults=[];word_counts=[]
    for i,page in enumerate(pdf,1):
        text=page.get_text();word_counts.append(len(text.split()))
        images=[pymupdf.Rect(b['bbox']) for b in page.get_text('dict')['blocks'] if b['type']==1]
        for block in page.get_text('dict')['blocks']:
            if block['type']!=0:continue
            for line in block['lines']:
                line_text=''.join(s['text'] for s in line['spans']).strip()
                match=re.match(r'^Fig\.\s*(\d+)',line_text)
                if match:
                    number=int(match[1]);captions[number]=i
                    if not images:faults.append(f'Figure {number} caption on page {i} without image')
                for span in line['spans']:
                    rect=pymupdf.Rect(span['bbox'])
                    if not page.rect.contains(rect):faults.append(f'page {i} text outside page: {span["text"][:60]}')
        for rect in images:
            if not page.rect.contains(rect):faults.append(f'page {i} image outside page')
        if len(text.split())<8:faults.append(f'empty or near-empty page {i}')
    assert sorted(captions)==list(range(1,12)),captions
    assert not faults,faults
    text=' '.join(' '.join(p.get_text().split()) for p in pdf)
    for token in ('1,446.69','4,764.26','16,384','131,072','31]'):
        assert token in text,token
    return dict(path=str(path),pages=len(pdf),figure_caption_pages=captions,word_counts=word_counts,geometry_faults=faults)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--render-root',type=Path);parser.add_argument('--output',type=Path)
    args=parser.parse_args();root=args.render_root
    ieee,ic=check_doc(ROOT/'paper/Revon_Research_Paper_IEEE.docx',False)
    discover,dc=check_doc(ROOT/'paper/Revon_Research_Paper_Discover_Computing.docx',True)
    assert content(ieee)==content(discover),'journal and IEEE scientific content differs'
    source=ROOT/'paper/older versions/sources/Revon_Reviewed_Public_Dataset_Manuscript.docx'
    assert hashlib.sha256(source.read_bytes()).hexdigest()=='22b8c68e42f2f305c1c43d095dbe719556ababa95c38722fed184c85d871b8d3','archived authoritative source changed'
    evidence=check_evidence()
    ieee_render='render-ieee-word' if root and (root/'render-ieee-word/Revon_Research_Paper_IEEE.pdf').exists() else 'render-ieee'
    paths=[root/ieee_render/'Revon_Research_Paper_IEEE.pdf',root/'render-discover/Revon_Research_Paper_Discover_Computing.pdf'] if root else [ROOT/'paper/Revon_Research_Paper_IEEE.pdf',ROOT/'paper/Revon_Research_Paper_Discover_Computing.pdf']
    pdfs=[check_pdf(p) for p in paths]
    review=root/'render-review/Revon_Discover_Computing_One_Page_Review.pdf' if root else ROOT/'output/docs/Revon_Discover_Computing_One_Page_Review.pdf'
    # Local review is intentionally ignored by Git and absent in CI.
    if review.exists():assert len(pymupdf.open(review))==1,'review is not one page'
    report=dict(passed=True,manuscripts=[ic,dc],evidence=evidence,pdfs=pdfs,limits='structural/content audit; rendered visual inspection remains required')
    if args.output:
        args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2))

if __name__=='__main__':main()
